#!/usr/bin/env python3
"""独立复核第2步数据审计的关键事实与冻结产物。

# 输入: data/附件1.xlsx, data/附件2.xlsx, output/data_audit/tables/*.csv
# 输出: output/data_audit/verification.json
# 职责: 不调用主审计脚本，独立读取原表并核对记录、公式、配方与产物哈希。
"""

from __future__ import annotations

import hashlib
import json
import math
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
DATA = ROOT / "data"
OUT = ROOT / "output" / "data_audit"
TABLES = OUT / "tables"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def check(name: str, passed: bool, observed: object, expected: object) -> dict[str, object]:
    return {
        "check": name,
        "passed": bool(passed),
        "observed": observed,
        "expected": expected,
    }


def main() -> None:
    att1 = DATA / "附件1.xlsx"
    att2 = DATA / "附件2.xlsx"
    clean1_path = TABLES / "cleaned_attachment1.csv"
    clean2_path = TABLES / "cleaned_attachment2.csv"
    recipes_path = TABLES / "catalyst_combinations.csv"
    manifest_path = OUT / "artifact_manifest.csv"

    wb1 = load_workbook(att1, read_only=True, data_only=True)
    rows1 = list(wb1.active.iter_rows(values_only=True))
    raw1 = []
    current_id = None
    for row_number, row in enumerate(rows1[1:], 2):
        if row[0] is not None:
            current_id = str(row[0]).strip()
        if row[2] is not None:
            raw1.append({
                "source_excel_row": row_number,
                "catalyst_id": current_id,
                "temperature_c": float(row[2]),
                "ethanol_conversion_pct": float(row[3]),
                "ethylene_selectivity_pct": float(row[4]),
                "c4_selectivity_pct": float(row[5]),
                "acetaldehyde_selectivity_pct": float(row[6]),
                "c4_12_alcohol_selectivity_pct": float(row[7]),
                "methylbenzaldehyde_methylbenzyl_alcohol_selectivity_pct": float(row[8]),
                "other_selectivity_pct": float(row[9]),
                "selectivity_sum_pct": sum(float(value) for value in row[4:10]),
            })

    wb2 = load_workbook(att2, read_only=True, data_only=True)
    rows2 = list(wb2.active.iter_rows(values_only=True))
    raw2 = [
        {
            "source_excel_row": row_number,
            "time_min": float(row[0]),
            "ethanol_conversion_pct": float(row[1]),
            "ethylene_selectivity_pct": float(row[2]),
            "c4_selectivity_pct": float(row[3]),
            "acetaldehyde_selectivity_pct": float(row[4]),
            "c4_12_alcohol_selectivity_pct": float(row[5]),
            "methylbenzaldehyde_methylbenzyl_alcohol_selectivity_pct": float(row[6]),
            "other_selectivity_pct": float(row[7]),
            "selectivity_sum_pct": sum(float(value) for value in row[2:8]),
        }
        for row_number, row in enumerate(rows2[3:], 4)
        if row[0] is not None
    ]

    clean1 = pd.read_csv(clean1_path)
    clean2 = pd.read_csv(clean2_path)
    recipes = pd.read_csv(recipes_path)
    manifest = pd.read_csv(manifest_path)
    modeling = pd.read_csv(TABLES / "modeling_dataset.csv")
    domains = pd.read_csv(TABLES / "field_domain_audit.csv")
    reconciliation = pd.read_csv(TABLES / "raw_clean_reconciliation.csv")
    recipe_review = pd.read_csv(TABLES / "recipe_parse_verification.csv")
    coverage_matrix = pd.read_csv(TABLES / "combination_temperature_coverage_matrix.csv")
    checks: list[dict[str, object]] = []

    checks.extend([
        check("附件1有效记录数", len(raw1) == len(clean1) == 114,
              {"raw": len(raw1), "clean": len(clean1)}, 114),
        check("附件1组合数", clean1.catalyst_id.nunique() == 21,
              int(clean1.catalyst_id.nunique()), 21),
        check("附件2有效记录数", len(raw2) == len(clean2) == 7,
              {"raw": len(raw2), "clean": len(clean2)}, 7),
        check("附件1组合温度键无重复",
              not clean1.duplicated(["catalyst_id", "temperature_c"]).any(),
              int(clean1.duplicated(["catalyst_id", "temperature_c"]).sum()), 0),
        check("正式模型接口只含获批字段",
              modeling.columns.tolist() == [
                  "source_excel_row", "catalyst_id", "temperature_c",
                  "ethanol_conversion_pct", "c4_selectivity_pct", "c4_yield_pct",
                  "temperature_level", "temperature_order",
              ],
              modeling.columns.tolist(), "批准的8个接口字段"),
        check("催化剂域为21种完整组合",
              set(modeling.catalyst_id) == {
                  *(f"A{i}" for i in range(1, 15)),
                  *(f"B{i}" for i in range(1, 8)),
              },
              int(modeling.catalyst_id.nunique()), 21),
        check("温度域为7个有序离散标签",
              sorted(modeling.temperature_c.unique().tolist())
              == [250, 275, 300, 325, 350, 400, 450],
              sorted(modeling.temperature_c.unique().tolist()),
              [250, 275, 300, 325, 350, 400, 450]),
    ])

    raw1_df = pd.DataFrame(raw1).sort_values("source_excel_row").reset_index(drop=True)
    clean1_ordered = clean1.sort_values("source_excel_row").reset_index(drop=True)
    yield_error = (
        clean1_ordered.ethanol_conversion_pct
        * clean1_ordered.c4_selectivity_pct / 100
        - clean1_ordered.c4_yield_pct
    ).abs().max()
    raw_clean_key_match = (
        raw1_df.catalyst_id.tolist() == clean1_ordered.catalyst_id.tolist()
        and raw1_df.temperature_c.tolist() == clean1_ordered.temperature_c.tolist()
    )
    checks.extend([
        check("附件1原表与清洁表行键一致", raw_clean_key_match,
              raw_clean_key_match, True),
        check("C4收率公式最大误差", yield_error < 1e-12,
              float(yield_error), "< 1e-12"),
        check("附件1选择性闭合最大误差",
              max(abs(value - 100) for value in raw1_df.selectivity_sum_pct) < 1e-10,
              float(max(abs(value - 100) for value in raw1_df.selectivity_sum_pct)),
              "< 1e-10个百分点"),
        check("附件2选择性闭合最大误差",
              max(abs(row["selectivity_sum_pct"] - 100) for row in raw2) < 1e-10,
              float(max(abs(row["selectivity_sum_pct"] - 100) for row in raw2)),
              "< 1e-10个百分点"),
    ])

    raw_fields1 = [
        "temperature_c", "ethanol_conversion_pct", "ethylene_selectivity_pct",
        "c4_selectivity_pct", "acetaldehyde_selectivity_pct",
        "c4_12_alcohol_selectivity_pct",
        "methylbenzaldehyde_methylbenzyl_alcohol_selectivity_pct",
        "other_selectivity_pct",
    ]
    raw_fields2 = [
        "time_min", "ethanol_conversion_pct", "ethylene_selectivity_pct",
        "c4_selectivity_pct", "acetaldehyde_selectivity_pct",
        "c4_12_alcohol_selectivity_pct",
        "methylbenzaldehyde_methylbenzyl_alcohol_selectivity_pct",
        "other_selectivity_pct",
    ]
    raw2_df = pd.DataFrame(raw2).sort_values("source_excel_row").reset_index(drop=True)
    clean2_ordered = clean2.sort_values("source_excel_row").reset_index(drop=True)
    full_reconciliation_error = max(
        max((raw1_df[field] - clean1_ordered[field]).abs().max() for field in raw_fields1),
        max((raw2_df[field] - clean2_ordered[field]).abs().max() for field in raw_fields2),
    )
    checks.extend([
        check("独立逐字段原始—清洁对账",
              full_reconciliation_error < 1e-12,
              float(full_reconciliation_error), "< 1e-12"),
        check("主脚本对账表无不一致",
              int(reconciliation.mismatch_count.sum()) == 0,
              int(reconciliation.mismatch_count.sum()), 0),
        check("全部定义域无违规",
              int(domains.violation_count.sum()) == 0,
              int(domains.violation_count.sum()), 0),
        check("21种配方说明全部进入解析复核表",
              len(recipe_review) == 21 and recipe_review.catalyst_id.nunique() == 21,
              {"rows": len(recipe_review),
               "combinations": int(recipe_review.catalyst_id.nunique())}, 21),
        check("21乘7覆盖矩阵与114条记录一致",
              coverage_matrix.shape == (21, 8)
              and int(coverage_matrix.drop(columns="catalyst_id").to_numpy().sum()) == 114,
              {"shape": list(coverage_matrix.shape),
               "observed_cells": int(coverage_matrix.drop(
                   columns="catalyst_id").to_numpy().sum())},
              {"shape": [21, 8], "observed_cells": 114}),
        check("辅助诊断未标为正式模型输入",
              set(manifest.loc[
                  manifest.artifact_type == "auxiliary_diagnostic", "status"
              ]) == {"auxiliary"}
              and not manifest.loc[
                  manifest.artifact_type == "auxiliary_diagnostic", "consumer"
              ].str.contains("正式模型默认读取").any(),
              sorted(manifest.artifact_type.unique().tolist()),
              "辅助诊断全部status=auxiliary"),
    ])

    first = clean1_ordered.iloc[0]
    hand_value = 2.06716944958221 * 34.05 / 100
    checks.append(check(
        "首行手算收率A1-250°C",
        math.isclose(first.c4_yield_pct, hand_value, rel_tol=0, abs_tol=1e-12),
        float(first.c4_yield_pct), hand_value,
    ))

    recipe_by_id = recipes.set_index("catalyst_id")
    a11 = recipe_by_id.loc["A11"]
    checks.extend([
        check("A11为50mg Co/SiO2、无HAP且含90mg石英砂",
              a11.cosio2_mass_mg == 50
              and a11.hap_mass_mg == 0 and a11.quartz_mass_mg == 90,
              {"cosio2_mass_mg": float(a11.cosio2_mass_mg),
               "hap_mass_mg": float(a11.hap_mass_mg),
               "quartz_mass_mg": float(a11.quartz_mass_mg)},
              {"cosio2_mass_mg": 50, "hap_mass_mg": 0, "quartz_mass_mg": 90}),
        check("A12与B1配方数值相同但装料方式不同",
              recipe_by_id.loc["A12", [
                  "cosio2_mass_mg", "co_loading_wt_pct", "hap_mass_mg",
                  "ethanol_condition_ml_min",
              ]].equals(recipe_by_id.loc["B1", [
                  "cosio2_mass_mg", "co_loading_wt_pct", "hap_mass_mg",
                  "ethanol_condition_ml_min",
              ]])
              and recipe_by_id.loc["A12", "loading_method"]
              != recipe_by_id.loc["B1", "loading_method"],
              "A12/B1", "数值配方相同、装料方式不同"),
        check("A9与B5配方数值相同但装料方式不同",
              recipe_by_id.loc["A9", [
                  "cosio2_mass_mg", "co_loading_wt_pct", "hap_mass_mg",
                  "ethanol_condition_ml_min",
              ]].equals(recipe_by_id.loc["B5", [
                  "cosio2_mass_mg", "co_loading_wt_pct", "hap_mass_mg",
                  "ethanol_condition_ml_min",
              ]])
              and recipe_by_id.loc["A9", "loading_method"]
              != recipe_by_id.loc["B5", "loading_method"],
              "A9/B5", "数值配方相同、装料方式不同"),
    ])

    missing_artifacts = []
    hash_mismatches = []
    for row in manifest.itertuples(index=False):
        path = ROOT / row.path
        if not path.is_file():
            missing_artifacts.append(row.path)
        elif sha256(path) != row.sha256:
            hash_mismatches.append(row.path)
    checks.extend([
        check("产物清单路径完整", not missing_artifacts, missing_artifacts, []),
        check("产物清单SHA-256一致", not hash_mismatches, hash_mismatches, []),
    ])

    result = {
        "schema_version": "1.0",
        "generated_at": datetime.now(ZoneInfo("Asia/Shanghai")).isoformat(),
        "verifier": "src/data_audit/02_verify_audit.py",
        "independence": "未导入或调用主审计脚本，独立读取官方工作簿与冻结产物",
        "source_sha256": {
            "data/附件1.xlsx": sha256(att1),
            "data/附件2.xlsx": sha256(att2),
        },
        "checks": checks,
        "passed_count": sum(item["passed"] for item in checks),
        "check_count": len(checks),
        "all_passed": all(item["passed"] for item in checks),
    }
    output_path = OUT / "verification.json"
    output_path.write_text(
        json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    if not result["all_passed"]:
        raise SystemExit(json.dumps(result, ensure_ascii=False, indent=2))
    print(json.dumps({
        "status": "ok",
        "passed": result["passed_count"],
        "total": result["check_count"],
        "output": str(output_path.relative_to(ROOT)),
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
