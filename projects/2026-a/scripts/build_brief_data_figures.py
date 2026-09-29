"""Refresh two explanatory figures from existing observations and Q2 results.

No solver is imported or run. The three appendix-3 properties are evaluated
explicitly and checked against the existing half-hour diagnostic records.
"""
from __future__ import annotations

import hashlib
import json
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.colors import to_rgba
from matplotlib.lines import Line2D
from matplotlib.text import Text
import numpy as np
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output/paper_figures/brief_refresh_20260911"
Q2_REL = "output/Q2/run_n10240_startsafe.npz"
Q2_SHA = "79f4b32f3a8858980b4649ed71ad2d3924b1c0312dc3de6260fea87866e9a59b"
ENV_REL = "output/ENV/q1_boundary.json"
REVIEW_REL = "output/diagnostics/q1_q2_parameterization_review.json"
BLUE = "#0072BD"
GRAY = "#333333"
CORAL = "#C76556"
WIDTH_CM = 15.8
DPI = 300


def identity(path: Path) -> dict:
    return {"path": path.relative_to(ROOT).as_posix(),
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "bytes": path.stat().st_size}


def configure() -> None:
    font = Path("C:/Windows/Fonts/msyh.ttc")
    font_manager.fontManager.addfont(str(font))
    plt.rcParams.update({
        "font.family": font_manager.FontProperties(fname=str(font)).get_name(),
        "font.size": 8.5, "axes.labelsize": 8.5,
        "xtick.labelsize": 8, "ytick.labelsize": 8,
        "axes.linewidth": .8, "xtick.direction": "in",
        "ytick.direction": "in", "axes.unicode_minus": False,
        "pdf.fonttype": 42, "ps.fonttype": 42, "svg.fonttype": "none",
        "mathtext.fontset": "stix", "mathtext.default": "it",
        "figure.facecolor": "white",
        "savefig.facecolor": "white", "path.simplify": False,
    })


def gray_color(value):
    r, g, b, a = to_rgba(value)
    luminance = .2126*r + .7152*g + .0722*b
    return (luminance, luminance, luminance, a)


def save_bundle(fig, name: str, metadata: dict) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    fig.canvas.draw()
    outputs = {}
    for ext in ("png", "pdf", "svg"):
        path = OUT / f"{name}.{ext}"
        fig.savefig(path, dpi=DPI)
        outputs[ext] = identity(path)

    # Freeze all axes positions before hiding text: data geometry is identical.
    positions = [(ax, ax.get_position().frozen()) for ax in fig.axes]
    texts = list(fig.findobj(Text))
    visible = [item.get_visible() for item in texts]
    for item in texts:
        item.set_visible(False)
    for ax, position in positions:
        ax.set_position(position)
    path = OUT / f"{name}_no_text.png"
    fig.savefig(path, dpi=DPI)
    outputs["no_text_png"] = identity(path)
    for item, old in zip(texts, visible):
        item.set_visible(old)

    # Render a grayscale version from the same artists, without editing pixels.
    for item in fig.findobj(Line2D):
        item.set_color(gray_color(item.get_color()))
        for getter, setter in [(item.get_markeredgecolor, item.set_markeredgecolor),
                               (item.get_markerfacecolor, item.set_markerfacecolor)]:
            try:
                setter(gray_color(getter()))
            except (TypeError, ValueError):
                pass
    for item in fig.findobj(Text):
        item.set_color(gray_color(item.get_color()))
    path = OUT / f"{name}_gray.png"
    fig.savefig(path, dpi=DPI)
    outputs["gray_png"] = identity(path)
    metadata.update({
        "schema_version": "1.0", "data_driven_visual": True,
        "generator": identity(Path(__file__)), "outputs": outputs,
        "display": {"width_cm": WIDTH_CM,
                    "height_cm": float(fig.get_figheight()*2.54),
                    "dpi": DPI, "font": "Microsoft YaHei",
                    "font_size_pt": 8.5, "main_linewidth_pt": 1.8,
                    "textfree": "same artist positions, all Text hidden",
                    "grayscale": "same artists with luminance-mapped line colors",
                    "extra_fit": False, "solver_executed": False},
    })
    (OUT/f"{name}_sources.json").write_text(
        json.dumps(metadata, ensure_ascii=False, indent=2)+"\n", encoding="utf-8")
    plt.close(fig)


def environment() -> None:
    env = json.loads((ROOT/ENV_REL).read_text(encoding="utf-8"))
    original = ROOT/env["source"]["path"]
    assert identity(original)["sha256"] == env["source"]["sha256"]
    nodes = env["nodes"]
    times = np.array([v["time_s"] for v in nodes], dtype=float)
    temp = np.array([v["temperature_C"] for v in nodes])
    moist = np.array([v["air_moisture_kg_kg"] for v in nodes])
    assert np.array_equal(times, np.arange(0, 1801, 60))
    book = load_workbook(original, read_only=True, data_only=True)
    rows = np.array(list(book[env["source"]["sheet"]].iter_rows(
        min_row=2, max_row=32, min_col=1, max_col=3, values_only=True)), dtype=float)
    book.close()
    np.testing.assert_allclose(rows, np.column_stack([times, temp, moist]), rtol=0, atol=0)
    fig, axes = plt.subplots(1, 2, figsize=(WIDTH_CM/2.54, 7.6/2.54))
    fig.subplots_adjust(left=.10, right=.99, bottom=.18, top=.89, wspace=.33)
    for j, (ax, values, color) in enumerate(zip(axes, [temp, moist], [CORAL, BLUE])):
        ax.plot(times/60, values, color=color, lw=1.8,
                label="分段线性重建", zorder=2)
        ax.plot(times/60, values, ls="none", marker="o", ms=3,
                color=GRAY, mfc="white", mew=.8, label="原始观测", zorder=3)
        ax.set_xlim(0, 30)
        ax.set_xticks([0, 10, 20, 30])
        ax.set_xlabel("时间 / min")
        if j == 0:
            ax.set_ylim(27.5, 42.3)
            ax.set_yticks([28, 32, 36, 40])
            ax.set_ylabel("烘\n房\n温\n度\n°C", rotation=0,
                          ha="center", va="center", labelpad=19)
        else:
            ax.set_ylim(.019, .0338)
            ax.set_yticks([.020, .024, .028, .032])
            ax.set_ylabel("空\n气\n水\n分\n浓\n度\n(kg/kg)", rotation=0,
                          ha="center", va="center", labelpad=24)
        ax.text(.0, 1.035, f"({'ab'[j]})", transform=ax.transAxes,
                weight="bold", fontsize=9.5, va="bottom")
        ax.grid(axis="y", color="#DBDFE2", linewidth=.55)
        ax.legend(loc="upper left", frameon=False, fontsize=7.4,
                  handlelength=2.4, borderaxespad=.4)

        inset = ax.inset_axes([.47, .13, .35, .32])
        selected = (times >= 120) & (times <= 240)
        local = values[selected]
        margin = (local.max()-local.min())*.20
        inset.plot(times[selected], local, color=color, lw=1.5)
        inset.plot(times[selected], local, ls="none", marker="o", ms=3.2,
                   mec=GRAY, mfc="white", mew=.8)
        inset.set(xlim=(110, 250), ylim=(local.min()-margin, local.max()+margin))
        inset.set_xticks([120, 180, 240])
        inset.set_yticks([local[0], local[-1]])
        inset.tick_params(labelsize=6.8, pad=2)
        inset.tick_params(axis="y", labelleft=False, labelright=True,
                          left=False, right=True)
        inset.set_xlabel("时间 / s", fontsize=7, labelpad=1)
        inset.text(.5, 1.03, "局部放大", ha="center", va="bottom",
                   transform=inset.transAxes, fontsize=7)
        inset.grid(color="#E7E9EB", lw=.45)
        inset.yaxis.set_major_formatter(plt.FuncFormatter(
            (lambda v, p: f"{v:.2f}") if j == 0 else (lambda v, p: f"{v:.5f}")))
    save_bundle(fig, "fig02_q1_environment", {
        "title": "问题1环境观测与分段线性重建",
        "caption": "(a)烘房温度；(b)空气水分浓度。圆点为0～1800 s的31个原始观测，实线连接相邻观测；嵌图放大120～240 s的三个原始节点及其连接线。只表示本问使用的环境输入，未平滑、拟合或外推。",
        "sources": [identity(ROOT/ENV_REL), identity(original)],
        "checks": {"all_31_rows_equal_original_attachment": True},
        "data": {"time_s": [0, 1800], "node_count": 31,
                 "time_unit_main": "min", "time_unit_insets": "s",
                 "inset_time_s": [120, 180, 240],
                 "interpolation": "piecewise linear between original observations"},
        "claim_boundary": "Q1 observed environment only; no Q2/Q3 extrapolation or internal material state is shown.",
    })


def coefficients() -> None:
    source = ROOT/Q2_REL
    assert identity(source)["sha256"] == Q2_SHA
    with np.load(source) as z:
        mask = (z["time_s"] >= 0) & (z["time_s"] <= 10800)
        time = z["time_s"][mask].copy()
        C = z["moisture_kg_kg"][mask][:, [0, -1]].copy()
        T = z["temperature_C"][mask][:, [0, -1]].copy()
        assert np.allclose(z["radius_m"][[0, -1]], [0, .02], atol=0, rtol=0)
    assert np.array_equal(time, np.arange(10801))
    assert np.isfinite(T).all() and (C > 0).all()
    # Explicit local appendix-3 formulas, independent of the solver module.
    k = .21 + .38*C/(1+C)
    s = (650+128*C)*(1450+2736*C/(1+C))
    D = 2.4e-3*np.exp(-.45/C-3850/(T+273.15))
    properties = {"k_W_mK": k, "s_J_m3K": s, "D_m2_s": D}
    ratios = {key: value/value[0] for key, value in properties.items()}
    review = json.loads((ROOT/REVIEW_REL).read_text(encoding="utf-8"))
    assert review["sources"]["q2_fields"]["sha256"] == Q2_SHA
    rows = review["q2_local_properties"]["rows"]
    assert len(rows) == 14
    max_rel = 0.
    for row in rows:
        i, j = int(row["time_s"]), {"center": 0, "surface": 1}[row["position"]]
        for key, values in properties.items():
            ref = row[key]
            max_rel = max(max_rel, abs(float(values[i, j])-ref)/abs(ref))
            np.testing.assert_allclose(values[i, j], ref, rtol=2e-14, atol=0)
            np.testing.assert_allclose(ratios[key][i, j], row["ratio_to_initial"][key],
                                       rtol=2e-14, atol=0)
        np.testing.assert_allclose([T[i, j], C[i, j]],
                                   [row["temperature_C"], row["moisture_kg_kg"]],
                                   rtol=0, atol=0)
    fig, axes = plt.subplots(1, 3, figsize=(WIDTH_CM/2.54, 6.8/2.54))
    fig.subplots_adjust(left=.065, right=.99, bottom=.225, top=.85, wspace=.36)
    for j, (ax, key, label) in enumerate(zip(axes, properties,
                                           [r"$k/k_0$", r"$s/s_0$", r"$D/D_0$"])):
        values = ratios[key]
        ax.plot(time/3600, values[:, 0], color=BLUE, lw=1.8, ls="-", label="中心")
        ax.plot(time/3600, values[:, 1], color=GRAY, lw=1.8,
                ls=(0, (4, 2.4)), label="表面")
        ax.axhline(1, color="#A4ABB0", ls=(0, (2, 2)), lw=.75, zorder=0)
        ax.set_xlim(0, 3)
        ax.set_xticks([0, 1, 2, 3])
        ax.set_xlabel("时间 / h")
        ax.text(0, 1.04, f"({'abc'[j]})", transform=ax.transAxes,
                fontsize=9.5, weight="bold", va="bottom")
        ax.text(.52, 1.04, label, transform=ax.transAxes, ha="center",
                fontsize=10, va="bottom")
        ax.grid(axis="y", color="#DFE3E6", linewidth=.55)
        if j == 0:
            ax.set_ylim(.80, 1.015); ax.set_yticks([.80, .85, .90, .95, 1.00])
        elif j == 1:
            ax.set_ylim(.62, 1.025); ax.set_yticks([.7, .8, .9, 1.0])
        else:
            ax.set_ylim(.86, 2.34); ax.set_yticks([1.0, 1.4, 1.8, 2.2])
    fig.legend(*axes[0].get_legend_handles_labels(), loc="lower center", ncol=2,
               bbox_to_anchor=(.53, -.006), frameon=False, fontsize=8,
               handlelength=3, columnspacing=2.8)
    summary = {}
    for key, values in ratios.items():
        summary[key] = {}
        for j, position in enumerate(["center", "surface"]):
            curve = values[:, j]
            summary[key][position] = {"initial": float(properties[key][0, j]),
                                      "at_3h_ratio": float(curve[-1]),
                                      "minimum_ratio": float(curve.min()),
                                      "minimum_time_s": int(time[curve.argmin()]),
                                      "maximum_ratio": float(curve.max()),
                                      "maximum_time_s": int(time[curve.argmax()])}
    save_bundle(fig, "fig09_q2_coefficients", {
        "title": "问题2中心与表面局部物性的相对变化",
        "caption": "(a)导热系数；(b)体积热容量；(c)水分扩散系数。各量除以题给均匀初态下的对应初值，横向细虚线表示初值比1；蓝色实线为中心，深灰虚线为表面。逐秒温度与含水率来自问题2同一主方案，物性按附录3逐点计算，未另作拟合。三面板纵轴范围不同。",
        "sources": [identity(source), identity(ROOT/REVIEW_REL), identity(ROOT/"A题.pdf")],
        "formulas": {"k": ".21+.38*C/(1+C)",
                     "s": "(650+128*C)*(1450+2736*C/(1+C))",
                     "D": "2.4e-3*exp(-.45/C-3850/(temperature_C+273.15))"},
        "data": {"time_s": [0, 10800], "time_count": len(time),
                 "radius_m": [0, .02], "normalization": "property/local uniform initial value",
                 "shown_curves": "center and surface only; not full-domain bounds"},
        "checks": {"independent_half_hour_rows": len(rows),
                   "maximum_property_relative_difference": max_rel,
                   "comparison_rtol": 2e-14, "all_rows_match": True},
        "summary": summary,
        "claim_boundary": "Actual local-property trajectories over Q2 0–3 h only. No frozen-coefficient control, stage partition, global monotonicity proof, sensitivity experiment, or new solver run is implied.",
    })


def main() -> None:
    configure()
    environment()
    coefficients()
    print("Rendered Q1 environment and Q2 local-property figures; original data and 14 diagnostic rows agree.")


if __name__ == "__main__":
    main()
