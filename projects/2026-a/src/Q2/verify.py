"""Independent verification for the frozen Q2-v1 mathematical contract.

The production System(n, boundary).rhs/jac is an object under test. All expected
coefficients, face fluxes, geometry, continuous manufactured forcing and Bessel
series below are independent of production helpers. No solver results are made
by importing this module. The eventual authorized CLI executes the small
controlled examples, reads already completed full-process runs, and writes one
report. It does not produce the primary answer or export result2.xlsx.

Production interface:
  System(n, boundary), boundary.evaluate(t) -> (temperature_C, moisture_kg_kg);
  state is (theta_0,C_0,...,theta_N,C_N), with no diagnostic state appended.
  A runs-config gives spatial_runs, tight_run and sensitivity_run descriptors,
  each with npz and metadata paths relative to the config, or absolute paths.
  Settings and endpoints are taken from the actual metadata, never filenames.
  Optional checks_report lists earlier reports with the same selected axes;
  their current source identities are checked before their evidence is reused.
  NPZ loads one field at a time, and comparisons then proceed in row blocks.

--checks-only --axes E1,E2 runs implementation and small analytic/MMS examples.
--checks-only --axes E3 runs the instantaneous structural checks only.
--runs-config ... --axes E1,E2 reads full-process grid/time results only.
--runs-config ... --axes E3 reads long-time structure and sensitivity only.
Axes select actual execution, rather than filtering an already executed suite.

Prespecified project targets: spatial T/C differences <=2.5e-5; temporal
differences <=5e-6; spatial crossing difference <=1s, temporal <=0.1s.
Integer endpoint agreement is reported separately, never inferred from these.
No per-cell correct-rounding or empirical material-validation claim is made.
"""
from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import sys
from pathlib import Path
from datetime import datetime, timezone

import numpy as np
from openpyxl import load_workbook
from scipy.integrate import solve_ivp
from scipy.optimize import brentq
from scipy.special import j0, j1, jn_zeros

R, H, HM = 0.02, 25.0, 8e-7
ROOT = Path(__file__).resolve().parents[2]
SPEC_SHA = "10d2ce44107c11c835ccf173c820e4100648f7c48fd7955c3890038a20f688a4"
BOUNDARY_SHA = "f30996b8fc3a132164ae1a8b688f683bd15120fa38d54155221d8839085f178f"
T0, C0, THRESHOLD = 28.0, 2.55, 0.15
TAIL_T, TAIL_C = 49.99591666666667, 0.04999041666666667
LAST_T, LAST_C = 50.165, 0.04986
RAW_SHA = "7ef32870abeef420b89560b2530ff60dfe4255917805151d89988d0311af9dd7"
OUTPUT_RADII = np.arange(21, dtype=float) * 0.001
SUMMARY_TIMES = np.arange(1800, 10801, 1800, dtype=float)


def digest(path):
    value = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            value.update(chunk)
    return value.hexdigest()


class Report:
    def __init__(self, axes, mode):
        self.checks, self.details, self.identities = [], {}, {}
        self.axes, self.mode = sorted(axes), mode

    def add(self, axis, name, passed, **details):
        if axis not in self.axes:
            raise ValueError("This runtime entry point cannot execute evidence on axis " + axis)
        self.checks.append(dict(axis=axis, name=name, passed=bool(passed), **details))

    def metric(self, axis, name, value, limit, **details):
        self.add(axis, name, np.isfinite(value) and value <= limit,
                 value=float(value), limit=float(limit), **details)

    def result(self):
        return dict(schema_version="1.0", question="Q2", axes=self.axes, mode=self.mode,
                    method_role="auxiliary_validator",
                    created_at=datetime.now(timezone.utc).isoformat(),
                    passed=bool(self.checks) and all(c["passed"] for c in self.checks),
                    identities=self.identities, checks=self.checks, details=self.details,
                    claim_boundary=("E1 implementation, E2 numerical agreement and selected "
                                    "E3 structural properties of the fixed-grid Q2 model. "
                                    "No continuous-PDE global well-posedness, rigorous total "
                                    "error bound, per-cell correct rounding or E4 empirical "
                                    "validation is established by this report."))


class RawBoundary:
    """Direct authoritative worksheet read and independent scalar interpolation."""
    def __init__(self, path, scenario="mean_tail"):
        self.path, self.scenario = Path(path), scenario
        if scenario not in ("mean_tail", "last_value"):
            raise ValueError("Undefined long-time scenario")
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = list(book["Sheet1"].iter_rows(min_row=2, max_row=242,
                                               min_col=1, max_col=3, values_only=True))
        finally:
            book.close()
        self.data = np.array(rows, dtype=float)
        if self.data.shape != (241, 3) or not np.isfinite(self.data).all():
            raise ValueError("Expected 241 finite original boundary nodes")
        if not np.array_equal(self.data[:, 0], np.arange(0, 14401, 60)):
            raise ValueError("Original boundary time axis is not 0:60:14400 s")
        self.mean_tail = np.sum(30.0 * (self.data[180:240, 1:] +
                                      self.data[181:241, 1:]), axis=0) / 3600.0
        self.theta_bounds = (min(T0, float(self.data[:, 1].min())),
                             max(T0, float(self.data[:, 1].max())))
        self.moisture_bounds = (min(C0, float(self.data[:, 2].min())),
                                max(C0, float(self.data[:, 2].max())))

    def evaluate(self, t):
        if not np.isfinite(t) or t < 0:
            raise ValueError("Boundary time must be finite and nonnegative")
        if t > 14400:
            tail = self.mean_tail if self.scenario == "mean_tail" else self.data[-1, 1:]
            return float(tail[0]), float(tail[1])
        index = min(int(t // 60), 239)
        weight = (float(t) - self.data[index, 0]) / 60.0
        values = self.data[index, 1:] * (1 - weight) + self.data[index + 1, 1:] * weight
        return float(values[0]), float(values[1])


class ConstantBoundary:
    def __init__(self, temperature, moisture):
        self.values = float(temperature), float(moisture)

    def evaluate(self, t):
        return self.values


def geometry(n):
    """Independent explicit shell formulas, including the two truncated cells."""
    spacing = R / n
    radius = np.arange(n + 1, dtype=float) * spacing
    volume = radius * spacing
    volume[0] = spacing ** 2 / 8.0
    volume[-1] = R * spacing / 2.0 - spacing ** 2 / 8.0
    return radius, volume


def material(theta, concentration):
    """Analytic arithmetic supports complex-step differentiation in this oracle."""
    rho = 650.0 + 128.0 * concentration
    cp = 1450.0 + 2736.0 * concentration / (1.0 + concentration)
    capacity = rho * cp
    conductivity = 0.21 + 0.38 * concentration / (1.0 + concentration)
    diffusion = 2.4e-3 * np.exp(-0.45 / concentration - 3850.0 / (theta + 273.15))
    kp = 0.38 / (1.0 + concentration) ** 2
    sp = 128.0 * cp + rho * 2736.0 / (1.0 + concentration) ** 2
    dc = 0.45 * diffusion / concentration ** 2
    dt = 3850.0 * diffusion / (theta + 273.15) ** 2
    return capacity, conductivity, diffusion, sp, kp, dc, dt


def independent_rhs(n, boundary, t, y):
    """Explicit per-face assembly, never importing production geometry/material."""
    state = np.asarray(y).reshape(n + 1, 2)
    theta, concentration = state[:, 0], state[:, 1]
    _, volume = geometry(n)
    capacity, k, diffusion, *_ = material(theta, concentration)
    flux_sum = np.zeros_like(state)
    for i in range(n):
        face_factor = i + 0.5
        heat = face_factor * (k[i] + k[i + 1]) / 2 * (theta[i + 1] - theta[i])
        water = (face_factor * (diffusion[i] + diffusion[i + 1]) / 2 *
                 (concentration[i + 1] - concentration[i]))
        flux_sum[i, 0] += heat
        flux_sum[i + 1, 0] -= heat
        flux_sum[i, 1] += water
        flux_sum[i + 1, 1] -= water
    ambient_t, ambient_c = boundary.evaluate(t)
    flux_sum[-1, 0] += R * H * (ambient_t - theta[-1])
    flux_sum[-1, 1] += R * HM * (ambient_c - concentration[-1])
    flux_sum[:, 0] /= volume * capacity
    flux_sum[:, 1] /= volume
    return flux_sum.reshape(-1)


def interleave(theta, concentration):
    return np.column_stack((theta, concentration)).reshape(-1)


def relative_error(actual, expected, floor=1e-25):
    return float(np.max(np.abs(actual - expected)) /
                 max(float(np.max(np.abs(expected))), floor))


def check_boundary(report, raw, production_boundary, boundary_path):
    report.add("E1", "raw_workbook_identity", digest(raw.path) == RAW_SHA)
    interface = json.loads(Path(boundary_path).read_text(encoding="utf-8"))
    frozen_nodes = np.column_stack([interface[key] for key in
                                    ("time_s", "temperature_C", "air_moisture_kg_kg")])
    report.add("E1", "all_frozen_nodes_equal_authoritative_worksheet",
               frozen_nodes.shape == raw.data.shape and np.array_equal(frozen_nodes, raw.data))
    report.metric("E1", "last_hour_integral_temperature", abs(raw.mean_tail[0] - TAIL_T), 2e-13)
    report.metric("E1", "last_hour_integral_moisture", abs(raw.mean_tail[1] - TAIL_C), 2e-16)
    queries = np.r_[raw.data[:, 0], raw.data[:-1, 0] + 0.25,
                    raw.data[:-1, 0] + 59.75, np.nextafter(14400., np.inf),
                    14400.25, 14460., 1e6]
    errors = np.asarray([np.abs(np.asarray(production_boundary.evaluate(float(t))) -
                                np.asarray(raw.evaluate(float(t)))) for t in queries])
    report.metric("E1", "full_boundary_interpolation_temperature", errors[:, 0].max(), 2e-12)
    report.metric("E1", "full_boundary_interpolation_moisture", errors[:, 1].max(), 2e-14)
    for label, value in (("negative", -1.), ("nan", np.nan), ("inf", np.inf)):
        try:
            production_boundary.evaluate(value)
        except ValueError:
            passed = True
        else:
            passed = False
        report.add("E1", "boundary_rejects_" + label, passed)


def check_implementation(report, make_system, boundary):
    n, t = 40, 137.25
    radius, volume = geometry(n)
    x = radius / R
    theta = 35.0 + 9.0 * x ** 2 + 0.4 * np.cos(3 * np.pi * x)
    concentration = 0.4 + 1.9 * (1 - x ** 2) + 0.11 * np.cos(2 * np.pi * x)
    y = interleave(theta, concentration)
    system = make_system(n, boundary)
    report.metric("E1", "independent_radius_geometry", np.max(np.abs(system.r - radius)), 1e-15)
    report.metric("E1", "independent_volume_geometry", np.max(np.abs(system.v - volume)), 1e-18)
    report.metric("E1", "independent_face_factors", np.max(np.abs(system.q - (np.arange(n) + .5))), 1e-12)
    report.add("E1", "unchanged_surface_transfer_coefficients", system.h == H and system.hm == HM)
    expected = independent_rhs(n, boundary, t, y)
    actual = np.asarray(system.rhs(t, y))
    report.add("E1", "interleaved_rhs_shape", actual.shape == y.shape)
    for field, index in (("temperature", 0), ("moisture", 1)):
        report.metric("E1", field + "_rhs_independent_faces",
                      relative_error(actual[index::2], expected[index::2]), 2e-11)

    # A small complete matrix, split into all four blocks, avoids missing a
    # cross-field entry through cancellation in a single random direction.
    jac = system.jac(t, y)
    report.add("E1", "jacobian_sparse_CSC", getattr(jac, "format", None) == "csc")
    full_expected = np.empty((y.size, y.size))
    for column in range(y.size):
        perturbation = np.zeros(y.size, dtype=complex)
        perturbation[column] = 1e-24j
        full_expected[:, column] = np.imag(independent_rhs(n, boundary, t,
                                                          y.astype(complex) + perturbation)) / 1e-24
    full_actual = jac.toarray()
    for input_index, input_name in ((0, "theta"), (1, "C")):
        for output_index, output_name in ((0, "theta"), (1, "C")):
            report.metric("E1", f"complete_jac_{output_name}_by_{input_name}",
                          relative_error(full_actual[output_index::2, input_index::2],
                                         full_expected[output_index::2, input_index::2]),
                          2e-10, oracle="complex step of independent face assembly")
    directions = (np.cos(0.73 * np.arange(n + 1)) + 0.3,
                  np.sin(1.11 * np.arange(n + 1)) + 0.8)

    # A deliberate omission must be numerically visible in the selected test.
    capacity, k, diffusion, sp, kp, dc, dt = material(theta, concentration)
    p = directions[0]
    omitted_sp = -sp / capacity * expected[0::2] * p
    omitted_kp = np.zeros(n + 1)
    omitted_dc = np.zeros(n + 1)
    omitted_dt = np.zeros(n + 1)
    for i in range(n):
        q = i + 0.5
        increments = (q / 2 * (kp[i] * p[i] + kp[i + 1] * p[i + 1]) * (theta[i + 1] - theta[i]),
                      q / 2 * (dc[i] * p[i] + dc[i + 1] * p[i + 1]) * (concentration[i + 1] - concentration[i]),
                      q / 2 * (dt[i] * p[i] + dt[i + 1] * p[i + 1]) * (concentration[i + 1] - concentration[i]))
        for destination, increment in zip((omitted_kp, omitted_dc, omitted_dt), increments):
            destination[i] += increment
            destination[i + 1] -= increment
    omitted_kp /= volume * capacity
    omitted_dc /= volume
    omitted_dt /= volume
    report.details["derivative_omission_signals"] = {
        name: float(np.max(np.abs(term)))
        for name, term in (("s_prime", omitted_sp), ("k_prime", omitted_kp),
                           ("D_C", omitted_dc), ("D_theta", omitted_dt))}
    report.add("E1", "chosen_state_exercises_all_coupling_derivatives",
               all(value > 1e-12 for value in report.details["derivative_omission_signals"].values()))

    # Exact quadratic flux balance: every shell, including both truncated ones,
    # has theta_dot=4*k*A/(s*R**2). This is independent of FVM assembly code.
    a, c = 3.0, np.full(n + 1, C0)
    theta_quadratic = 35.0 + a * x ** 2
    s0, k0, *_ = material(35.0, C0)
    manufactured_ambient = ConstantBoundary(theta_quadratic[-1] + k0 * 2 * a / (R * H), C0)
    quadratic_system = make_system(n, manufactured_ambient)
    quadratic_rhs = np.asarray(quadratic_system.rhs(0.0, interleave(theta_quadratic, c)))
    scalar = 4 * k0 * a / (s0 * R ** 2)
    report.metric("E1", "exact_quadratic_all_shells_including_axis_surface",
                  float(np.max(np.abs(quadratic_rhs[0::2] - scalar))) / abs(scalar), 5e-10)
    report.metric("E1", "quadratic_temperature_has_zero_moisture_rhs",
                  np.max(np.abs(quadratic_rhs[1::2])), 1e-14)
    for label, index, value in (("zero_C", 1, 0.), ("negative_C", 1, -.1),
                                ("absolute_zero_temperature", 0, -273.15),
                                ("nan", 0, np.nan), ("infinity", 1, np.inf)):
        invalid = y.copy()
        invalid[index] = value
        try:
            system.rhs(t, invalid)
        except ValueError:
            rejected = True
        else:
            rejected = False
        report.add("E1", "rhs_rejects_" + label, rejected)


def check_structure(report, make_system, raw):
    n, t = 20, 16000.0
    radius, volume = geometry(n)
    x = radius / R
    theta = 29.0 + 16 * x ** 2
    concentration = 0.20 + 1.4 * (1 - x ** 2)
    boundary = ConstantBoundary(TAIL_T, TAIL_C)
    y = interleave(theta, concentration)
    system = make_system(n, boundary)
    f = np.asarray(system.rhs(t, y)).reshape(n + 1, 2)
    capacity, k, diffusion, *_ = material(theta, concentration)
    for name, weighted, outer in (
        ("water", volume * f[:, 1], R * HM * (TAIL_C - concentration[-1])),
        ("heat_power", volume * capacity * f[:, 0], R * H * (TAIL_T - theta[-1]))):
        scale = max(float(np.sum(np.abs(weighted))) + abs(outer), 1e-30)
        report.metric("E3", name + "_instantaneous_boundary_balance",
                      abs(float(np.sum(weighted)) - outer) / scale, 5e-12)

    z = concentration - TAIL_C
    lhs = float(np.dot(volume * z, f[:, 1]))
    dissipation = R * HM * z[-1] ** 2
    for i in range(n):
        dissipation += (i + 0.5) * (diffusion[i] + diffusion[i + 1]) / 2 * (z[i + 1] - z[i]) ** 2
    report.metric("E3", "water_variance_dissipation_with_nonuniform_temperature",
                  abs(lhs + dissipation) / max(dissipation, 1e-30), 5e-12)
    report.add("E3", "water_variance_dissipation_is_strict_here", lhs < 0)
    zero_flux_system = make_system(n, ConstantBoundary(theta[-1], concentration[-1]))
    zero_f = np.asarray(zero_flux_system.rhs(t, y)).reshape(n + 1, 2)
    for name, terms in (("water", volume * zero_f[:, 1]),
                        ("heat_power", volume * capacity * zero_f[:, 0])):
        report.metric("E3", "zero_surface_flux_instantaneous_" + name,
                      abs(float(np.sum(terms))) / max(float(np.sum(np.abs(terms))), 1e-30), 5e-12)

    equilibrium = interleave(np.full(n + 1, TAIL_T), np.full(n + 1, TAIL_C))
    report.metric("E3", "constant_equilibrium", np.max(np.abs(system.rhs(t, equilibrium))), 1e-14)

    # All coordinate facets are sampled on a nonuniform base state. The separate
    # mathematical proof, not these finite samples, establishes invariance.
    ranges = (raw.theta_bounds, raw.moisture_bounds)
    base = interleave(np.linspace(*ranges[0], n + 1), np.linspace(*ranges[1], n + 1))
    maximum_outward = 0.0
    for field in range(2):
        for i in range(n + 1):
            for side in range(2):
                state = base.copy()
                state[2 * i + field] = ranges[field][side]
                value = float(system.rhs(t, state)[2 * i + field])
                maximum_outward = max(maximum_outward, value if side else -value)
    report.metric("E3", "rectangle_facets_have_inward_rhs", maximum_outward, 1e-14)
    maximum = np.max(concentration)
    active = concentration == maximum
    report.metric("E3", "maximum_C_Dini_derivative_after_constant_tail",
                  max(0., float(np.max(f[active, 1]))), 1e-14)
    report.details["heat_capacity_boundary"] = (
        "Only instantaneous sum(v*s(C)*theta_dot) equals heat input. No assertion "
        "that d(sum(v*s(C)*theta))/dt equals it, or that a variable-capacity "
        "temperature quadratic energy is monotonically decreasing.")


def integrate_controlled(system, initial, times, source=None):
    """Authorized validation-only integration; no primary result is produced."""
    initial = np.asarray(initial, dtype=float)
    absolute = np.empty(initial.size)
    absolute[0::2], absolute[1::2] = 2e-11, 2e-13
    if source is None:
        fun = system.rhs
    else:
        def fun(t, y):
            return system.rhs(t, y) + source(t)
    solution = solve_ivp(fun, (float(times[0]), float(times[-1])), initial,
                         method="BDF", jac=system.jac, t_eval=times,
                         rtol=2e-11, atol=absolute, max_step=10.0)
    if not solution.success or solution.y.shape != (initial.size, len(times)):
        raise RuntimeError("Controlled validation solve failed: " + solution.message)
    return solution.y.T.reshape(len(times), -1, 2)


def bessel_constant_step(times, radii, modes):
    s0, k0, *_ = material(T0, C0)
    bi, alpha = H * R / k0, k0 / s0
    zero0, zero1 = jn_zeros(0, modes), jn_zeros(1, modes)
    roots = np.empty(modes)
    for i in range(modes):
        left = 0.0 if i == 0 else zero1[i - 1]
        roots[i] = brentq(lambda mu: mu * j1(mu) - bi * j0(mu), left, zero0[i],
                          xtol=1e-14, rtol=1e-14)
    coefficient = 2 * j1(roots) / (roots * (j0(roots) ** 2 + j1(roots) ** 2))
    decay = np.exp(-np.outer(times, alpha * roots ** 2 / R ** 2))
    basis = j0(np.outer(roots, radii / R))
    return 50.0 + (T0 - 50.0) * ((decay * coefficient) @ basis)


def check_bessel(report, make_system):
    times = np.array([0.0, 30.0, 120.0, 600.0, 1800.0])
    lower = bessel_constant_step(times[1:], OUTPUT_RADII, 128)
    series_error, modes = np.inf, 256
    while modes <= 2048:
        reference = bessel_constant_step(times[1:], OUTPUT_RADII, modes)
        series_error = float(np.max(np.abs(reference - lower)))
        if series_error <= 1e-9:
            break
        lower, modes = reference, modes * 2
    report.metric("E2", "bessel_series_truncation_difference", series_error, 1e-9,
                  modes=modes, excludes_initial_robin_corner=True)
    errors = []
    for n in (160, 320, 640):
        boundary = ConstantBoundary(50.0, C0)
        initial = interleave(np.full(n + 1, T0), np.full(n + 1, C0))
        state = integrate_controlled(make_system(n, boundary), initial, times)
        errors.append(float(np.max(np.abs(state[1:, ::n // 20, 0] - reference))))
        report.metric("E2", f"bessel_C_remains_uniform_n{n}",
                      np.max(np.abs(state[:, :, 1] - C0)), 1e-11)
    report.metric("E2", "bessel_finest_temperature", errors[-1], 5e-5,
                  grids=[160, 320, 640], grid_errors=errors)
    report.add("E2", "bessel_spatial_error_decreases", errors[-1] < errors[0] or max(errors) <= 1e-8)


def manufactured_fields(t, radii):
    """Smooth regular radial polynomials; derivatives are continuous formulas."""
    x = radii / R
    values, temporal, radial, laplacian = [], [], [], []
    coefficients = (
        (40 + .7 * np.sin(t / 200), 3 * np.exp(-t / 400), .4 * np.cos(t / 170),
         .7 / 200 * np.cos(t / 200), -3 / 400 * np.exp(-t / 400), -.4 / 170 * np.sin(t / 170)),
        (1.15 + .08 * np.sin(t / 260), .25 * np.exp(-t / 500), .035 * np.cos(t / 180),
         .08 / 260 * np.cos(t / 260), -.25 / 500 * np.exp(-t / 500), -.035 / 180 * np.sin(t / 180)))
    for base, a, b, base_t, a_t, b_t in coefficients:
        values.append(base + a * x ** 2 + b * x ** 4)
        temporal.append(base_t + a_t * x ** 2 + b_t * x ** 4)
        radial.append((2 * a * x + 4 * b * x ** 3) / R)
        laplacian.append((4 * a + 16 * b * x ** 2) / R ** 2)
    return tuple(np.asarray(item) for item in (values, temporal, radial, laplacian))


class ManufacturedBoundary:
    def evaluate(self, t):
        value, _, gradient, _ = manufactured_fields(t, np.array([R]))
        theta, concentration = value[:, 0]
        _, k, diffusion, *_ = material(theta, concentration)
        return (float(theta + k * gradient[0, 0] / H),
                float(concentration + diffusion * gradient[1, 0] / HM))


def manufactured_source(t, radii):
    value, temporal, gradient, laplacian = manufactured_fields(t, radii)
    theta, concentration = value
    s, k, diffusion, _, kp, dc, dt = material(theta, concentration)
    heat_operator = (k * laplacian[0] + kp * gradient[1] * gradient[0]) / s
    water_operator = (diffusion * laplacian[1] + dc * gradient[1] ** 2 +
                      dt * gradient[0] * gradient[1])
    # This is a prescribed additive rate source after dividing the temperature
    # PDE by s(C), so it is state-independent and contributes zero to the Jacobian.
    return interleave(temporal[0] - heat_operator, temporal[1] - water_operator)


def check_manufactured(report, make_system):
    times = np.array([0., 30., 60., 120., 240.])
    errors = []
    for n in (20, 40, 80):
        radii, _ = geometry(n)
        initial = manufactured_fields(0., radii)[0].T.reshape(-1)
        system = make_system(n, ManufacturedBoundary())
        state = integrate_controlled(system, initial, times,
                                     source=lambda t, r=radii: manufactured_source(t, r))
        exact = np.array([manufactured_fields(t, radii)[0].T for t in times])
        errors.append(np.max(np.abs(state[1:] - exact[1:]), axis=(0, 1)))
    errors = np.array(errors)
    # Modest budgets are chosen before any run; the expected smooth-solution
    # order is two, with >=1.5 checked away from the integration error floor.
    for field, col, absolute_limit, floor in (("temperature", 0, 5e-4, 1e-8),
                                              ("moisture", 1, 5e-6, 1e-10)):
        column = errors[:, col]
        order = float(np.log2(column[-2] / max(column[-1], 1e-300)))
        report.metric("E2", "manufactured_" + field + "_fine_error", column[-1], absolute_limit,
                      grids=[20, 40, 80], errors=column.tolist(), observed_order=order,
                      source="independent continuous radial divergence, not DUT residual")
        report.add("E2", "manufactured_" + field + "_convergence",
                   (column[-1] < column[-2] < column[-3] and order >= 1.5)
                   or max(column[-2], column[-1]) <= floor)


class RunView:
    """Read-only adapter for the formal Q2 NPZ and metadata contract."""
    def __init__(self, descriptor, base):
        self.descriptor, self.base = descriptor, Path(base)
        self.path = self.base / descriptor["npz"]
        self.meta_path = self.base / descriptor["metadata"]
        self.meta = json.loads(self.meta_path.read_text(encoding="utf-8"))
        self.n = int(self.meta["grid_n"])
        self.scenario = self.meta["scenario"]
        self.crossing = float(self.meta["t_cross_s"])
        self.stop = int(self.meta["n_end"])
        self.archive = np.load(self.path, allow_pickle=False)
        self.snapshot_cache = None

    def array(self, key):
        aliases = {"node_max_moisture": "max_moisture_kg_kg",
                   "node_argmax_radius_m": "max_moisture_radius_m"}
        if key == "snapshot_state":
            if self.snapshot_cache is None:
                self.snapshot_cache = np.stack((self.archive["temperature_snapshots"],
                                                self.archive["moisture_snapshots"]), axis=-1)
            return self.snapshot_cache
        return self.archive[aliases.get(key, key)]

    def state(self, t):
        times = self.array("snapshot_time_s")
        # Exact binary64 times survive JSON/NPY round trips. Do not conflate an
        # event 5e-8 s from an integer with that integer's separate snapshot.
        location = np.flatnonzero(times == float(t))
        if location.size == 0:
            raise ValueError(f"Run N={self.n} is missing a full-grid snapshot at {t}")
        values = self.array("snapshot_state")
        selected = np.asarray(values[int(location[0])])
        if any(not np.array_equal(selected, values[int(i)]) for i in location[1:]):
            raise ValueError("Coincident snapshot roles disagree")
        return selected

    def info(self, key):
        aliases = {"n": "grid_n", "crossing_time_s": "t_cross_s", "stop_integer_s": "n_end",
                   "max_step_before_4h": "max_step_observed_s", "max_step_after_4h": "max_step_tail_s"}
        return self.meta[aliases.get(key, key)]

    def close(self):
        if self.archive is not None:
            self.archive.close()


def compare_fields(first, second, end):
    outcome = {}
    for key in ("temperature_C", "moisture_kg_kg"):
        a, b = first.array(key), second.array(key)
        winner = dict(max_abs=-1.0)
        for start in range(0, end + 1, 4096):
            stop = min(end + 1, start + 4096)
            if not np.isfinite(a[start:stop]).all() or not np.isfinite(b[start:stop]).all():
                raise ValueError("A field comparison contains nonfinite values")
            difference = np.abs(a[start:stop] - b[start:stop])
            row, col = np.unravel_index(int(np.argmax(difference)), difference.shape)
            if difference[row, col] > winner["max_abs"]:
                winner = dict(max_abs=float(difference[row, col]), time_s=int(start + row),
                              radius_m=float(OUTPUT_RADII[col]),
                              first=float(a[start + row, col]), second=float(b[start + row, col]))
        outcome[key] = winner
        del a, b
    return outcome


def check_run(report, run, raw, common_end):
    label = f"{run.scenario}_n{run.n}_{run.descriptor.get('role', 'nominal')}"
    times, radii = run.array("time_s"), run.array("radius_m")
    report.add("E1", label + "_mesh_has_exact_output_nodes", run.n > 0 and run.n % 20 == 0)
    report.add("E1", label + "_complete_second_axis",
               times.ndim == 1 and len(times) > common_end and
               np.array_equal(times, np.arange(len(times), dtype=float)),
               required_common_end_s=int(common_end), actual_end_s=int(times[-1]))
    report.metric("E1", label + "_output_radii", np.max(np.abs(radii - OUTPUT_RADII)), 1e-15)
    for field, limits, suffix, allowance in (("temperature", raw.theta_bounds, "C", 5e-7),
                                            ("moisture", raw.moisture_bounds, "kg_kg", 5e-9)):
        low = float(run.meta["observed_extrema"][f"{field}_min_{suffix}"])
        high = float(run.meta["observed_extrema"][f"{field}_max_{suffix}"])
        report.add("E2", label + "_accepted_and_sampled_all_node_" + field + "_range",
                   np.isfinite([low, high]).all() and limits[0] - allowance <= low <= high <= limits[1] + allowance,
                   observed_range=[low, high], theoretical_range=list(limits),
                   numerical_allowance=allowance, scope=run.meta["extrema_scope"])
    maxima = run.array("node_max_moisture")
    positions = run.array("node_argmax_radius_m")
    node_indices = run.array("max_moisture_node")
    report.add("E1", label + "_all_node_maximum_axis", maxima.shape == times.shape and positions.shape == times.shape)
    report.add("E2", label + "_all_node_maximum_finite_and_bounded",
               np.isfinite(maxima).all() and maxima.min() >= raw.moisture_bounds[0] - 5e-9 and
               maxima.max() <= raw.moisture_bounds[1] + 5e-9)
    report.add("E1", label + "_argmax_is_a_mesh_location",
               np.isfinite(positions).all() and np.min(positions) >= 0 and np.max(positions) <= R and
               np.max(np.abs(positions / (R / run.n) - np.round(positions / (R / run.n)))) <= 1e-8)
    report.add("E1", label + "_argmax_node_indices",
               node_indices.shape == times.shape and node_indices.dtype.kind in "iu" and
               np.min(node_indices) >= 0 and np.max(node_indices) <= run.n and
               np.max(np.abs(positions - node_indices * R / run.n)) <= 1e-15)
    report.metric("E1", label + "_full_mesh_radius",
                  np.max(np.abs(run.array("mesh_radius_m") - geometry(run.n)[0])), 1e-15)
    for key, index, initial, limits, allowance in (
        ("temperature_C", 0, T0, raw.theta_bounds, 5e-7),
        ("moisture_kg_kg", 1, C0, raw.moisture_bounds, 5e-9)):
        values = run.array(key)
        report.add("E1", label + "_" + key + "_shape_dtype",
                   values.shape == (len(times), 21) and values.dtype == np.dtype("float64"))
        low, high, finite, max_sample_gap = np.inf, -np.inf, True, 0.0
        for start in range(0, len(times), 4096):
            block = values[start:start + 4096]
            finite = finite and bool(np.isfinite(block).all())
            low, high = min(low, float(np.min(block))), max(high, float(np.max(block)))
            if index == 1:
                max_sample_gap = max(max_sample_gap,
                                     float(np.max(np.max(block, axis=1) - maxima[start:start + len(block)])))
        report.add("E2", label + "_" + key + "_finite_and_invariant_range",
                   finite and low >= limits[0] - allowance and high <= limits[1] + allowance,
                   observed_range=[low, high], theoretical_range=list(limits), numerical_allowance=allowance)
        report.metric("E1", label + "_" + key + "_initial", np.max(np.abs(values[0] - initial)), 1e-13)
        if index == 1:
            report.metric("E1", label + "_max_over_all_nodes_dominates_21_samples", max_sample_gap, 1e-13)
        del values

    snapshot_times = run.array("snapshot_time_s")
    snapshots = run.array("snapshot_state")
    report.add("E1", label + "_full_snapshot_shape", snapshots.shape == (len(snapshot_times), run.n + 1, 2))
    roles = run.meta["snapshot_roles"]
    required_roles = {"initial": 0., "observed_end": 14400., "before_n_end": float(run.stop - 1),
                      "n_end": float(run.stop), "threshold_crossing": run.crossing,
                      "common_end": float(run.meta["common_end_s"])}
    required_roles.update({f"summary_{int(t)}s": t for t in SUMMARY_TIMES})
    report.add("E1", label + "_snapshot_role_time_mapping",
               all(role in roles and isinstance(roles[role], int) and
                   0 <= roles[role] < len(snapshot_times) and snapshot_times[roles[role]] == expected
                   for role, expected in required_roles.items()))
    required_times = np.unique(np.r_[0., 14400., SUMMARY_TIMES, run.stop - 1, run.stop, common_end])
    selected_snapshots = []
    for t in required_times:
        state = run.state(float(t))
        selected_snapshots.append(state)
        report.add("E2", label + f"_snapshot_{t:g}_physical", np.isfinite(state).all() and
                   state[:, 1].min() >= raw.moisture_bounds[0] - 5e-9 and
                   state[:, 1].max() <= raw.moisture_bounds[1] + 5e-9 and
                   state[:, 0].min() >= raw.theta_bounds[0] - 5e-7 and
                   state[:, 0].max() <= raw.theta_bounds[1] + 5e-7)
        report.metric("E1", label + f"_snapshot_{t:g}_all_node_max",
                      abs(np.max(state[:, 1]) - maxima[int(t)]), 1e-13)
        report.metric("E1", label + f"_snapshot_{t:g}_argmax_value",
                      abs(state[int(round(positions[int(t)] / (R / run.n))), 1] - maxima[int(t)]), 1e-13)
    selected_snapshots = np.asarray(selected_snapshots)
    report.metric("E1", label + "_full_initial_temperature",
                  np.max(np.abs(run.state(0.)[:, 0] - T0)), 1e-13)
    report.metric("E1", label + "_full_initial_moisture",
                  np.max(np.abs(run.state(0.)[:, 1] - C0)), 1e-13)
    _, independent_volumes = geometry(run.n)
    for key, index in (("temperature_C", 0), ("moisture_kg_kg", 1)):
        # Load a compressed field once, rather than once per snapshot.
        samples = run.array(key)
        expected = selected_snapshots[:, ::run.n // 20, index]
        report.metric("E1", label + f"_all_required_snapshot_projections_{index}",
                      np.max(np.abs(expected - samples[required_times.astype(int)])), 1e-12)
        del samples
        mean_key = "temperature_mean" if index == 0 else "moisture_mean"
        means = run.array(mean_key)
        expected_means = np.sum(selected_snapshots[:, :, index] * independent_volumes, axis=1) * 2 / R ** 2
        report.add("E1", label + "_" + mean_key + "_axis_finite",
                   means.shape == times.shape and np.isfinite(means).all())
        report.metric("E1", label + "_" + mean_key + "_independent_snapshot_weights",
                      np.max(np.abs(means[required_times.astype(int)] - expected_means)), 2e-12)

    first_below = np.flatnonzero(maxima < THRESHOLD)
    report.add("E1", label + "_first_strict_integer_endpoint",
               first_below.size > 0 and int(first_below[0]) == run.stop and
               maxima[run.stop - 1] >= THRESHOLD and maxima[run.stop] < THRESHOLD,
               stop_integer_s=run.stop, previous_max=float(maxima[run.stop - 1]),
               current_max=float(maxima[run.stop]), threshold=THRESHOLD)
    event_state = run.state(run.crossing)
    report.metric("E2", label + "_event_uses_all_nodes",
                  abs(np.max(event_state[:, 1]) - THRESHOLD), 5e-10,
                  crossing_time_s=run.crossing, argmax_radius_m=float(np.argmax(event_state[:, 1]) * R / run.n))
    report.add("E1", label + "_event_precedes_strict_integer", run.stop - 1 - 1e-7 <= run.crossing <= run.stop + 1e-7)
    report.add("E1", label + "_metadata_common_end",
               run.meta["common_end_s"] == int(times[-1]))


def check_tail_structure(report, run):
    """E3 calculation, called only by the model-axis entry point."""
    label = f"{run.scenario}_n{run.n}_{run.descriptor.get('role', 'nominal')}"
    maxima = run.array("node_max_moisture")
    if not np.isfinite(maxima).all():
        raise ValueError("Nonfinite all-node maximum in structural comparison")
    # On the fixed-tail interval the mathematical maximum is nonincreasing.
    tail = maxima[14400:]
    active = tail[:-1] >= TAIL_C if run.scenario == "mean_tail" else tail[:-1] >= LAST_C
    upward = np.diff(tail)[active]
    report.metric("E3", label + "_constant_tail_maximum_nonincrease",
                  max(0., float(upward.max())) if len(upward) else 0., 5e-9)


def check_full_implementation(report, runs, tight, raw):
    if [run.n for run in runs] != [160 * 2 ** i for i in range(len(runs))] or len(runs) < 3:
        raise ValueError("Spatial sequence must start at160, double, and include at least three levels")
    if tight.n != runs[-1].n:
        raise ValueError("Tight comparison must use the accepted finest mesh")
    report.add("E1", "primary_and_tight_scenario_contract",
               all(run.scenario == "mean_tail" and run.meta["tight"] is False for run in runs) and
               tight.scenario == "mean_tail" and tight.meta["tight"] is True)
    for key, upper in (("rtol", 2e-9), ("atol_temperature", 2e-9), ("atol_moisture", 2e-11),
                       ("max_step_before_4h", 5.), ("max_step_after_4h", 300.)):
        report.add("E1", "nominal_settings_within_specification_" + key,
                   0 < runs[-1].info(key) <= upper)
    for key in ("rtol", "atol_temperature", "atol_moisture"):
        ratio = float(tight.info(key)) / float(runs[-1].info(key))
        report.metric("E2", "tight_metadata_" + key, abs(ratio - .1), 1e-14)
    for key in ("max_step_before_4h", "max_step_after_4h"):
        ratio = float(tight.info(key)) / float(runs[-1].info(key))
        report.metric("E2", "tight_metadata_" + key, abs(ratio - .5), 1e-14)
    for key in ("rtol", "atol_temperature", "atol_moisture", "max_step_before_4h", "max_step_after_4h"):
        report.add("E1", "spatial_sequence_same_time_settings_" + key,
                   all(run.info(key) == runs[-1].info(key) for run in runs))
    ends = {int(run.meta["common_end_s"]) for run in [*runs, tight]}
    minimum_end = max([r.stop for r in runs] + [tight.stop])
    if len(ends) != 1 or min(ends) < minimum_end:
        raise ValueError("Grid/time runs need the same actual common_end_s covering every strict endpoint")
    common_end = ends.pop()
    for run in [*runs, tight]:
        check_run(report, run, raw, common_end)
    differences = [compare_fields(a, b, common_end) for a, b in zip(runs[:-1], runs[1:])]
    report.details["spatial_comparisons"] = [dict(coarse=a.n, fine=b.n, **d)
                                             for a, b, d in zip(runs[:-1], runs[1:], differences)]
    for key in ("temperature_C", "moisture_kg_kg"):
        report.metric("E2", "full_process_spatial_" + key, differences[-1][key]["max_abs"], 2.5e-5,
                      comparison=differences[-1][key])
        previous, last = differences[-2][key]["max_abs"], differences[-1][key]["max_abs"]
        report.add("E2", "spatial_trend_" + key, last < previous or max(previous, last) <= 5e-7,
                   previous_difference=previous, final_difference=last,
                   observed_ratio=previous / max(last, 1e-300))
    report.metric("E2", "spatial_crossing_time_difference", abs(runs[-1].crossing - runs[-2].crossing), 1.0)
    temporal = compare_fields(runs[-1], tight, common_end)
    for key in ("temperature_C", "moisture_kg_kg"):
        report.metric("E2", "full_process_temporal_" + key, temporal[key]["max_abs"], 5e-6,
                      comparison=temporal[key])
    report.metric("E2", "temporal_crossing_time_difference", abs(runs[-1].crossing - tight.crossing), .1)
    endpoints = [runs[-2].stop, runs[-1].stop, tight.stop]
    report.details["integer_endpoint_stability"] = dict(
        stable=len(set(endpoints)) == 1, coarse_fine_tight_endpoints=endpoints,
        status="stable_observed" if len(set(endpoints)) == 1 else "integer_boundary_unresolved",
        claim="The nominal path has a strict all-node integer endpoint; cross-grid integer stability is separate.")



def check_full_structure(report, nominal, sensitivity):
    """Only E3 calculations: tail structure and the two accepted future inputs."""
    if sensitivity.n != nominal.n:
        raise ValueError("Scenario comparison must use the accepted finest mesh")
    for key in ("rtol", "atol_temperature", "atol_moisture", "max_step_before_4h", "max_step_after_4h"):
        report.add("E3", "sensitivity_same_numerical_settings_" + key,
                   sensitivity.info(key) == nominal.info(key))
    check_tail_structure(report, nominal)
    check_tail_structure(report, sensitivity)
    # Sensitivity uses a genuinely different future boundary, not an arbitrary
    # parameter percentage. No smallness threshold is invented for its outcome.
    comparison_end = min(len(nominal.array("time_s")), len(sensitivity.array("time_s"))) - 1
    if comparison_end < 14400:
        raise ValueError("Sensitivity comparison does not cover the observed boundary")
    if sensitivity.scenario != "last_value" or nominal.scenario != "mean_tail":
        raise ValueError("Incorrect sensitivity scenarios")
    causal = compare_fields(nominal, sensitivity, 10800)
    for key, limit in (("temperature_C", 5e-9), ("moisture_kg_kg", 5e-11)):
        report.metric("E3", "unchanged_first_three_hours_" + key, causal[key]["max_abs"], limit)
    report.details["long_time_boundary_sensitivity"] = dict(
        common_comparison_end_s=comparison_end,
        field_differences=compare_fields(nominal, sensitivity, comparison_end),
        crossing_time_difference_s=sensitivity.crossing - nominal.crossing,
        integer_endpoint_difference_s=sensitivity.stop - nominal.stop,
        scope="Only the two user-selected future environments; not statistical uncertainty or E4 validation.")


def load_dut(path):
    spec = importlib.util.spec_from_file_location("q2_primary_under_test", path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def bind_common(report):
    artifacts = ((ROOT / "planning/Q2/model_spec.md", SPEC_SHA),
                 (ROOT / "output/ENV/q2_boundary.json", BOUNDARY_SHA),
                 (ROOT / "附件/附件1.xlsx", RAW_SHA),
                 (ROOT / "src/Q2/solve.py", None), (Path(__file__).resolve(), None))
    for path, expected in artifacts:
        actual = digest(path)
        report.identities[str(path.resolve())] = actual
        if expected is not None and expected != actual:
            raise ValueError("Frozen artifact identity changed: " + str(path))


def bind_run(report, run):
    expected = {"spec_sha256": SPEC_SHA, "boundary_sha256": BOUNDARY_SHA,
                "raw_environment_sha256": RAW_SHA,
                "solver_sha256": report.identities[str((ROOT / "src/Q2/solve.py").resolve())]}
    for key, value in expected.items():
        if run.meta.get(key) != value:
            raise ValueError(f"Run {run.path.name} has a different {key}")
    result_hash = digest(run.path)
    if run.meta["result_sha256"] != result_hash or run.meta["success"] is not True:
        raise ValueError("Run result identity or success state is invalid")
    report.details.setdefault("run_identities", []).append(dict(
        result=str(run.path.resolve()), result_sha256=result_hash,
        metadata=str(run.meta_path.resolve()), metadata_sha256=digest(run.meta_path)))


def reuse_checks_report(report, paths, base):
    if isinstance(paths, str):
        paths = [paths]
    for path in paths:
        path = Path(base) / path
        previous = json.loads(path.read_text(encoding="utf-8"))
        if previous.get("mode") != "checks_only" or previous.get("axes") != report.axes:
            raise ValueError("A reused quick report must have the same actual axes")
        if previous.get("passed") is not True:
            raise ValueError("Cannot reuse a failed quick report")
        if any(previous.get("identities", {}).get(key) != value for key, value in report.identities.items()):
            raise ValueError("A reused quick report is stale for the current artifacts")
        if any(check.get("axis") not in report.axes for check in previous["checks"]):
            raise ValueError("A quick report contains evidence from another runtime axis")
        report.checks.extend(previous["checks"])
        report.details.setdefault("reused_checks_reports", []).append(dict(
            path=str(path.resolve()), sha256=digest(path), details=previous.get("details", {})))


def json_ready(value):
    """Keep failed numerical diagnostics serializable without disguising them."""
    if isinstance(value, dict):
        return {str(key): json_ready(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [json_ready(item) for item in value]
    if isinstance(value, np.ndarray):
        return json_ready(value.tolist())
    if isinstance(value, np.generic):
        return json_ready(value.item())
    if isinstance(value, float) and not np.isfinite(value):
        return str(value)
    return value


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument("--checks-only", action="store_true")
    mode.add_argument("--runs-config", type=Path)
    parser.add_argument("--axes", required=True, choices=("E1,E2", "E3"))
    parser.add_argument("--report", required=True, type=Path)
    args = parser.parse_args()
    if args.report.exists():
        raise FileExistsError("A verification report must not overwrite an earlier identity")
    axes = set(args.axes.split(","))
    report = Report(axes, "checks_only" if args.checks_only else "full_process")
    views = []
    try:
        bind_common(report)
        if args.checks_only:
            dut = load_dut(ROOT / "src/Q2/solve.py")
            make_system = lambda n, boundary: dut.System(n, boundary)
            raw = RawBoundary(ROOT / "附件/附件1.xlsx")
            if axes == {"E1", "E2"}:
                primary = dut.Boundary(ROOT / "output/ENV/q2_boundary.json")
                check_boundary(report, raw, primary, ROOT / "output/ENV/q2_boundary.json")
                # Independently check both publicly exposed long-time choices.
                last_raw = RawBoundary(ROOT / "附件/附件1.xlsx", "last_value")
                last_primary = dut.Boundary(ROOT / "output/ENV/q2_boundary.json", "last_value")
                last_error = max(np.max(np.abs(np.asarray(last_primary.evaluate(t)) - last_raw.evaluate(t)))
                                 for t in (0., 14400., np.nextafter(14400., np.inf), 14460., 1e6))
                report.metric("E1", "last_value_boundary_extension", last_error, 2e-12)
                check_implementation(report, make_system, raw)
                check_bessel(report, make_system)
                check_manufactured(report, make_system)
                report.details["coverage"] = "E1 implementation and E2 independent small examples; full-process refinement remains separate."
            else:
                check_structure(report, make_system, raw)
                report.details["coverage"] = "E3 instantaneous structure; full-process tail and scenario results remain separate."
        else:
            config = json.loads(args.runs_config.read_text(encoding="utf-8"))
            base = args.runs_config.parent
            report.details["runs_config_identity"] = dict(path=str(args.runs_config.resolve()),
                                                          sha256=digest(args.runs_config))
            reuse_checks_report(report, config.get("checks_report", []), base)
            if axes == {"E1", "E2"}:
                runs = []
                for item in config["spatial_runs"]:
                    view = RunView(item, base)
                    views.append(view)
                    bind_run(report, view)
                    runs.append(view)
                tight = RunView(config["tight_run"], base)
                views.append(tight)
                bind_run(report, tight)
                raw = RawBoundary(ROOT / "附件/附件1.xlsx")
                check_full_implementation(report, runs, tight, raw)
                if "sensitivity_run" in config:
                    sensitivity = RunView(config["sensitivity_run"], base)
                    views.append(sensitivity)
                    bind_run(report, sensitivity)
                    check_run(report, sensitivity, RawBoundary(raw.path, "last_value"),
                              len(sensitivity.array("time_s")) - 1)
                report.details["coverage"] = "E1 result contract and E2 full-process refinement; earlier quick evidence is included only if explicitly bound."
            else:
                nominal = RunView(config["spatial_runs"][-1], base)
                views.append(nominal)
                bind_run(report, nominal)
                sensitivity = RunView(config["sensitivity_run"], base)
                views.append(sensitivity)
                bind_run(report, sensitivity)
                check_full_structure(report, nominal, sensitivity)
                report.details["coverage"] = "E3 tail structure and two-scenario sensitivity; no E1/E2 auxiliary solver was executed by this entry point."
    except Exception as exc:
        report.add(report.axes[0], "verification_execution_failed", False,
                   exception=type(exc).__name__, message=str(exc))
    finally:
        for view in views:
            view.close()
    args.report.parent.mkdir(parents=True, exist_ok=True)
    with args.report.open("x", encoding="utf-8") as handle:
        handle.write(json.dumps(json_ready(report.result()), ensure_ascii=False,
                                indent=2, allow_nan=False) + "\n")
    if not report.result()["passed"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
