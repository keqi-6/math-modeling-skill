"""Preserve approved Q1 environment nodes and their explicit interpolation contract."""
from pathlib import Path
import hashlib
import json
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]


def main():
    source = ROOT / '附件/附件1.xlsx'
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    book = load_workbook(source, read_only=True, data_only=False)
    try:
        sheet = book['Sheet1']
        values = list(sheet.iter_rows(min_row=2, max_row=32, max_col=3, values_only=True))
    finally:
        book.close()
    nodes = []
    for index, (t, temperature, moisture) in enumerate(values):
        if t != index * 60 or not all(isinstance(x, (int, float)) for x in (t, temperature, moisture)):
            raise ValueError('Official node schema or time changed')
        if not (0 < temperature < 100 and 0 < moisture < 1):
            raise ValueError('Official values outside the audited domain')
        nodes.append({'time_s': t, 'temperature_C': temperature,
                      'air_moisture_kg_kg': moisture, 'source_row': index + 2})
    if len(nodes) != 31 or hashlib.sha256(source.read_bytes()).hexdigest() != before:
        raise ValueError('Incomplete or changed input')
    record = {'schema_version': '1.0', 'scope': 'Q1_environment_0_to_1800s',
              'source': {'path': '附件/附件1.xlsx', 'sha256': before,
                         'sheet': 'Sheet1', 'range': 'A2:C32'},
              'units': {'time_s': 's', 'temperature_C': 'degC',
                        'air_moisture_kg_kg': 'kg/kg (air; per official statement)'},
              'interpolation': 'piecewise_linear', 'valid_time_s': [0, 1800],
              'extrapolation': 'error', 'rounding': 'none', 'nodes': nodes,
              'approval': '用户于2026-09-10确认上一轮成组呈现的原始31节点分段线性方案。',
              'generator_sha256': hashlib.sha256(Path(__file__).read_bytes()).hexdigest()}
    target = ROOT / 'output/ENV/q1_boundary.json'
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(record, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(f'Preserved {len(nodes)} nodes: {target}')


if __name__ == '__main__':
    main()
