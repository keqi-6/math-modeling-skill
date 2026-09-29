"""Independent Q1-v1 verification; only the report JSON is written.

The production RHS/Jacobian are called as objects under test. Expected geometry,
balances, raw-input interpolation and Bessel references are implemented here
without importing the production coefficients, geometry or parsing helpers.
The constant-D check deliberately executes one controlled production solve.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import platform
import traceback
from datetime import datetime, timezone
from functools import lru_cache
from pathlib import Path

import numpy as np
import scipy
from openpyxl import load_workbook
from scipy.optimize import brentq
from scipy.special import j0, j1, jn_zeros

import solve as dut

ROOT = Path(__file__).resolve().parents[2]
R = 0.02
CAPACITY = 820.0 * 2600.0
K, H, HM = 0.36, 25.0, 8e-7
D_STAR = 7e-9 * np.exp(-0.89 / 2.55)
SPEC_HASH = "c9aa1d7771b729e4659a6ae7d92bf5523b3fefca4ea28c95630ed62dd350285a"
BOUNDARY_HASH = "4c80135753189b839bcbe3ccc667a65118364a07471135695180fe7e8c6f923d"
RAW_HASH = "7ef32870abeef420b89560b2530ff60dfe4255917805151d89988d0311af9dd7"
TIMES = np.arange(1801)
RADII = np.arange(21) * 0.001
SNAP_TIMES = np.array([0, 1, 60, 1800])


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


class Report:
    def __init__(self):
        self.checks = []
        self.details = {}

    def add(self, name, passed, **details):
        self.checks.append(dict(name=name, passed=bool(passed), **details))

    def metric(self, name, value, limit, **details):
        self.add(name, np.isfinite(value) and value <= limit,
                 value=float(value), limit=float(limit), **details)

    def reject_call(self, name, callback):
        try:
            callback()
        except ValueError as exc:
            self.add(name, True, exception=type(exc).__name__, message=str(exc))
        except Exception as exc:
            self.add(name, False, exception=type(exc).__name__, message=str(exc))
        else:
            self.add(name, False, message="Invalid input was accepted")

    def result(self):
        return dict(
            schema_version="1.0", question="Q1", role="auxiliary_validator",
            created_at=datetime.now(timezone.utc).isoformat(),
            passed=bool(self.checks) and all(c["passed"] for c in self.checks),
            checks=self.checks, details=self.details,
            claim_boundary=(
                "Implementation, numerical consistency and selected structural checks "
                "for Q1-v1; no experimental validation or per-cell correct-rounding claim. "
                "Accumulated flux from the same integrator is an internal balance check."
            ),
        )


class RawBoundary:
    """Independent direct worksheet reader and scalar linear interpolator."""
    def __init__(self, path):
        self.path = Path(path)
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = list(workbook["Sheet1"].iter_rows(
                min_row=2, max_row=32, min_col=1, max_col=3, values_only=True))
        finally:
            workbook.close()
        self.data = np.asarray(rows, dtype=float)
        if self.data.shape != (31, 3) or not np.isfinite(self.data).all():
            raise ValueError("Raw boundary requires 31 finite three-column records")
        if not np.array_equal(self.data[:, 0], np.arange(0, 1801, 60)):
            raise ValueError("Raw boundary time nodes differ from the specification")

    def evaluate(self, t):
        if not np.isfinite(t) or not 0 <= t <= 1800:
            raise ValueError("Raw reference query is outside its time domain")
        index = min(int(t // 60), 29)
        weight = (float(t) - self.data[index, 0]) / 60.0
        out = (1.0 - weight) * self.data[index, 1:] + weight * self.data[index + 1, 1:]
        return float(out[0]), float(out[1])


class ConstantBoundary:
    def __init__(self, temperature=35.0, moisture=1.3):
        self.value = (temperature, moisture)

    def evaluate(self, t):
        return self.value


def geometry(n):
    dr = R / n
    radii = np.arange(n + 1) * dr
    volume = radii * dr
    volume[0] = dr * dr / 8.0
    volume[-1] = R * dr / 2.0 - dr * dr / 8.0
    return radii, volume


def field_parameters(field):
    return (CAPACITY, H, 28.0) if field == "temperature" else (1.0, HM, 2.55)


def independent_rhs(n, field, raw_boundary, t, y, transfer=None, constant_d=False):
    """Explicit per-face assembly, with independently derived nodal volumes."""
    _, volume = geometry(n)
    capacity, default_transfer, _ = field_parameters(field)
    transfer = default_transfer if transfer is None else transfer
    u = np.asarray(y[:-1])
    if field == "temperature":
        diffusion = np.full(n + 1, K)
        ambient = raw_boundary.evaluate(t)[0]
    else:
        diffusion = (np.full(n + 1, D_STAR) if constant_d
                     else 7e-9 * np.exp(-0.89 / u))
        ambient = raw_boundary.evaluate(t)[1]
    out = np.zeros(n + 2)
    for i in range(n):
        incoming = (i + 0.5) * (diffusion[i] + diffusion[i + 1]) / 2 * (u[i + 1] - u[i])
        out[i] += incoming / (capacity * volume[i])
        out[i + 1] -= incoming / (capacity * volume[i + 1])
    boundary_flux = R * transfer * (ambient - u[-1])
    out[-2] += boundary_flux / (capacity * volume[-1])
    out[-1] = boundary_flux
    return out


def compare_grid_values(first, second):
    differences = np.abs(np.asarray(first)[1:] - np.asarray(second)[1:])
    row, col = np.unravel_index(int(np.argmax(differences)), differences.shape)
    return dict(max_abs=float(differences[row, col]), time_s=int(row + 1),
                radius_m=float(RADII[col]), first=float(first[row + 1, col]),
                second=float(second[row + 1, col]),
                differing_four_decimal_cells=int(np.count_nonzero(
                    np.round(first[1:], 4) != np.round(second[1:], 4))))


def check_inputs(report, raw, boundary_path):
    report.add("frozen_spec_identity", digest(ROOT / "planning/Q1/model_spec.md") == SPEC_HASH)
    report.add("frozen_boundary_identity", digest(boundary_path) == BOUNDARY_HASH)
    report.add("raw_workbook_identity", digest(raw.path) == RAW_HASH)
    frozen = json.loads(Path(boundary_path).read_text(encoding="utf-8"))
    frozen_nodes = np.array([[node[k] for k in
                             ("time_s", "temperature_C", "air_moisture_kg_kg")]
                            for node in frozen["nodes"]], dtype=float)
    report.add("frozen_nodes_equal_raw_worksheet",
               np.array_equal(frozen_nodes, raw.data),
               maximum_difference=float(np.max(np.abs(frozen_nodes - raw.data))))
    primary = dut.Boundary(boundary_path)
    queries = np.unique(np.r_[TIMES, np.arange(0, 1800, 60) + 0.25,
                              np.arange(0, 1800, 60) + 59.75])
    errors = np.array([np.abs(np.asarray(primary.evaluate(float(t))) -
                             raw.evaluate(float(t))) for t in queries])
    report.metric("boundary_temperature_interpolation", np.max(errors[:, 0]), 1e-12)
    report.metric("boundary_moisture_interpolation", np.max(errors[:, 1]), 1e-14)
    for label, value in (("negative", -1.0), ("above_end", 1800.01),
                         ("nan", float("nan")), ("infinity", float("inf"))):
        report.reject_call("boundary_rejects_" + label, lambda v=value: primary.evaluate(v))
    report.reject_call("boundary_rejects_wrong_identity",
                       lambda: dut.Boundary(ROOT / "planning/Q1/model_spec.md"))
    for value in (0, 10, 21, 20.0):
        report.reject_call("grid_rejects_" + str(value),
                           lambda v=value: dut.FieldSystem(v, "temperature", primary))
    report.reject_call("field_rejects_unknown_name",
                       lambda: dut.FieldSystem(20, "unknown", primary))
    return primary


def check_implementation(report, raw):
    n, t = 40, 137.25
    radii, volume = geometry(n)
    x = radii / R
    for field in ("temperature", "moisture"):
        system = dut.FieldSystem(n, field, raw)
        capacity, transfer, initial = field_parameters(field)
        u = (28.0 + 5.0 * x**2 + 0.3 * np.cos(np.pi*x)
             if field == "temperature" else 0.8 + 1.7*(1-x*x) + 0.07*np.cos(3*np.pi*x))
        y = np.r_[u, 0.0]
        f = system.rhs(t, y)
        expected = independent_rhs(n, field, raw, t, y)
        report.metric(field + "_geometry_volume", np.max(np.abs(system.v-volume)), 1e-18)
        report.metric(field + "_geometry_radius", np.max(np.abs(system.r-radii)), 1e-16)
        report.add(field + "_coefficients_contract",
                   system.capacity == capacity and system.transfer == transfer)
        report.metric(field + "_rhs_independent_formula",
                      np.max(np.abs(f-expected)) / max(np.max(np.abs(expected)), 1e-30), 1e-11)
        ambient = raw.evaluate(t)[0 if field == "temperature" else 1]
        flux = R * transfer * (ambient-u[-1])
        terms = capacity * volume * f[:-1]
        algebra_error = abs(np.sum(terms)-flux) / max(np.sum(np.abs(terms))+abs(flux), 1e-30)
        report.metric(field + "_algebraic_conservation", algebra_error, 1e-10)
        report.metric(field + "_integral_rhs", abs(f[-1]-flux) / max(abs(flux), 1e-30), 1e-12)

        directions = [np.cos(1.37*np.arange(n+2)) + 0.2*np.sin(0.71*np.arange(n+2)),
                      np.zeros(n+2)]
        directions[1][[0, 1, n, n+1]] = [1.0, -0.5, 0.8, 0.37]
        step = 1e-5 * max(1.0, float(np.max(np.abs(u))))
        for index, p in enumerate(directions):
            p = p / np.max(np.abs(p))
            jp = np.asarray(system.jac(t, y) @ p)
            fd = (system.rhs(t, y+step*p)-system.rhs(t, y-step*p)) / (2*step)
            error = np.max(np.abs(jp-fd)) / max(np.max(np.abs(jp)), np.max(np.abs(fd)), 1e-12)
            report.metric(field + "_jacobian_direction_" + str(index), error, 1e-5,
                          step=float(step), worst_component=int(np.argmax(np.abs(jp-fd))),
                          direction_norm="infinity norm 1")

        bad = y.copy()
        bad[3] = np.nan
        report.reject_call(field + "_rejects_nonfinite", lambda s=system, v=bad: s.rhs(t, v))
        if field == "moisture":
            for value in (0.0, -0.01):
                bad = y.copy()
                bad[3] = value
                report.reject_call("moisture_rejects_" + str(value),
                                   lambda s=system, v=bad: s.rhs(t, v))

        constant = ConstantBoundary()
        fixed = dut.FieldSystem(n, field, constant)
        g = constant.evaluate(0)[0 if field == "temperature" else 1]
        equilibrium = fixed.rhs(0, np.r_[np.full(n+1, g), 0.0])
        report.add(field + "_constant_equilibrium", np.array_equal(equilibrium, np.zeros(n+2)),
                   maximum_rhs=float(np.max(np.abs(equilibrium))))
        fixed_rhs = fixed.rhs(0, y)
        diffusion = np.full(n+1, K) if field == "temperature" else 7e-9*np.exp(-0.89/u)
        dissipation = -np.sum((np.arange(n)+0.5)*(diffusion[:-1]+diffusion[1:])/2 *
                              np.diff(u)**2) - R*transfer*(u[-1]-g)**2
        energy_terms = capacity*volume*(u-g)*fixed_rhs[:-1]
        residual = abs(np.sum(energy_terms)-dissipation) / max(
            np.sum(np.abs(energy_terms))+abs(dissipation), 1e-30)
        report.metric(field + "_constant_environment_dissipation_identity", residual, 1e-10)
        report.add(field + "_constant_environment_dissipation_sign",
                   np.sum(energy_terms) <= 0 and dissipation <= 0)

        insulated = dut.FieldSystem(n, field, constant)
        insulated.transfer = 0.0
        no_flux_rhs = insulated.rhs(0, y)
        terms = capacity*volume*no_flux_rhs[:-1]
        report.metric(field + "_zero_exchange_conservation",
                      abs(np.sum(terms)) / max(np.sum(np.abs(terms)), 1e-30), 1e-10)
        report.add(field + "_zero_exchange_accumulator", no_flux_rhs[-1] == 0.0)
        constant_no_flux = insulated.rhs(0, np.r_[np.full(n+1, initial), 0.0])
        report.add(field + "_zero_exchange_constant_rhs", np.all(constant_no_flux == 0))


def load_run(path, report, label):
    path = Path(path).resolve()
    meta_path = path.with_suffix(".json")
    meta = json.loads(meta_path.read_text(encoding="utf-8"))
    with np.load(path, allow_pickle=False) as archive:
        arrays = {key: np.array(archive[key]) for key in archive.files}
    n = len(arrays["mesh_radius_m"]) - 1
    if n < 20 or n % 20 or int(meta["n"]) != n:
        raise ValueError("Run mesh and metadata N disagree or violate output-node alignment")
    required_shapes = {"time_s": (1801,), "radius_m": (21,),
                       "mesh_radius_m": (n+1,), "snapshot_time_s": (4,)}
    for field, values_key in (("temperature", "temperature_C"), ("moisture", "moisture_kg_kg")):
        required_shapes.update({values_key: (1801, 21), field+"_mean": (1801,),
                                field+"_integral": (1801,), field+"_snapshots": (4, n+1),
                                field+"_integral_snapshots": (4,)})
    for key, shape in required_shapes.items():
        if key not in arrays or arrays[key].shape != shape or not np.isfinite(arrays[key]).all():
            raise ValueError("Missing, wrong-shaped or nonfinite run array: " + key)
    report.add(label + "_time_axis", np.array_equal(arrays["time_s"], TIMES))
    report.add(label + "_snapshot_times", np.array_equal(arrays["snapshot_time_s"], SNAP_TIMES))
    report.metric(label + "_output_radius", np.max(np.abs(arrays["radius_m"]-RADII)), 1e-16)
    report.metric(label + "_mesh_radius", np.max(np.abs(arrays["mesh_radius_m"]-geometry(n)[0])), 1e-16)
    identities = {"result_sha256": digest(path), "spec_sha256": SPEC_HASH,
                  "boundary_sha256": BOUNDARY_HASH, "solver_sha256": digest(ROOT/"src/Q1/solve.py")}
    for key, expected in identities.items():
        report.add(label + "_" + key, meta.get(key) == expected,
                   recorded=meta.get(key), expected=expected)
    report.add(label + "_method", meta.get("method") == "nodal_finite_volume_segmented_BDF")
    report.add(label + "_max_step", meta.get("max_step_s") == 5.0)
    return dict(path=str(path), sha256=identities["result_sha256"],
                metadata_path=str(meta_path), metadata_sha256=digest(meta_path),
                n=n, meta=meta, arrays=arrays, label=label)


def check_run_fields(run, raw, report):
    n, arrays, label = run["n"], run["arrays"], run["label"]
    _, volume = geometry(n)
    sample_indices = np.arange(21)*(n//20)
    for field, key, margin in (("temperature", "temperature_C", 1e-7),
                               ("moisture", "moisture_kg_kg", 1e-8)):
        capacity, _, initial = field_parameters(field)
        values = arrays[key]
        snapshots = arrays[field+"_snapshots"]
        accum = arrays[field+"_integral"]
        accum_snap = arrays[field+"_integral_snapshots"]
        mean = arrays[field+"_mean"]
        nodes = raw.data[:, 1 if field == "temperature" else 2]
        low, high = min(initial, float(nodes.min())), max(initial, float(nodes.max()))
        actual_low = min(float(values.min()), float(snapshots.min()))
        actual_high = max(float(values.max()), float(snapshots.max()))
        report.add(label + "_" + field + "_bounds",
                   actual_low >= low-margin and actual_high <= high+margin,
                   observed_min=actual_low, observed_max=actual_high,
                   allowed_min=low-margin, allowed_max=high+margin)
        report.add(label + "_" + field + "_initial",
                   np.all(values[0] == initial) and np.all(snapshots[0] == initial) and
                   mean[0] == initial and accum[0] == 0.0 and accum_snap[0] == 0.0)
        report.add(label + "_" + field + "_snapshot_samples",
                   np.array_equal(snapshots[:, sample_indices], values[SNAP_TIMES]))
        report.add(label + "_" + field + "_snapshot_accumulator",
                   np.array_equal(accum_snap, accum[SNAP_TIMES]))
        reconstructed_mean = 2/R**2 * (snapshots @ volume)
        report.metric(label + "_" + field + "_independent_snapshot_mean",
                      np.max(np.abs(reconstructed_mean-mean[SNAP_TIMES])), 1e-10)
        storage = capacity*((snapshots-initial) @ volume)
        scale = capacity*(R**2/2)*max(1.0, abs(initial))
        balance = np.abs(storage-accum_snap)/scale
        report.metric(label + "_" + field + "_snapshot_integral_balance",
                      np.max(balance), 1e-7, worst_time_s=int(SNAP_TIMES[np.argmax(balance)]),
                      interpretation="Independent full-grid storage; same-integrator accumulated flux")
        internal = np.abs(capacity*(R**2/2)*(mean-initial)-accum)/scale
        report.metric(label + "_" + field + "_all_time_internal_balance",
                      np.max(internal), 1e-7, worst_time_s=int(np.argmax(internal)),
                      interpretation="Internal consistency of stored mean and accumulated flux")


@lru_cache(maxsize=12)
def eigenmodes(biot, modes):
    """Roots are bracketed without the poles introduced by dividing by J0."""
    zeros0 = jn_zeros(0, modes)
    zeros1 = jn_zeros(1, max(1, modes-1))
    roots = np.empty(modes)
    for i in range(modes):
        left = 0.0 if i == 0 else float(zeros1[i-1])
        right = float(zeros0[i])
        roots[i] = brentq(lambda z: z*j1(z)-biot*j0(z), left, right,
                          xtol=2e-13, rtol=4*np.finfo(float).eps)
    b0, b1 = j0(roots), j1(roots)
    coefficients = 2*b1/(roots*(b0*b0+b1*b1))
    return roots, coefficients


def modal_reference(raw, field, modes):
    """Continuous-space Robin eigenfunctions with exact linear-input updates."""
    capacity, transfer, initial = field_parameters(field)
    diffusivity = K/capacity if field == "temperature" else D_STAR
    biot = transfer*R/K if field == "temperature" else transfer*R/D_STAR
    roots, coefficients = eigenmodes(float(biot), int(modes))
    rates = diffusivity*roots**2/R**2
    basis = j0(roots[:, None]*(RADII/R)[None, :])
    environment = raw.data[:, 1 if field == "temperature" else 2]
    z = np.full(modes, initial-environment[0])
    out = np.empty((1801, 21))
    out[0] = initial
    dt = np.arange(1.0, 61.0)[:, None]
    decay = np.exp(-dt*rates[None, :])
    ramp = -np.expm1(-dt*rates[None, :])/rates[None, :]
    for segment in range(30):
        slope = (environment[segment+1]-environment[segment])/60.0
        modal_states = decay*z[None, :] - slope*ramp
        out[segment*60+1:segment*60+61] = (
            environment[segment] + slope*dt +
            (modal_states*coefficients[None, :]) @ basis)
        z = modal_states[-1].copy()
    residual = np.max(np.abs(roots*j1(roots)-biot*j0(roots)))
    return out, float(residual)


def converged_reference(raw, field, report):
    modes = 256
    previous, _ = modal_reference(raw, field, modes)
    history = []
    while modes < 16384:
        modes *= 2
        current, root_residual = modal_reference(raw, field, modes)
        delta = compare_grid_values(previous, current)
        history.append(dict(modes=modes, **delta, maximum_root_residual=root_residual))
        if delta["max_abs"] <= 1e-7:
            report.metric(field + "_reference_modal_truncation", delta["max_abs"], 1e-7,
                          modes=modes, prior_modes=modes//2, history=history)
            return current, modes
        previous = current
    report.add(field + "_reference_modal_truncation", False, history=history,
               message="Resource cap reached; reference is not accepted; increase cap deliberately")
    raise RuntimeError("Bessel reference did not meet the frozen truncation criterion")


def verify(args, report):
    boundary_path = ROOT/"output/ENV/q1_boundary.json"
    raw = RawBoundary(ROOT/"附件/附件1.xlsx")
    report.details["identities"] = {
        "specification": {"path": str(ROOT/"planning/Q1/model_spec.md"),
                          "sha256": digest(ROOT/"planning/Q1/model_spec.md")},
        "boundary": {"path": str(boundary_path), "sha256": digest(boundary_path)},
        "raw_environment": {"path": str(raw.path), "sha256": digest(raw.path)},
        "verifier": {"path": str(Path(__file__).resolve()), "sha256": digest(__file__)},
        "solver": {"path": str(ROOT/"src/Q1/solve.py"), "sha256": digest(ROOT/"src/Q1/solve.py")},
    }
    report.details["environment"] = dict(python=platform.python_version(),
                                         numpy=np.__version__, scipy=scipy.__version__)
    print("Checking raw inputs, formulas and boundary branches", flush=True)
    primary_boundary = check_inputs(report, raw, boundary_path)
    check_implementation(report, raw)
    runs = [load_run(path, report, "run_"+str(i)) for i, path in enumerate(args.runs)]
    runs.sort(key=lambda item: item["n"])
    if len(runs) < 3 or len({run["n"] for run in runs}) != len(runs):
        raise ValueError("At least three distinct increasing mesh sizes are required")
    report.add("spatial_refinement_ratio",
               all(b["n"] == 2*a["n"] for a, b in zip(runs[:-1], runs[1:])),
               mesh_sizes=[run["n"] for run in runs])
    for run in runs:
        report.add(run["label"]+"_baseline_tolerances",
                   np.isclose(run["meta"]["rtol"], 2e-9, rtol=1e-12, atol=0) and
                   np.isclose(run["meta"]["atol_scale"], 1.0, rtol=1e-12, atol=0))
        check_run_fields(run, raw, report)
    tight = load_run(args.tight, report, "tight")
    check_run_fields(tight, raw, report)
    fine = runs[-1]
    report.add("tight_mesh_matches", tight["n"] == fine["n"])
    report.add("tight_tolerances",
               np.isclose(tight["meta"]["rtol"], fine["meta"]["rtol"]/10, rtol=1e-12, atol=0) and
               np.isclose(tight["meta"]["atol_scale"], fine["meta"]["atol_scale"]/10,
                          rtol=1e-12, atol=0))
    report.details["runs"] = [{k: run[k] for k in
                              ("path", "sha256", "metadata_path", "metadata_sha256", "n", "meta")}
                             for run in runs+[tight]]
    for field, key in (("temperature", "temperature_C"), ("moisture", "moisture_kg_kg")):
        comparisons = []
        for coarse, refined in zip(runs[:-1], runs[1:]):
            comparisons.append(dict(coarse_n=coarse["n"], fine_n=refined["n"],
                                    **compare_grid_values(coarse["arrays"][key], refined["arrays"][key])))
        errors = [item["max_abs"] for item in comparisons]
        report.add(field+"_spatial_differences_decrease",
                   all(b < a or (a == 0 and b == 0) for a, b in zip(errors[:-1], errors[1:])),
                   comparisons=comparisons,
                   observed_orders=[float(np.log2(a/b)) if a > 0 and b > 0 else None
                                    for a, b in zip(errors[:-1], errors[1:])])
        report.metric(field+"_spatial_final_difference", errors[-1], 2e-5,
                      **{k: v for k, v in comparisons[-1].items() if k != "max_abs"})
        timing = compare_grid_values(fine["arrays"][key], tight["arrays"][key])
        report.metric(field+"_time_tolerance_difference", timing["max_abs"], 5e-6,
                      **{k: v for k, v in timing.items() if k != "max_abs"})

    print("Constructing independent actual-input temperature reference", flush=True)
    temperature_reference, temperature_modes = converged_reference(raw, "temperature", report)
    error = compare_grid_values(fine["arrays"]["temperature_C"], temperature_reference)
    report.metric("actual_temperature_bessel_reference", error["max_abs"], 5e-5,
                  modes=temperature_modes, **{k: v for k, v in error.items() if k != "max_abs"})

    print("Solving controlled constant-D moisture case for independent comparison", flush=True)
    constant_case = dut.solve_field(fine["n"], "moisture", primary_boundary,
                                   rtol=2e-9, atol_scale=1.0, constant_diffusivity=True)
    if constant_case["values"].shape != (1801, 21) or not np.isfinite(constant_case["values"]).all():
        raise ValueError("Constant-D controlled solve returned invalid values")
    moisture_reference, moisture_modes = converged_reference(raw, "moisture", report)
    error = compare_grid_values(constant_case["values"], moisture_reference)
    report.metric("constant_D_moisture_bessel_reference", error["max_abs"], 5e-5,
                  modes=moisture_modes, n=fine["n"], D_star=float(D_STAR),
                  **{k: v for k, v in error.items() if k != "max_abs"})
    constant_arrays = {"moisture_kg_kg": constant_case["values"],
                       "moisture_snapshots": constant_case["snapshots"],
                       "moisture_integral_snapshots": constant_case["integral_snapshots"]}
    _, volume = geometry(fine["n"])
    residual = np.abs((constant_arrays["moisture_snapshots"]-2.55) @ volume -
                      constant_arrays["moisture_integral_snapshots"]) / ((R**2/2)*2.55)
    report.metric("constant_D_snapshot_integral_balance", np.max(residual), 1e-7)
    report.details["constant_D_run"] = {k: constant_case[k] for k in
                                       ("n", "nfev", "njev", "nlu", "elapsed_s")}
    report.details["reference_method"] = dict(
        geometry="Continuous cylindrical radius; independent of finite-volume mesh",
        root_equation="mu*J1(mu)=Bi*J0(mu)",
        coefficient="2*J1(mu)/(mu*(J0(mu)^2+J1(mu)^2))",
        modal_update="z'=-(diffusivity*mu^2/R^2)*z-g'(t); exact within each linear input segment",
        source_urls=["https://dlmf.nist.gov/10.6#E3", "https://dlmf.nist.gov/10.22#E38"],
        source_scope="NIST derivative identity and weighted Robin-root orthogonality",
        initial_time="t=0 returned exactly; reference errors measured at 1..1800 seconds",
    )


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--runs", nargs="+", type=Path, required=True)
    parser.add_argument("--tight", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    if args.output.suffix.lower() != ".json":
        parser.error("--output must name a JSON report")
    if args.output.exists():
        parser.error("Choose a new report path; existing verification evidence is not overwritten")
    report = Report()
    try:
        verify(args, report)
    except Exception as exc:
        report.add("verification_execution", False, exception=type(exc).__name__,
                   message=str(exc), traceback=traceback.format_exc())
    result = report.result()
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2, allow_nan=False)+"\n",
                           encoding="utf-8")
    failures = [item["name"] for item in report.checks if not item["passed"]]
    print(json.dumps(dict(report=str(args.output), passed=result["passed"],
                          check_count=len(report.checks), failed_checks=failures),
                     ensure_ascii=False), flush=True)
    return 0 if result["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())

