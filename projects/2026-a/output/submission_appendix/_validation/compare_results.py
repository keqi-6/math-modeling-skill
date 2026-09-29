"""Internal portability verification; never part of the submission support files."""
from pathlib import Path
import ast
import argparse
import json
import hashlib
import numpy as np
import openpyxl

ROOT = Path(__file__).resolve().parents[3]
HERE = Path(__file__).resolve().parent
CODE = HERE.parent / "support/code"
COMPONENT = "Q1"
CASE = "q1"


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def ast_checks():
    result = []
    for source, target, names in [
        ("src/Q1/solve.py", "q1_model.py", ["FieldSystem", "solve_field"]),
        ("src/Q2/solve.py", "q23_model.py", ["System", "Recorder", "solve"]),
        ("src/Q4/solve.py", "q4_model.py", ["System", "Recorder", "solve"]),
    ]:
        if target != {"q1":"q1_model.py", "q23":"q23_model.py", "q4":"q4_model.py"}[CASE]:
            continue
        old = {n.name: n for n in ast.parse((ROOT/source).read_text(encoding="utf-8")).body
               if isinstance(n, (ast.ClassDef, ast.FunctionDef))}
        new = {n.name: n for n in ast.parse((CODE/target).read_text(encoding="utf-8")).body
               if isinstance(n, (ast.ClassDef, ast.FunctionDef))}
        for name in names:
            diagnostic_strings = {
                "computed_pending_independent_verification": "computed",
                "The full result2 table exceeds Excel's single-sheet row limit":
                    "Dense sampled trajectory exceeds the configured sample limit",
            }
            for node in ast.walk(old[name]):
                if isinstance(node, ast.Constant) and isinstance(node.value, str):
                    node.value = diagnostic_strings.get(node.value, node.value)
            same = (ast.dump(old[name], include_attributes=False)
                    == ast.dump(new[name], include_attributes=False))
            result.append(dict(source=source, target=target, object=name, equal=same))
            assert same, (target, name)
    return result


def numerical_checks():
    results = []
    configs = [
        ("q1", "output/Q1/run_n5120.npz", "output/Q1/run_n5120.json"),
        ("q23", "output/Q2/run_n10240_startsafe.npz", "output/Q2/run_n10240_startsafe.json"),
        ("q4", "output/Q4/run_n10240-v1-20260911.npz", "output/Q4/run_n10240-v1-20260911.json"),
    ]
    for case, original, settings in configs:
        if case != CASE:
            continue
        fresh = HERE / case / f"{case}.npz"
        meta = json.loads((HERE/case/f"{case}.json").read_text(encoding="utf-8"))
        original_meta = json.loads((ROOT/settings).read_text(encoding="utf-8"))
        entry = dict(case=case, elapsed_s=meta["elapsed_s"], fields={}, snapshots={},
                     new_sha256=digest(fresh), original_sha256=digest(ROOT/original))
        with np.load(fresh) as new, np.load(ROOT/original) as old:
            active = new["time_s"] <= 10800 if COMPONENT == "Q2" else np.ones(new["time_s"].shape, dtype=bool)
            indices = np.searchsorted(old["time_s"], new["time_s"][active])
            assert np.array_equal(old["time_s"][indices], new["time_s"][active])
            fields = ["temperature_C", "moisture_kg_kg"]
            fields += (["temperature_mean", "moisture_mean", "temperature_integral", "moisture_integral"]
                       if case == "q1" else ["max_moisture_kg_kg", "temperature_mean", "moisture_mean"])
            if case == "q4":
                fields += ["surface_radius_m", "surface_temperature_C", "surface_moisture_kg_kg"]
                assert np.array_equal(new["inside_mask"], old["inside_mask"][indices])
            for key in fields:
                x, y = new[key][active], old[key][indices]
                assert np.array_equal(np.isnan(x), np.isnan(y)), (case, key, "NaN mask")
                error = float(np.nanmax(np.abs(x-y)))
                entry["fields"][key] = error
                assert error <= (5e-6 if "temperature" in key else 5e-8), (case, key, error)
            time_key = "snapshot_time_s" if case == "q1" else "snapshot_time_s"
            for j, t in enumerate(new[time_key]):
                if COMPONENT == "Q2" and t > 10800:
                    continue
                match = np.flatnonzero(np.isclose(old[time_key], t, rtol=0, atol=1e-6))
                if not match.size:
                    continue
                for key in ("temperature_snapshots", "moisture_snapshots"):
                    error = float(np.max(np.abs(new[key][j]-old[key][int(match[0])])))
                    entry["snapshots"][f"{key}@{t:g}"] = error
                    assert error <= (5e-6 if key.startswith("temperature") else 5e-8), (case, key, t, error)
        if COMPONENT in ("Q3", "Q4"):
            entry["n_end"] = meta["n_end"]
            entry["continuous_crossing_difference_s"] = abs(meta["t_cross_s"]-original_meta["t_cross_s"])
            assert meta["n_end"] == original_meta["n_end"]
            assert entry["continuous_crossing_difference_s"] <= 1e-6
        results.append(entry)
    return results


def workbook_checks():
    results = []
    for number, case in [(1, "q1"), (2, "q23"), (3, "q23"), (4, "q4")]:
        if f"Q{number}" != COMPONENT:
            continue
        original = ROOT/f"output/Q{number}/result{number}.xlsx"
        fresh = HERE/case/f"result{number}.xlsx"
        old = openpyxl.load_workbook(original, read_only=True, data_only=True)
        new = openpyxl.load_workbook(fresh, read_only=True, data_only=True)
        assert old.sheetnames == new.sheetnames
        count, differences = 0, []
        for a, b in zip(old.worksheets, new.worksheets):
            # Write-only originals can omit XML dimension metadata; count cells.
            rowsa, rowsb = list(a.values), list(b.values)
            assert len(rowsa) == len(rowsb), (number, a.title, "row count")
            assert [len(row) for row in rowsa] == [len(row) for row in rowsb], (number, a.title, "column count")
            for r, (rowa, rowb) in enumerate(zip(rowsa, rowsb), 1):
                for c, (x, y) in enumerate(zip(rowa, rowb), 1):
                    count += 1
                    if x != y:
                        differences.append(dict(sheet=a.title, row=r, column=c, old=x, new=y))
        old.close()
        new.close()
        results.append(dict(workbook=f"result{number}.xlsx", compared_cells=count,
                            difference_count=len(differences), examples=differences[:12]))
    return results


def main():
    global COMPONENT, CASE
    parser = argparse.ArgumentParser()
    parser.add_argument("--component", choices=("Q1","Q2","Q3","Q4"), required=True)
    COMPONENT = parser.parse_args().component
    CASE = {"Q1":"q1","Q2":"q23","Q3":"q23","Q4":"q4"}[COMPONENT]
    ast_result = ast_checks()
    report = dict(ast=ast_result, numerical=numerical_checks(), workbooks=workbook_checks())
    report["pass"] = all(x["difference_count"] == 0 for x in report["workbooks"])
    (HERE/f"comparison_{COMPONENT}.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"pass": report["pass"], "numerical": report["numerical"], "workbooks": report["workbooks"]}, ensure_ascii=False))
    assert report["pass"], "Workbook values differ; inspect the saved report"


if __name__ == "__main__":
    main()
