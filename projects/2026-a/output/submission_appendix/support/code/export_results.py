"""Write the four prescribed tables from unrounded numerical arrays."""
from __future__ import annotations
from copy import copy
import json
from pathlib import Path
import numpy as np
import openpyxl
from common import DATA


def four(value):
    return round(float(value), 4) if np.isfinite(value) else None


def write_table(number, output_dir, times, fields, *, surface=None, data_dir=DATA):
    """Expand the original ellipsis to every prescribed 0.1 cm position."""
    output_dir = Path(output_dir)
    target = output_dir / f"result{number}.xlsx"
    if target.exists():
        raise FileExistsError(f"Refusing to overwrite {target.name}")
    book = openpyxl.load_workbook(Path(data_dir) / "templates" / target.name)
    if len(book.worksheets) != len(fields):
        raise ValueError("Template worksheet count differs from requested fields")
    radii = [round(j / 10, 1) for j in range(21)]
    for sheet, values in zip(book.worksheets, fields):
        header, style = sheet.cell(1, 1).value, copy(sheet.cell(1, 2)._style)
        if np.asarray(values).shape != (len(times), 21):
            raise ValueError("Expected one value per time and physical radial position")
        sheet.delete_rows(1, sheet.max_row)
        sheet.append([header, *radii] + (["药材表面"] if surface is not None else []))
        for cell in sheet[1]:
            cell._style = copy(style)
        for j, (t, row) in enumerate(zip(times, values)):
            sheet.append([int(t), *(four(v) for v in row)]
                         + ([four(surface[j])] if surface is not None else []))
        for row in sheet.iter_rows(min_row=2, min_col=2):
            for cell in row:
                cell.number_format = "0.0000"
        sheet.freeze_panes = "B2"
        sheet.column_dimensions["A"].width = 27
    book.save(target)
    book.close()
    return target


def export_case(case, output_dir, data_dir=DATA):
    output_dir = Path(output_dir)
    with np.load(output_dir / f"{case}.npz", allow_pickle=False) as a:
        meta = json.loads((output_dir / f"{case}.json").read_text(encoding="utf-8"))
        times = a["time_s"]
        if case == "q1":
            use = (times >= 1) & (times <= 1800)
            return [write_table(1, output_dir, times[use],
                                [a["temperature_C"][use], a["moisture_kg_kg"][use]], data_dir=data_dir)]
        if case == "q23":
            early = (times >= 1) & (times <= 10800)
            files = [write_table(2, output_dir, times[early],
                                 [a["temperature_C"][early], a["moisture_kg_kg"][early]], data_dir=data_dir)]
            end = int(meta["n_end"])
            use = (times > 0) & (times <= end) & ((times % 60 == 0) | (times == end))
            files.append(write_table(3, output_dir, times[use], [a["moisture_kg_kg"][use]], data_dir=data_dir))
            return files
        if case == "q4":
            use = a["official_output_mask"]
            moisture = a["moisture_kg_kg"][use].copy()
            moisture[~a["inside_mask"][use]] = np.nan
            return [write_table(4, output_dir, times[use], [moisture],
                                surface=a["surface_moisture_kg_kg"][use], data_dir=data_dir)]
    raise ValueError("Unknown calculation")
