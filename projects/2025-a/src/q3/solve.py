"""Q3 working solver: shared-track service curve and three-event scheduling.

The physical model remains the eight-variable FY1/three-bomb problem. For a
fixed UAV track and desired observation time, intersection with the
missile--target centre sightline gives explosion time and LOS fraction; smoke
sinking then gives fuse delay and release time. The resulting service curve is
scheduled under the one-second release-gap constraint. Finalists return to the
original eight variables and complete visible-cylinder joint evaluator.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import platform
import shutil
import time
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable, Sequence

import numpy as np
import scipy
from openpyxl import load_workbook

try:
    from . import selection_pilot as kernel
except ImportError:
    import selection_pilot as kernel


ROOT = Path(__file__).resolve().parents[2]
SPEC_ID = "SPEC-Q3-WORKING-2.0"
RESULT_ID = "Q3-SPEC-Q3-WORKING-2.0-TRACK-SERVICE-SCHEDULE"
OFFICIAL_SHA256 = "a37f6aad30b16ea09da9cb320ce194b12d3a28874008cf9ff19549f6c6079447"
TEMPLATE_SHA256 = "af04b16e6a4719628971bcf5a03d230c9da6738e67eebac9276d254fdd4df1a7"

PROXY_STEP_S = 0.05
OBSERVATION_STEP_S = 0.08
EVENTS_PER_TRACK = 28
SCHEDULES_PER_TRACK = 3
INITIAL_TRACK_KEEP = 10
STRICT_TRACK_KEEP = 18
REFINEMENT_STARTS = 6
FINAL_PRECISE = 3
RELEASE_GAP_S = 1.0

FAST = {
    "name": "fast_strict",
    "scan_step_s": 0.05,
    "n_theta": 64,
    "n_levels": 5,
    "root_time_tolerance_s": 2e-6,
    "margin_tolerance_m": 2e-6,
    "surface_refinement": False,
}

LEGACY_INCUMBENT = {
    "heading_rad": 0.11095698283407587,
    "speed_mps": 126.40484451116781,
    "release_times_s": (0.0, 1.0, 2.0),
    "fuse_delays_s": (0.026525697111554575, 0.0, 0.12235718186154557),
    "reference_duration_s": 6.391870728368502,
}


class Q3UnifiedError(RuntimeError):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code


@dataclass(frozen=True)
class ServiceEvent:
    observation_time_s: float
    explosion_time_s: float
    sight_fraction: float
    smoke_age_s: float
    release_time_s: float
    fuse_delay_s: float
    explosion_height_m: float
    horizontal_residual_m: float
    vertical_residual_m: float


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def ready(value):
    if isinstance(value, np.ndarray):
        return value.tolist()
    if isinstance(value, np.generic):
        return value.item()
    if isinstance(value, dict):
        return {key: ready(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [ready(item) for item in value]
    return value


def verify_identities() -> dict[str, str]:
    paths = {
        "official_input_sha256": ROOT / "data/A题.pdf",
        "excel_template_sha256": ROOT / "data/附件/result1.xlsx",
        "working_spec_sha256": ROOT / "planning/20_q3_unified_solver_spec.md",
        "q1_model_sha256": ROOT / "src/q1/model.py",
        "q2_result_sha256": ROOT / "docs/q2_result.json",
        "joint_kernel_sha256": ROOT / "src/q3/selection_pilot.py",
    }
    observed = {key: sha256(path) for key, path in paths.items()}
    if observed["official_input_sha256"] != OFFICIAL_SHA256:
        raise Q3UnifiedError("Q3_IDENTITY_MISMATCH", "official problem identity changed")
    if observed["excel_template_sha256"] != TEMPLATE_SHA256:
        raise Q3UnifiedError("Q3_IDENTITY_MISMATCH", "official result1 template identity changed")
    return observed


def global_contribution_time_cap() -> float:
    p = kernel.base_parameters()
    direction = p.missile_initial / np.linalg.norm(p.missile_initial)
    missile_x_speed = p.missile_speed * float(direction[0])
    relative_speed = missile_x_speed - 140.0
    if relative_speed <= 0.0:
        return kernel.arrival_time()
    cap = (
        float(p.missile_initial[0]) + p.smoke_radius + kernel.LENGTH_TOL
        - float(p.uav_initial[0])
    ) / relative_speed
    target_x_max = float(p.target_base_center[0]) + p.target_radius
    if float(kernel.q1.missile_position(cap, p)[0]) < target_x_max:
        return kernel.arrival_time()
    return float(min(max(cap, 0.0), kernel.arrival_time()))


def service_event_for_track(
    heading_rad: float,
    speed_mps: float,
    observation_time_s: float,
) -> ServiceEvent | None:
    """Invert one centre-sightline service state on a fixed UAV track."""
    p = kernel.base_parameters()
    t_m = kernel.arrival_time()
    if not (0.0 < observation_time_s < min(t_m, global_contribution_time_cap())):
        return None
    if not (0.0 <= heading_rad < 2.0 * math.pi and 70.0 <= speed_mps <= 140.0):
        return None

    missile = kernel.q1.missile_position(observation_time_s, p)
    target = p.target_base_center + np.array([0.0, 0.0, 0.5 * p.target_height])
    sight_xy = target[:2] - missile[:2]
    direction = np.array([math.cos(heading_rad), math.sin(heading_rad)], dtype=float)
    matrix = np.column_stack((speed_mps * direction, -sight_xy))
    if abs(float(np.linalg.det(matrix))) <= 1e-8:
        return None
    explosion_time, fraction = np.linalg.solve(matrix, missile[:2] - p.uav_initial[:2])
    explosion_time = float(explosion_time)
    fraction = float(fraction)
    age = float(observation_time_s - explosion_time)
    if not (0.0 < fraction < 1.0 and 0.0 <= age <= p.smoke_duration):
        return None
    if not (0.0 < explosion_time <= min(observation_time_s, t_m)):
        return None

    desired_cloud = missile + fraction * (target - missile)
    explosion_height = float(desired_cloud[2] + p.smoke_sink_speed * age)
    if not (0.0 <= explosion_height <= p.uav_initial[2]):
        return None
    fuse_delay = math.sqrt(
        max(0.0, 2.0 * (float(p.uav_initial[2]) - explosion_height) / p.gravity)
    )
    release_time = explosion_time - fuse_delay
    if release_time < -kernel.TIME_TOL:
        return None
    release_time = max(0.0, float(release_time))

    explosion_xy = p.uav_initial[:2] + speed_mps * explosion_time * direction
    horizontal_residual = float(np.linalg.norm(explosion_xy - desired_cloud[:2]))
    reconstructed_height = float(p.uav_initial[2] - 0.5 * p.gravity * fuse_delay * fuse_delay)
    vertical_residual = abs(
        reconstructed_height - p.smoke_sink_speed * age - float(desired_cloud[2])
    )
    if horizontal_residual > 1e-6 or vertical_residual > 1e-7:
        return None
    return ServiceEvent(
        observation_time_s=float(observation_time_s),
        explosion_time_s=explosion_time,
        sight_fraction=fraction,
        smoke_age_s=age,
        release_time_s=release_time,
        fuse_delay_s=float(fuse_delay),
        explosion_height_m=explosion_height,
        horizontal_residual_m=horizontal_residual,
        vertical_residual_m=vertical_residual,
    )


def event_record(event: ServiceEvent) -> dict:
    return {
        "observation_time_s": event.observation_time_s,
        "explosion_time_s": event.explosion_time_s,
        "sight_fraction": event.sight_fraction,
        "smoke_age_s": event.smoke_age_s,
        "release_time_s": event.release_time_s,
        "fuse_delay_s": event.fuse_delay_s,
        "explosion_height_m": event.explosion_height_m,
        "horizontal_residual_m": event.horizontal_residual_m,
        "vertical_residual_m": event.vertical_residual_m,
    }


def service_curve(heading_rad: float, speed_mps: float) -> list[ServiceEvent]:
    stop = global_contribution_time_cap()
    observations = np.arange(0.20, stop + 0.5 * OBSERVATION_STEP_S, OBSERVATION_STEP_S)
    values: list[ServiceEvent] = []
    observed: set[tuple[int, int]] = set()
    for observation in observations:
        event = service_event_for_track(heading_rad, speed_mps, float(observation))
        if event is None:
            continue
        key = (round(event.release_time_s / 0.01), round(event.fuse_delay_s / 0.01))
        if key in observed:
            continue
        observed.add(key)
        values.append(event)
    return sorted(values, key=lambda item: (item.observation_time_s, item.release_time_s))


def approximate_service_mask(
    heading_rad: float,
    speed_mps: float,
    event: ServiceEvent,
    times: np.ndarray,
) -> tuple[np.ndarray, np.ndarray]:
    """Conservative centre-bundle proxy used only for schedule generation."""
    p = kernel.base_parameters()
    direction = np.array([math.cos(heading_rad), math.sin(heading_rad), 0.0])
    explosion = p.uav_initial + speed_mps * event.explosion_time_s * direction
    explosion[2] = event.explosion_height_m
    centers = np.repeat(explosion.reshape(1, 3), len(times), axis=0)
    centers[:, 2] -= p.smoke_sink_speed * (times - event.explosion_time_s)
    missiles = np.vstack([kernel.q1.missile_position(float(t), p) for t in times])
    target = p.target_base_center + np.array([0.0, 0.0, 0.5 * p.target_height])
    vectors = target.reshape(1, 3) - missiles
    denominator = np.einsum("ij,ij->i", vectors, vectors)
    offsets = centers - missiles
    lam = np.clip(np.einsum("ij,ij->i", offsets, vectors) / denominator, 0.0, 1.0)
    nearest = missiles + lam[:, None] * vectors
    centre_distance = np.linalg.norm(centers - nearest, axis=1)
    target_bound = math.hypot(p.target_radius, 0.5 * p.target_height)
    margin = centre_distance + lam * target_bound - p.smoke_radius
    active = (
        (times >= event.explosion_time_s - kernel.TIME_TOL)
        & (times <= min(event.explosion_time_s + p.smoke_duration, kernel.arrival_time()) + kernel.TIME_TOL)
    )
    return active & (margin <= 0.0), margin


def select_curve_events(
    heading_rad: float,
    speed_mps: float,
    curve: Sequence[ServiceEvent],
    times: np.ndarray,
) -> list[dict]:
    records = []
    for event in curve:
        mask, margin = approximate_service_mask(heading_rad, speed_mps, event, times)
        records.append({
            "event": event,
            "mask": mask,
            "mask_count": int(np.sum(mask)),
            "minimum_proxy_margin_m": float(np.min(margin)),
        })
    if len(records) <= EVENTS_PER_TRACK:
        return records

    chosen: dict[tuple[int, int], dict] = {}
    for item in sorted(
        records,
        key=lambda value: (-value["mask_count"], value["minimum_proxy_margin_m"]),
    )[:8]:
        event = item["event"]
        chosen[(round(event.release_time_s / 0.01), round(event.fuse_delay_s / 0.01))] = item
    order = np.argsort([item["event"].observation_time_s for item in records])
    for group in np.array_split(order, EVENTS_PER_TRACK - len(chosen)):
        if len(group) == 0:
            continue
        item = max(
            (records[int(index)] for index in group),
            key=lambda value: (value["mask_count"], -value["minimum_proxy_margin_m"]),
        )
        event = item["event"]
        chosen[(round(event.release_time_s / 0.01), round(event.fuse_delay_s / 0.01))] = item
    return sorted(
        chosen.values(),
        key=lambda item: (item["event"].release_time_s, item["event"].observation_time_s),
    )[:EVENTS_PER_TRACK]


def release_compatible(events: Sequence[ServiceEvent]) -> bool:
    release = sorted(event.release_time_s for event in events)
    return all(
        right - left >= RELEASE_GAP_S - kernel.TIME_TOL
        for left, right in zip(release[:-1], release[1:])
    )


def decision_from_events(
    heading_rad: float,
    speed_mps: float,
    events: Sequence[ServiceEvent],
    source: str,
    proposal_index: int,
) -> kernel.Decision:
    ordered = sorted(events, key=lambda item: item.release_time_s)
    if len(ordered) != 3 or not release_compatible(ordered):
        raise Q3UnifiedError("Q3_EVENT_SCHEDULE_FAILURE", "three release-compatible events required")
    decision = kernel.Decision(
        heading_rad=float(heading_rad % (2.0 * math.pi)),
        speed_mps=float(speed_mps),
        release_times_s=tuple(float(item.release_time_s) for item in ordered),
        fuse_delays_s=tuple(float(item.fuse_delay_s) for item in ordered),
        source=source,
        proposal_index=int(proposal_index),
    )
    try:
        kernel.validate_decision(decision)
    except kernel.PilotError as error:
        raise Q3UnifiedError(error.code, str(error)) from error
    return decision


def schedule_curve(
    heading_rad: float,
    speed_mps: float,
    selected: Sequence[dict],
    times: np.ndarray,
    source: str,
) -> list[dict]:
    if len(selected) < 3:
        return []
    pair_beam = []
    for first in range(len(selected)):
        for second in range(first + 1, len(selected)):
            events = [selected[first]["event"], selected[second]["event"]]
            if not release_compatible(events):
                continue
            union = selected[first]["mask"] | selected[second]["mask"]
            pair_beam.append((
                int(np.sum(union)),
                selected[first]["mask_count"] + selected[second]["mask_count"],
                -sum(event.release_time_s for event in events),
                first, second, union,
            ))
    pair_beam.sort(key=lambda item: item[:3], reverse=True)
    pair_beam = pair_beam[:80]

    triples = []
    observed: set[tuple[int, int, int]] = set()
    for pair in pair_beam:
        first, second, pair_union = pair[3], pair[4], pair[5]
        for third in range(len(selected)):
            indices = tuple(sorted((first, second, third)))
            if len(set(indices)) != 3 or indices in observed:
                continue
            events = [selected[index]["event"] for index in indices]
            if not release_compatible(events):
                continue
            observed.add(indices)
            union = pair_union | selected[third]["mask"]
            active_times = times[union]
            span = float(active_times[-1] - active_times[0]) if len(active_times) else 0.0
            triples.append((
                int(np.sum(union)),
                sum(selected[index]["mask_count"] for index in indices),
                span,
                -sum(event.release_time_s for event in events),
                indices,
            ))
    triples.sort(key=lambda item: item[:4], reverse=True)
    output = []
    for rank, item in enumerate(triples[:SCHEDULES_PER_TRACK], 1):
        events = [selected[index]["event"] for index in item[-1]]
        decision = decision_from_events(
            heading_rad, speed_mps, events, f"{source}_schedule_{rank}", rank,
        )
        output.append({
            "decision": decision,
            "events": [event_record(event) for event in sorted(events, key=lambda value: value.release_time_s)],
            "proxy_union_duration_s": item[0] * PROXY_STEP_S,
            "proxy_individual_total_s": item[1] * PROXY_STEP_S,
            "proxy_span_s": item[2],
        })
    return output


def load_q2_track() -> tuple[float, float, kernel.Decision]:
    data = json.loads((ROOT / "docs/q2_result.json").read_text(encoding="utf-8"))
    best = data["formal_best"]["decision"]
    tau = float(best["release_time_s"])
    delta = float(best["fuse_delay_s"])
    decision = kernel.Decision(
        float(best["heading_rad"]), float(best["speed_mps"]),
        (tau, tau + 1.0, tau + 2.0), (delta, delta, delta),
        "q2_staggered_regression", 0,
    )
    kernel.validate_decision(decision)
    return decision.heading_rad, decision.speed_mps, decision


def legacy_decision() -> kernel.Decision:
    decision = kernel.Decision(
        LEGACY_INCUMBENT["heading_rad"], LEGACY_INCUMBENT["speed_mps"],
        tuple(LEGACY_INCUMBENT["release_times_s"]),
        tuple(LEGACY_INCUMBENT["fuse_delays_s"]),
        "legacy_feasible_incumbent", 0,
    )
    kernel.validate_decision(decision)
    return decision


def track_key(heading_rad: float, speed_mps: float) -> tuple[int, int]:
    return (round((heading_rad % (2.0 * math.pi)) / 2e-4), round(speed_mps / 0.05))


def initial_track_bank() -> list[dict]:
    q2_heading, q2_speed, _ = load_q2_track()
    legacy = legacy_decision()
    tracks: dict[tuple[int, int], dict] = {}

    def add(heading: float, speed: float, source: str) -> None:
        if not 70.0 <= speed <= 140.0:
            return
        heading %= 2.0 * math.pi
        tracks.setdefault(track_key(heading, speed), {
            "heading_rad": float(heading), "speed_mps": float(speed), "source": source,
        })

    add(q2_heading, q2_speed, "q2_track")
    add(legacy.heading_rad, legacy.speed_mps, "legacy_q3_track")
    for base_heading, base_speed, label in (
        (q2_heading, q2_speed, "q2_neighbourhood"),
        (legacy.heading_rad, legacy.speed_mps, "legacy_neighbourhood"),
    ):
        for d_heading in (-0.04, -0.02, 0.0, 0.02, 0.04):
            for d_speed in (-20.0, -10.0, 0.0, 10.0, 20.0):
                add(base_heading + d_heading, base_speed + d_speed, label)

    p = kernel.base_parameters()
    for observation in np.linspace(0.8, 12.8, 13):
        for age in (0.0, 0.4, 1.0, 2.0, 4.0):
            if age >= observation:
                continue
            intervals = kernel.q2_inverse.q_intervals(float(observation), float(age), p)
            for left, right in intervals:
                for position in (0.25, 0.5, 0.75):
                    fraction = float(left + position * (right - left))
                    try:
                        anchor = kernel.inverse_service_state(
                            float(observation), float(age), fraction,
                            "analytic_service_anchor", len(tracks),
                        )
                    except kernel.PilotError:
                        continue
                    add(anchor.heading_rad, anchor.speed_mps, "analytic_service_anchor")

    for heading in np.linspace(0.0, 2.0 * math.pi, 16, endpoint=False):
        for speed in (70.0, 105.0, 140.0):
            add(float(heading), speed, "full_heading_guard")
    return list(tracks.values())


def score_track(track: dict, times: np.ndarray) -> dict | None:
    curve = service_curve(track["heading_rad"], track["speed_mps"])
    selected = select_curve_events(track["heading_rad"], track["speed_mps"], curve, times)
    schedules = schedule_curve(
        track["heading_rad"], track["speed_mps"], selected, times, track["source"],
    )
    if not schedules:
        return None
    return {
        **track,
        "service_curve_event_count": len(curve),
        "selected_event_count": len(selected),
        "best_proxy_union_duration_s": schedules[0]["proxy_union_duration_s"],
        "best_proxy_span_s": schedules[0]["proxy_span_s"],
        "schedules": schedules,
    }


def neighbourhood_tracks(records: Sequence[dict]) -> list[dict]:
    values: dict[tuple[int, int], dict] = {}
    for rank, record in enumerate(records[:INITIAL_TRACK_KEEP], 1):
        for d_heading in (-0.012, -0.006, 0.0, 0.006, 0.012):
            for d_speed in (-8.0, -4.0, 0.0, 4.0, 8.0):
                speed = float(record["speed_mps"] + d_speed)
                if not 70.0 <= speed <= 140.0:
                    continue
                heading = float((record["heading_rad"] + d_heading) % (2.0 * math.pi))
                values.setdefault(track_key(heading, speed), {
                    "heading_rad": heading,
                    "speed_mps": speed,
                    "source": f"track_neighbourhood_rank_{rank}",
                })
    return list(values.values())


def evaluate(
    decision: kernel.Decision,
    settings: dict,
    mesh: kernel.SurfaceMesh,
    cloud_indices: Sequence[int] = (0, 1, 2),
    restrict_windows: Sequence[Sequence[float]] | None = None,
) -> dict:
    windows = [[0.0, global_contribution_time_cap()]] if restrict_windows is None else restrict_windows
    try:
        return kernel.solve_intervals(
            decision, settings, cloud_indices=cloud_indices, mesh=mesh,
            restrict_windows=windows,
        )
    except kernel.PilotError as error:
        raise Q3UnifiedError(error.code, str(error)) from error


def decision_to_unit(decision: kernel.Decision) -> np.ndarray:
    t_m = kernel.arrival_time()
    p = kernel.base_parameters()
    delta_cap = math.sqrt(2.0 * p.uav_initial[2] / p.gravity)
    tau = np.asarray(decision.release_times_s, dtype=float)
    delta = np.asarray(decision.fuse_delays_s, dtype=float)
    release_u = np.array([
        tau[0] / (t_m - 2.0),
        (tau[1] - tau[0] - 1.0) / (t_m - tau[0] - 2.0),
        (tau[2] - tau[1] - 1.0) / (t_m - tau[1] - 1.0),
    ])
    fuse_caps = np.minimum(delta_cap, t_m - tau)
    fuse_u = np.divide(delta, fuse_caps, out=np.zeros(3), where=fuse_caps > 0.0)
    return np.clip(np.array([
        decision.heading_rad / (2.0 * math.pi),
        (decision.speed_mps - 70.0) / 70.0,
        *release_u, *fuse_u,
    ]), 0.0, 1.0)


def unit_to_decision(values: Sequence[float], source: str, index: int) -> kernel.Decision:
    z = np.clip(np.asarray(values, dtype=float), 0.0, 1.0)
    if z.size != 8:
        raise Q3UnifiedError("Q3_CONSTRAINT_FAILURE", "eight unit coordinates required")
    t_m = kernel.arrival_time()
    p = kernel.base_parameters()
    delta_cap = math.sqrt(2.0 * p.uav_initial[2] / p.gravity)
    tau1 = (t_m - 2.0) * z[2]
    tau2 = tau1 + 1.0 + (t_m - tau1 - 2.0) * z[3]
    tau3 = tau2 + 1.0 + (t_m - tau2 - 1.0) * z[4]
    tau = np.array([tau1, tau2, tau3])
    caps = np.minimum(delta_cap, t_m - tau)
    delta = z[5:8] * caps
    decision = kernel.Decision(
        heading_rad=float((2.0 * math.pi * z[0]) % (2.0 * math.pi)),
        speed_mps=float(70.0 + 70.0 * z[1]),
        release_times_s=tuple(float(value) for value in tau),
        fuse_delays_s=tuple(float(value) for value in delta),
        source=source,
        proposal_index=int(index),
    )
    try:
        kernel.validate_decision(decision)
    except kernel.PilotError as error:
        raise Q3UnifiedError(error.code, str(error)) from error
    return decision


def pattern_refine(
    start: kernel.Decision,
    start_duration_s: float,
    mesh: kernel.SurfaceMesh,
    ordinal: int,
) -> dict:
    current_z = decision_to_unit(start)
    current = start
    current_result = evaluate(start, FAST, mesh)
    cache: dict[tuple[float, ...], tuple[kernel.Decision, dict]] = {
        tuple(np.round(current_z, 12)): (current, current_result)
    }
    calls = 1
    stages = []

    def get(z: np.ndarray) -> tuple[kernel.Decision, dict]:
        nonlocal calls
        clipped = np.clip(z, 0.0, 1.0)
        key = tuple(np.round(clipped, 12))
        if key not in cache:
            candidate = unit_to_decision(clipped, f"unified_pattern_{ordinal}", ordinal)
            cache[key] = (candidate, evaluate(candidate, FAST, mesh))
            calls += 1
        return cache[key]

    for radius in (0.015, 0.006):
        accepted = 0
        for coordinate in range(8):
            options = [(current, current_result, current_z.copy())]
            for sign in (-1.0, 1.0):
                proposal = current_z.copy()
                proposal[coordinate] = np.clip(
                    proposal[coordinate] + sign * radius, 0.0, 1.0,
                )
                decision, result = get(proposal)
                options.append((decision, result, proposal))
            chosen = max(
                options,
                key=lambda item: (
                    float(item[1]["effective_duration_s"]),
                    item[1]["minimum_scan_margin_m"] * -1.0,
                ),
            )
            if float(chosen[1]["effective_duration_s"]) > float(current_result["effective_duration_s"]) + 1e-8:
                current, current_result, current_z = chosen
                accepted += 1
        stages.append({
            "radius": radius,
            "accepted_coordinate_moves": accepted,
            "duration_s": float(current_result["effective_duration_s"]),
        })
    return {
        "start_decision": kernel.decision_record(start),
        "start_fast_duration_s": float(start_duration_s),
        "chosen_decision": kernel.decision_record(current),
        "chosen_fast": current_result,
        "function_evaluations": calls,
        "stages": stages,
        "_decision": current,
    }


def strict_candidate_record(
    decision: kernel.Decision,
    origin: str,
    proxy: dict | None,
    mesh: kernel.SurfaceMesh,
) -> dict:
    result = evaluate(decision, FAST, mesh)
    return {
        "origin": origin,
        "decision": decision,
        "proxy": proxy,
        "fast": result,
        "minimum_constraint_margin": kernel.minimum_constraint_margin(decision),
    }


def deduplicate_candidates(records: Iterable[dict]) -> list[dict]:
    values: dict[tuple[float, ...], dict] = {}
    for record in records:
        key = kernel.decision_key(record["decision"])
        incumbent = values.get(key)
        if incumbent is None or float(record["fast"]["effective_duration_s"]) > float(incumbent["fast"]["effective_duration_s"]):
            values[key] = record
    return list(values.values())


def expanded(intervals: Sequence[Sequence[float]], pad: float) -> list[list[float]]:
    cap = global_contribution_time_cap()
    return [[max(0.0, float(a) - pad), min(cap, float(b) + pad)] for a, b in intervals]


def final_diagnostics(
    decision: kernel.Decision,
    precise_mesh: kernel.SurfaceMesh,
    dense_mesh: kernel.SurfaceMesh,
) -> dict:
    joint = evaluate(decision, kernel.PRECISE, precise_mesh)
    individuals = [evaluate(decision, kernel.PRECISE, precise_mesh, (index,)) for index in range(3)]
    individual_union = kernel.merge_intervals(
        interval for item in individuals for interval in item["intervals_s"]
    )
    pure = kernel.interval_difference(joint["intervals_s"], individual_union)
    leave_out = []
    for omitted in range(3):
        indices = tuple(index for index in range(3) if index != omitted)
        result = evaluate(decision, kernel.PRECISE, precise_mesh, indices)
        leave_out.append({
            "omitted_bomb": omitted + 1,
            "active_cloud_indices": list(indices),
            "intervals_s": result["intervals_s"],
            "duration_s": result["effective_duration_s"],
            "marginal_contribution_s": float(joint["effective_duration_s"] - result["effective_duration_s"]),
        })

    dense_joint = evaluate(
        decision, kernel.DENSIFIED, dense_mesh,
        restrict_windows=expanded(joint["intervals_s"], kernel.DENSIFIED["scan_step_s"]),
    ) if joint["intervals_s"] else {
        "intervals_s": [], "effective_duration_s": 0.0, "root_fallback_count": 0,
    }
    dense_individuals = []
    for index, item in enumerate(individuals):
        if item["intervals_s"]:
            dense_individuals.append(evaluate(
                decision, kernel.DENSIFIED, dense_mesh, (index,),
                restrict_windows=expanded(item["intervals_s"], kernel.DENSIFIED["scan_step_s"]),
            ))
        else:
            dense_individuals.append({
                "intervals_s": [], "effective_duration_s": 0.0, "root_fallback_count": 0,
            })
    dense_union = kernel.merge_intervals(
        interval for item in dense_individuals for interval in item["intervals_s"]
    )
    dense_pure = kernel.interval_difference(dense_joint["intervals_s"], dense_union)
    individual_sum = float(sum(item["effective_duration_s"] for item in individuals))
    union_duration = kernel.interval_measure(individual_union)
    return {
        "precise": {
            "joint": joint,
            "individuals": individuals,
            "individual_union_intervals_s": individual_union,
            "pure_synergy_intervals_s": pure,
            "pure_synergy_duration_s": kernel.interval_measure(pure),
        },
        "densified": {
            "joint": dense_joint,
            "individuals": dense_individuals,
            "individual_union_intervals_s": dense_union,
            "pure_synergy_intervals_s": dense_pure,
            "pure_synergy_duration_s": kernel.interval_measure(dense_pure),
        },
        "densified_joint_change_s": float(
            dense_joint["effective_duration_s"] - joint["effective_duration_s"]
        ),
        "leave_one_out": leave_out,
        "individual_duration_sum_s": individual_sum,
        "individual_union_duration_s": union_duration,
        "individual_overlap_excess_s": individual_sum - union_duration,
        "resource_monotonicity_pass": bool(
            joint["effective_duration_s"] + 1e-8
            >= max(item["effective_duration_s"] for item in individuals)
        ),
    }


def write_excel(
    decision: kernel.Decision,
    individual_durations: Sequence[float],
    excel_path: Path,
    precise_mesh: kernel.SurfaceMesh,
) -> dict:
    template = ROOT / "data/附件/result1.xlsx"
    excel_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(template, excel_path)
    workbook = load_workbook(excel_path)
    sheet = workbook[workbook.sheetnames[0]]
    record = kernel.decision_record(decision)
    for index in range(3):
        values = [
            record["heading_deg"], record["speed_mps"], index + 1,
            *record["release_points_m"][index],
            *record["explosion_points_m"][index],
            float(individual_durations[index]),
        ]
        for column, value in enumerate(values, 1):
            sheet.cell(row=index + 2, column=column, value=int(value) if column == 3 else float(value))
    workbook.save(excel_path)

    reread = load_workbook(excel_path, data_only=True)
    sheet = reread[reread.sheetnames[0]]
    rows = [[sheet.cell(row=row, column=column).value for column in range(1, 11)] for row in range(2, 5)]
    headings = np.radians(np.asarray([row[0] for row in rows], dtype=float))
    speeds = np.asarray([row[1] for row in rows], dtype=float)
    if float(np.ptp(headings)) > 1e-12 or float(np.ptp(speeds)) > 1e-10:
        raise Q3UnifiedError("Q3_EXCEL_ROUNDTRIP_FAILURE", "shared track changed on readback")
    heading = float(headings[0] % (2.0 * math.pi))
    speed = float(speeds[0])
    direction = np.array([math.cos(heading), math.sin(heading)])
    p = kernel.base_parameters()
    release_points = np.asarray([row[3:6] for row in rows], dtype=float)
    explosion_points = np.asarray([row[6:9] for row in rows], dtype=float)
    tau = (release_points[:, :2] - p.uav_initial[:2]) @ direction / speed
    explosion_times = (explosion_points[:, :2] - p.uav_initial[:2]) @ direction / speed
    delta = explosion_times - tau
    reconstructed = kernel.Decision(
        heading, speed,
        tuple(float(value) for value in tau), tuple(float(value) for value in delta),
        "excel_roundtrip", 0,
    )
    try:
        kernel.validate_decision(reconstructed)
    except kernel.PilotError as error:
        raise Q3UnifiedError("Q3_EXCEL_ROUNDTRIP_FAILURE", str(error)) from error
    reread_joint = evaluate(reconstructed, kernel.PRECISE, precise_mesh)
    reread_individual = [evaluate(reconstructed, kernel.PRECISE, precise_mesh, (index,)) for index in range(3)]
    cell_durations = np.asarray([row[9] for row in rows], dtype=float)
    recomputed = np.asarray([item["effective_duration_s"] for item in reread_individual])
    if float(np.max(np.abs(cell_durations - recomputed))) > 1e-6:
        raise Q3UnifiedError("Q3_EXCEL_ROUNDTRIP_FAILURE", "row duration changed on readback")
    return {
        "path": excel_path.relative_to(ROOT).as_posix(),
        "sha256": sha256(excel_path),
        "status": "pass",
        "reconstructed_decision": kernel.decision_record(reconstructed),
        "recomputed_joint": reread_joint,
        "row_duration_max_residual_s": float(np.max(np.abs(cell_durations - recomputed))),
    }


def build_report(result: dict, result_path: Path, result_sha: str) -> str:
    best = result["working_best"]
    diagnostics = result["final_diagnostics"]
    individual = [item["effective_duration_s"] for item in diagnostics["precise"]["individuals"]]
    marginal = [item["marginal_contribution_s"] for item in diagnostics["leave_one_out"]]
    search = result["track_search"]
    return f"""# Q3 统一航迹—服务事件调度工作结果

> 本报告绑定 `{SPEC_ID}`，是 Q3 S4 invalidated 状态下的工作结果，不是 S5 准入或最终答案。

## 唯一求解主链

固定共享航迹 `(theta,v)` 后，对观察时刻 `t` 联立无人机水平航迹与 M1—圆柱中心视线，
解析得到起爆时刻和视线比例；再把烟幕每秒下沉 3 m 代入竖向关系，反解引信延迟与投放
时刻。每条航迹因此生成一条物理可行服务事件曲线，内层只从曲线上调度三个相邻投放至少
间隔 1 s 的事件。候选最终返回原八变量，由完整可见圆柱、有限视线段和连续时间集合统一
评价。三弹边际和纯协同只在结果阶段诊断。

## 搜索规模

- 初始航迹：`{search['initial_track_count']}`；存在可调度服务曲线：`{search['initial_schedulable_track_count']}`。
- 优质航迹邻域：`{search['neighbour_track_count']}`；存在可调度服务曲线：`{search['neighbour_schedulable_track_count']}`。
- 严格快速筛选候选：`{search['strict_candidate_count']}`；原八变量局部精修起点：`{search['refinement_start_count']}`。
- 高精度复算候选：`{search['final_precise_count']}`。

## 当前工作最优解

- 航向：`{best['decision']['heading_rad']:.12f} rad = {best['decision']['heading_deg']:.9f}°`。
- 速度：`{best['decision']['speed_mps']:.12f} m/s`。
- 投放时刻：`{best['decision']['release_times_s']} s`。
- 引信延迟：`{best['decision']['fuse_delays_s']} s`。
- 联合遮蔽区间：`{diagnostics['precise']['joint']['intervals_s']} s`。
- 联合时长：`{diagnostics['precise']['joint']['effective_duration_s']:.12f} s`。
- 三弹独立时长：`{individual} s`。
- 删除单弹后的边际贡献：`{marginal} s`。
- 纯协同时长：`{diagnostics['precise']['pure_synergy_duration_s']:.12f} s`。

## 约束、加密与回读

- 三次投放的相邻间隔为 `{best['release_gaps_s']} s`，均不小于 1 s；云团运动显式包含 3 m/s 下沉。
- 加密表面联合时长为 `{diagnostics['densified']['joint']['effective_duration_s']:.12f} s`，相对变化 `{diagnostics['densified_joint_change_s']:.3e} s`。
- 资源单调性：`{diagnostics['resource_monotonicity_pass']}`。
- Excel 工作副本 `{result['excel_roundtrip']['path']}` 写后回读复算：`{result['excel_roundtrip']['status']}`。

## 与旧结果的关系及边界

旧 Q3 的 `6.391870728369 s` 可行解作为 incumbent 进入同一个严格筛选器，防止架构重写造成
无证据退化；它不再作为独立方案或主方法。当前算法是否提高了该下界，以机器结果中的
`comparison_to_legacy` 为准。

机器结果：`{result_path.relative_to(ROOT).as_posix()}`，SHA-256 `{result_sha}`。

本轮只能声称固定候选与精修预算下得到可复现的高质量可行工作解。服务曲线是高命中率
候选生成，不是对原八维可行域的完备参数化；Q3 尚未执行正式 E1—E4，不能声称连续域全局
最优或现实部署有效。
"""


def run(output_path: Path, report_path: Path, excel_path: Path) -> dict:
    started = time.perf_counter()
    identities = verify_identities()
    cap = global_contribution_time_cap()
    proxy_times = np.arange(0.0, cap + 0.5 * PROXY_STEP_S, PROXY_STEP_S)
    fast_mesh = kernel.surface_mesh(FAST["n_theta"], FAST["n_levels"])
    screen_mesh = kernel.surface_mesh(kernel.SCREEN["n_theta"], kernel.SCREEN["n_levels"])
    precise_mesh = kernel.surface_mesh(kernel.PRECISE["n_theta"], kernel.PRECISE["n_levels"])
    dense_mesh = kernel.surface_mesh(kernel.DENSIFIED["n_theta"], kernel.DENSIFIED["n_levels"])

    initial_tracks = initial_track_bank()
    initial_records = []
    for index, track in enumerate(initial_tracks, 1):
        record = score_track(track, proxy_times)
        if record is not None:
            initial_records.append(record)
        if index % 40 == 0:
            print(json.dumps({
                "stage": "initial_track_proxy", "processed": index,
                "total": len(initial_tracks), "schedulable": len(initial_records),
            }), flush=True)
    initial_records.sort(
        key=lambda item: (item["best_proxy_union_duration_s"], item["best_proxy_span_s"]),
        reverse=True,
    )

    neighbour_tracks = neighbourhood_tracks(initial_records)
    existing = {track_key(item["heading_rad"], item["speed_mps"]) for item in initial_tracks}
    neighbour_tracks = [
        item for item in neighbour_tracks
        if track_key(item["heading_rad"], item["speed_mps"]) not in existing
    ]
    neighbour_records = []
    for index, track in enumerate(neighbour_tracks, 1):
        record = score_track(track, proxy_times)
        if record is not None:
            neighbour_records.append(record)
        if index % 50 == 0:
            print(json.dumps({
                "stage": "neighbour_track_proxy", "processed": index,
                "total": len(neighbour_tracks), "schedulable": len(neighbour_records),
            }), flush=True)

    all_track_records = sorted(
        [*initial_records, *neighbour_records],
        key=lambda item: (item["best_proxy_union_duration_s"], item["best_proxy_span_s"]),
        reverse=True,
    )
    strict_records = []
    for track_rank, track in enumerate(all_track_records[:STRICT_TRACK_KEEP], 1):
        for schedule_rank, schedule in enumerate(track["schedules"], 1):
            strict_records.append(strict_candidate_record(
                schedule["decision"],
                f"service_curve_track_{track_rank}_schedule_{schedule_rank}",
                {
                    "track": {
                        "heading_rad": track["heading_rad"],
                        "speed_mps": track["speed_mps"],
                        "source": track["source"],
                    },
                    "events": schedule["events"],
                    "proxy_union_duration_s": schedule["proxy_union_duration_s"],
                    "proxy_span_s": schedule["proxy_span_s"],
                }, fast_mesh,
            ))

    q2_decision = load_q2_track()[2]
    legacy = legacy_decision()
    strict_records.extend([
        strict_candidate_record(legacy, "legacy_incumbent_regression", None, fast_mesh),
        strict_candidate_record(q2_decision, "q2_staggered_regression", None, fast_mesh),
    ])
    strict_records = deduplicate_candidates(strict_records)
    strict_records.sort(
        key=lambda item: (
            float(item["fast"]["effective_duration_s"]), item["minimum_constraint_margin"],
        ), reverse=True,
    )
    if not strict_records or float(strict_records[0]["fast"]["effective_duration_s"]) <= 0.0:
        raise Q3UnifiedError("Q3_UNIFIED_ALL_ZERO", "all strict track schedules are zero")

    refinements = []
    for ordinal, record in enumerate(strict_records[:REFINEMENT_STARTS], 1):
        print(json.dumps({
            "stage": "original8d_pattern_refinement", "ordinal": ordinal,
            "total": min(REFINEMENT_STARTS, len(strict_records)),
            "start_duration_s": record["fast"]["effective_duration_s"],
        }), flush=True)
        refinements.append(pattern_refine(
            record["decision"], record["fast"]["effective_duration_s"], fast_mesh, ordinal,
        ))

    screen_records = []
    for index, refinement in enumerate(refinements, 1):
        decision = refinement["_decision"]
        screen_records.append({
            "refinement_index": index,
            "decision": decision,
            "screen": evaluate(decision, kernel.SCREEN, screen_mesh),
            "refinement": refinement,
            "minimum_constraint_margin": kernel.minimum_constraint_margin(decision),
        })
    screen_records.sort(
        key=lambda item: (
            float(item["screen"]["effective_duration_s"]), item["minimum_constraint_margin"],
        ), reverse=True,
    )

    precise_records = []
    for rank, item in enumerate(screen_records[:FINAL_PRECISE], 1):
        print(json.dumps({"stage": "final_precise", "rank": rank}), flush=True)
        precise_records.append({
            "rank_by_screen": rank,
            "decision": item["decision"],
            "screen": item["screen"],
            "precise": evaluate(item["decision"], kernel.PRECISE, precise_mesh),
            "refinement_index": item["refinement_index"],
            "minimum_constraint_margin": item["minimum_constraint_margin"],
        })
    precise_records.sort(
        key=lambda item: (
            float(item["precise"]["effective_duration_s"]), item["minimum_constraint_margin"],
        ), reverse=True,
    )
    best_object = precise_records[0]
    best_decision = best_object["decision"]
    diagnostics = final_diagnostics(best_decision, precise_mesh, dense_mesh)
    if not diagnostics["resource_monotonicity_pass"]:
        raise Q3UnifiedError("Q3_SURFACE_EVALUATOR_FAILURE", "resource monotonicity failed")
    individual_durations = [
        float(item["effective_duration_s"])
        for item in diagnostics["precise"]["individuals"]
    ]
    excel_roundtrip = write_excel(best_decision, individual_durations, excel_path, precise_mesh)
    roundtrip_change = float(
        excel_roundtrip["recomputed_joint"]["effective_duration_s"]
        - diagnostics["precise"]["joint"]["effective_duration_s"]
    )
    if abs(roundtrip_change) > 1e-6:
        raise Q3UnifiedError("Q3_EXCEL_ROUNDTRIP_FAILURE", "joint duration changed on readback")

    best_record = kernel.decision_record(best_decision)
    release = best_record["release_times_s"]
    working_best = {
        "source_family": "unified_track_service_schedule",
        "decision": best_record,
        "screen": best_object["screen"],
        "precise": best_object["precise"],
        "minimum_constraint_margin": best_object["minimum_constraint_margin"],
        "release_gaps_s": [float(release[1] - release[0]), float(release[2] - release[1])],
    }
    duration = float(diagnostics["precise"]["joint"]["effective_duration_s"])
    result = {
        "schema_version": "2.0-working",
        "result_id": RESULT_ID,
        "question_id": "Q3",
        "spec_id": SPEC_ID,
        "status": "working",
        "purpose": "q3_invalidated_s4_working_replacement_not_formal_result",
        "identities": identities,
        "solver_contract": {
            "primary_solver": "physics_derived_shared_track_service_curve_then_three_event_schedule",
            "formal_variables": ["heading", "speed", "tau_1", "tau_2", "tau_3", "delta_1", "delta_2", "delta_3"],
            "release_gap_s": RELEASE_GAP_S,
            "smoke_sink_speed_mps": kernel.base_parameters().smoke_sink_speed,
            "proxy_step_s": PROXY_STEP_S,
            "observation_step_s": OBSERVATION_STEP_S,
            "events_per_track": EVENTS_PER_TRACK,
            "schedules_per_track": SCHEDULES_PER_TRACK,
            "fast": FAST,
            "screen": kernel.SCREEN,
            "precise": kernel.PRECISE,
            "densified": kernel.DENSIFIED,
            "global_geometry_cap_s": cap,
            "legacy_incumbent_role": "feasible_regression_lower_bound_only",
        },
        "track_search": {
            "initial_track_count": len(initial_tracks),
            "initial_schedulable_track_count": len(initial_records),
            "neighbour_track_count": len(neighbour_tracks),
            "neighbour_schedulable_track_count": len(neighbour_records),
            "strict_track_keep": STRICT_TRACK_KEEP,
            "strict_candidate_count": len(strict_records),
            "refinement_start_count": len(refinements),
            "final_precise_count": len(precise_records),
            "top_track_records": [{
                "heading_rad": item["heading_rad"],
                "speed_mps": item["speed_mps"],
                "source": item["source"],
                "service_curve_event_count": item["service_curve_event_count"],
                "selected_event_count": item["selected_event_count"],
                "best_proxy_union_duration_s": item["best_proxy_union_duration_s"],
                "best_proxy_span_s": item["best_proxy_span_s"],
            } for item in all_track_records[:STRICT_TRACK_KEEP]],
        },
        "strict_candidates": [{
            "origin": item["origin"],
            "decision": kernel.decision_record(item["decision"]),
            "proxy": item["proxy"],
            "fast": item["fast"],
            "minimum_constraint_margin": item["minimum_constraint_margin"],
        } for item in strict_records],
        "refinement_records": [{
            key: value for key, value in item.items() if not key.startswith("_")
        } for item in refinements],
        "final_precise_candidates": [{
            "rank_by_screen": item["rank_by_screen"],
            "decision": kernel.decision_record(item["decision"]),
            "screen": item["screen"],
            "precise": item["precise"],
            "refinement_index": item["refinement_index"],
            "minimum_constraint_margin": item["minimum_constraint_margin"],
        } for item in precise_records],
        "working_best": working_best,
        "formal_best": working_best,
        "formal_best_compatibility_note": "alias only; Q3 remains invalidated S4 and this is not a frozen formal result",
        "final_diagnostics": diagnostics,
        "comparison_to_legacy": {
            "legacy_reference_duration_s": LEGACY_INCUMBENT["reference_duration_s"],
            "working_duration_s": duration,
            "change_s": duration - LEGACY_INCUMBENT["reference_duration_s"],
            "incumbent_not_lost": duration >= LEGACY_INCUMBENT["reference_duration_s"] - 1e-5,
        },
        "excel_roundtrip": {**excel_roundtrip, "joint_duration_change_s": roundtrip_change},
        "limitations": [
            "This is a working S4 result under a new solver specification; E1-E4 and formal checkpoint remain open.",
            "The service curve is a physics-derived high-hit candidate manifold, not a proof that every eight-dimensional optimum lies on it.",
            "A deterministic anchor bank and finite neighbourhood budget do not prove global optimality.",
            "Three-bomb contribution and pure synergy are diagnostics, not constraints or search objectives.",
            "Wind, diffusion, concentration decay, UAV turn dynamics, and execution errors are outside the stated idealized model.",
        ],
        "run_identity": {
            "executed_at": datetime.now(timezone.utc).isoformat(),
            "execution_kind": "deterministic_physics_derived_track_and_event_schedule",
            "python": platform.python_version(),
            "numpy": np.__version__,
            "scipy": scipy.__version__,
            "openpyxl": __import__("openpyxl").__version__,
            "solver_sha256": sha256(ROOT / "src/q3/solve.py"),
            "elapsed_seconds": time.perf_counter() - started,
        },
        "output_paths": {
            "machine_result": output_path.relative_to(ROOT).as_posix(),
            "implementation_report": report_path.relative_to(ROOT).as_posix(),
            "excel_working_copy": excel_path.relative_to(ROOT).as_posix(),
        },
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(ready(result), ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    result_sha = sha256(output_path)
    report_path.write_text(build_report(result, output_path, result_sha), encoding="utf-8")
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--report", type=Path, required=True)
    parser.add_argument("--excel", type=Path, required=True)
    args = parser.parse_args()
    try:
        result = run(ROOT / args.output, ROOT / args.report, ROOT / args.excel)
    except (Q3UnifiedError, kernel.PilotError) as error:
        print(json.dumps({
            "status": "fail",
            "failure_code": getattr(error, "code", "Q3_UNIFIED_FAILURE"),
            "message": str(error),
        }, ensure_ascii=False), flush=True)
        return 2
    print(json.dumps({
        "status": result["status"],
        "result_id": result["result_id"],
        "joint_duration_s": result["working_best"]["precise"]["effective_duration_s"],
        "output": args.output.as_posix(),
        "report": args.report.as_posix(),
        "excel": args.excel.as_posix(),
        "elapsed_seconds": result["run_identity"]["elapsed_seconds"],
    }, ensure_ascii=False), flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
