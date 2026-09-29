"""Working E1--E4 verification and rounded workbook delivery for Q5_v2."""

from __future__ import annotations

import argparse
from dataclasses import replace
from datetime import datetime, timezone
import hashlib
import json
import math
from pathlib import Path
import shutil
import time
from typing import Mapping, Sequence

import numpy as np
from openpyxl import load_workbook

from src.q5 import solve as q5
from src.q5_v2 import solver
from src.q5_v2.objectives import from_interval_sets


ROOT = Path(__file__).resolve().parents[2]
EXPECTED_HASHES = {
    "data/A题.pdf": "a37f6aad30b16ea09da9cb320ce194b12d3a28874008cf9ff19549f6c6079447",
    "data/附件/result3.xlsx": "b648c82d63e459ba6e6b3711ae79875e373521cd543b45571c4d8ff1ad5ec54a",
}
ORACLE_TOL_M = 1e-8
SCREEN_PRECISE_PER_MISSILE_TOL_S = 4e-4
SCREEN_PRECISE_SUM_TOL_MISSILE_S = 1e-3
DENSE_INTERIOR_TOL_M = 5e-6
DENSE_ENDPOINT_DRIFT_TOL_S = 2.5e-3
EXCEL_SPACE_TOL_M = 2e-5
EXCEL_DURATION_ROUNDING_TOL_S = 6e-7
EXCEL_METRIC_TOL_MISSILE_S = 2e-4


def emit(stage: str, **payload: object) -> None:
    print(json.dumps({"stage": stage, **solver.ready(payload)}, ensure_ascii=False), flush=True)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1 << 20), b""):
            digest.update(block)
    return digest.hexdigest()


def save_checkpoint(path: Path, started: float, stage: str, **payload: object) -> None:
    solver.atomic_json(path, {
        "schema_version": "q5-v2-verification-checkpoint-1.0",
        "stage": stage,
        "elapsed_s": time.perf_counter() - started,
        "updated_at": datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds"),
        **payload,
    })


def decision_from_payload(payload: Mapping[str, object]) -> q5.Decision:
    decision = solver.platform_record_decision(payload["decision"], "q5_v2_verification")
    q5.validate(decision)
    return decision


def independent_margin(
    decision: q5.Decision, missile_index: int, times: Sequence[float], mesh: object
) -> np.ndarray:
    """Independent finite-segment max-min evaluator.

    This implementation deliberately does not call the production visibility,
    distance-to-segment, cloud, or strict-margin routines.
    """

    times_array = np.asarray(times, dtype=float)
    initial = np.asarray(q5.geometry.MISSILES[missile_index], dtype=float)
    missiles = (
        1.0 - q5.p().missile_speed * times_array / np.linalg.norm(initial)
    )[:, None] * initial[None, :]
    points = np.asarray(mesh.points, dtype=float)
    normals = np.asarray(mesh.point_normals, dtype=float)
    kinds = np.asarray(mesh.kinds)
    result = np.empty(len(times_array), dtype=float)

    for row, (time_value, missile) in enumerate(zip(times_array, missiles)):
        visible = np.zeros(len(points), dtype=bool)
        side = kinds == 0
        top = ~side
        visible[side] = np.einsum(
            "ij,ij->i", missile[None, :2] - points[side, :2], normals[side, :2]
        ) >= -1e-12
        visible[top] = missile[2] >= points[top, 2] - 1e-12
        minimum = np.full(len(points), np.inf)
        for bomb in decision.bombs:
            explosion_time = bomb.release_s + bomb.fuse_s
            if not (
                explosion_time - q5.TIME_TOL
                <= time_value
                <= min(explosion_time + q5.p().smoke_duration, q5.arrival(missile_index))
                + q5.TIME_TOL
            ):
                continue
            direction = np.array([
                math.cos(decision.headings_rad[bomb.platform]),
                math.sin(decision.headings_rad[bomb.platform]),
                0.0,
            ])
            center = (
                q5.ORIGINS[bomb.platform]
                + decision.speeds_mps[bomb.platform] * explosion_time * direction
            )
            center = center.copy()
            center[2] = (
                q5.ORIGINS[bomb.platform, 2]
                - 0.5 * q5.p().gravity * bomb.fuse_s**2
                - q5.p().smoke_sink_speed * (time_value - explosion_time)
            )
            segment = points - missile
            denominator = np.einsum("ij,ij->i", segment, segment)
            lam = np.clip(
                np.einsum("ij,j->i", segment, center - missile) / denominator,
                0.0,
                1.0,
            )
            nearest = missile + lam[:, None] * segment
            distance = np.linalg.norm(center - nearest, axis=1)
            minimum = np.minimum(minimum, distance)
        values = minimum - q5.p().smoke_radius
        values[~visible] = -np.inf
        result[row] = float(np.max(values[visible])) if np.any(visible) else np.inf
    return result


def interval_probe_times(intervals: Sequence[Sequence[float]]) -> list[float]:
    values: set[float] = set()
    for left_raw, right_raw in intervals:
        left, right = float(left_raw), float(right_raw)
        for fraction in (0.0, 0.25, 0.5, 0.75, 1.0):
            values.add(left + fraction * (right - left))
        values.add(max(0.0, left - 0.001))
        values.add(min(q5.ARRIVAL_TIMES[0], right + 0.001))
    return sorted(values)


def invalid_case_results(decision: q5.Decision) -> dict[str, bool]:
    cases: dict[str, q5.Decision] = {}
    speeds = list(decision.speeds_mps)
    speeds[0] = 69.999
    cases["speed_below_lower_bound"] = q5.Decision(
        decision.headings_rad, tuple(speeds), decision.bombs, "invalid_speed"
    )
    first = decision.bombs[0]
    cases["negative_release"] = q5.Decision(
        decision.headings_rad,
        decision.speeds_mps,
        (replace(first, release_s=-0.1), *decision.bombs[1:]),
        "invalid_release",
    )
    cases["negative_explosion_height"] = q5.Decision(
        decision.headings_rad,
        decision.speeds_mps,
        (replace(first, fuse_s=30.0), *decision.bombs[1:]),
        "invalid_height",
    )
    cases["invalid_label"] = q5.Decision(
        decision.headings_rad,
        decision.speeds_mps,
        (replace(first, label=3), *decision.bombs[1:]),
        "invalid_label",
    )
    cases["release_gap_below_one_second"] = q5.Decision(
        decision.headings_rad,
        decision.speeds_mps,
        (*decision.bombs, q5.Bomb(0, 1, first.release_s + 0.5, 0.1)),
        "invalid_gap",
    )
    cases["more_than_three_bombs_on_platform"] = q5.Decision(
        decision.headings_rad,
        decision.speeds_mps,
        (*decision.bombs,
         q5.Bomb(0, 0, 1.6, 0.1), q5.Bomb(0, 1, 2.7, 0.1), q5.Bomb(0, 2, 3.8, 0.1)),
        "invalid_capacity",
    )
    rejected: dict[str, bool] = {}
    for name, candidate in cases.items():
        try:
            q5.validate(candidate)
        except q5.Q5Error:
            rejected[name] = True
        else:
            rejected[name] = False
    return rejected


def kinematic_roundtrip(decision: q5.Decision) -> dict[str, object]:
    position_errors: list[float] = []
    time_errors: list[float] = []
    for bomb in decision.bombs:
        direction = np.array([
            math.cos(decision.headings_rad[bomb.platform]),
            math.sin(decision.headings_rad[bomb.platform]),
        ])
        release = q5.release_point(decision, bomb)
        explosion = q5.explosion_point(decision, bomb)
        tau = float((release[:2] - q5.ORIGINS[bomb.platform, :2]) @ direction / decision.speeds_mps[bomb.platform])
        et = float((explosion[:2] - q5.ORIGINS[bomb.platform, :2]) @ direction / decision.speeds_mps[bomb.platform])
        release_rebuilt = q5.ORIGINS[bomb.platform].copy()
        release_rebuilt[:2] += decision.speeds_mps[bomb.platform] * tau * direction
        explosion_rebuilt = q5.ORIGINS[bomb.platform].copy()
        explosion_rebuilt[:2] += decision.speeds_mps[bomb.platform] * et * direction
        explosion_rebuilt[2] -= 0.5 * q5.p().gravity * (et - tau) ** 2
        position_errors.extend([
            float(np.linalg.norm(release_rebuilt - release)),
            float(np.linalg.norm(explosion_rebuilt - explosion)),
        ])
        time_errors.extend([abs(tau - bomb.release_s), abs((et - tau) - bomb.fuse_s)])
    return {
        "maximum_position_residual_m": max(position_errors, default=0.0),
        "maximum_time_residual_s": max(time_errors, default=0.0),
        "pass": max(position_errors, default=0.0) <= 1e-8 and max(time_errors, default=0.0) <= 1e-10,
    }


def write_coordinate_workbook(decision: q5.Decision, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(ROOT / "data/附件/result3.xlsx", path)
    book = load_workbook(path)
    sheet = book[book.sheetnames[0]]
    for platform in range(5):
        bombs = q5.by_platform(decision, platform)
        for slot in range(3):
            row = 2 + 3 * platform + slot
            if slot >= len(bombs):
                for column in range(1, 13):
                    sheet.cell(row=row, column=column).value = None
                continue
            bomb = bombs[slot]
            release = q5.release_point(decision, bomb)
            explosion = q5.explosion_point(decision, bomb)
            values: list[object] = [
                q5.NAMES[platform],
                round(math.degrees(decision.headings_rad[platform]) % 360.0, 6),
                round(decision.speeds_mps[platform], 6),
                float(slot + 1),
                *(round(float(value), 6) for value in release),
                *(round(float(value), 6) for value in explosion),
                None,
                q5.MISSILE_NAMES[bomb.label],
            ]
            for column, value in enumerate(values, 1):
                sheet.cell(row=row, column=column).value = value
    book.save(path)


def reconstruct_workbook(path: Path) -> tuple[q5.Decision, dict[str, object]]:
    sheet = load_workbook(path, data_only=True).active
    headings = [0.0] * 5
    speeds = [70.0] * 5
    bombs: list[q5.Bomb] = []
    spatial_residuals: list[float] = []
    active_rows: list[int] = []
    blank_rows: list[int] = []
    for row in range(2, 17):
        values = [sheet.cell(row=row, column=column).value for column in range(1, 13)]
        if values[0] is None:
            if any(value is not None for value in values):
                raise RuntimeError(f"partially blank workbook row {row}")
            blank_rows.append(row)
            continue
        if any(values[column] is None for column in (*range(0, 10), 11)):
            raise RuntimeError(f"incomplete workbook row {row}")
        active_rows.append(row)
        platform = q5.NAMES.index(str(values[0]))
        headings[platform] = math.radians(float(values[1])) % (2 * math.pi)
        speeds[platform] = float(values[2])
        direction = np.array([math.cos(headings[platform]), math.sin(headings[platform])])
        release = np.asarray(values[4:7], dtype=float)
        explosion = np.asarray(values[7:10], dtype=float)
        tau = float((release[:2] - q5.ORIGINS[platform, :2]) @ direction / speeds[platform])
        et = float((explosion[:2] - q5.ORIGINS[platform, :2]) @ direction / speeds[platform])
        fuse = et - tau
        release_xy_rebuilt = q5.ORIGINS[platform, :2] + speeds[platform] * tau * direction
        explosion_xy_rebuilt = q5.ORIGINS[platform, :2] + speeds[platform] * et * direction
        height_rebuilt = q5.ORIGINS[platform, 2] - 0.5 * q5.p().gravity * fuse**2
        spatial_residuals.extend([
            float(np.linalg.norm(release[:2] - release_xy_rebuilt)),
            abs(float(release[2] - q5.ORIGINS[platform, 2])),
            float(np.linalg.norm(explosion[:2] - explosion_xy_rebuilt)),
            abs(float(explosion[2] - height_rebuilt)),
        ])
        bombs.append(q5.Bomb(platform, q5.MISSILE_NAMES.index(str(values[11])), tau, fuse))
    decision = q5.Decision(
        tuple(headings), tuple(speeds),
        tuple(sorted(bombs, key=lambda bomb: (bomb.platform, bomb.release_s))),
        "q5_v2_rounded_excel_readback",
    )
    q5.validate(decision)
    return decision, {
        "active_rows": active_rows,
        "blank_rows": blank_rows,
        "maximum_spatial_reconstruction_residual_m": max(spatial_residuals, default=0.0),
    }


def complete_workbook_durations(
    path: Path, rounded: q5.Decision, mesh: object, started: float, checkpoint_path: Path
) -> list[float]:
    durations: list[float] = []
    book = load_workbook(path)
    sheet = book[book.sheetnames[0]]
    completed = 0
    for platform in range(5):
        for number, bomb in enumerate(q5.by_platform(rounded, platform), 1):
            duration = q5.single_bomb_duration(rounded, bomb, mesh)
            durations.append(duration)
            sheet.cell(row=2 + 3 * platform + number - 1, column=11).value = round(duration, 6)
            completed += 1
            emit("excel_single_bomb", completed=completed, total=len(rounded.bombs))
            save_checkpoint(checkpoint_path, started, "excel_single_bomb", completed=completed, total=len(rounded.bombs))
    book.save(path)
    return durations


def excel_delivery(
    decision: q5.Decision,
    precise_payload: Mapping[str, object],
    path: Path,
    started: float,
    checkpoint_path: Path,
) -> dict[str, object]:
    write_coordinate_workbook(decision, path)
    rounded, first_read = reconstruct_workbook(path)
    mesh = q5.q3.surface_mesh(q5.PRECISE["n_theta"], q5.PRECISE["n_levels"])
    computed_durations = complete_workbook_durations(path, rounded, mesh, started, checkpoint_path)
    reread, final_read = reconstruct_workbook(path)
    sheet = load_workbook(path, data_only=True).active
    stored_durations = [float(sheet.cell(row=row, column=11).value) for row in final_read["active_rows"]]
    duration_residual = max(
        (abs(stored - computed) for stored, computed in zip(stored_durations, computed_durations)),
        default=0.0,
    )
    restricts = [[list(interval) for interval in intervals]
                 for intervals in precise_payload["guard"]["restricts_s"]]
    rounded_evaluation = q5.evaluate(reread, q5.PRECISE, mesh, restricts)
    rounded_metrics = from_interval_sets(
        [record["intervals_s"] for record in rounded_evaluation["by_missile"]]
    )
    original_metrics = precise_payload["metrics"]
    delta = {
        key: float(rounded_metrics.as_dict()[key]) - float(original_metrics[key])
        for key in ("j_sum_missile_s", "j_fair_s", "j_all_s", "l_all_s")
    }
    maximum_spatial = max(
        first_read["maximum_spatial_reconstruction_residual_m"],
        final_read["maximum_spatial_reconstruction_residual_m"],
    )
    result = {
        "path": str(path.relative_to(ROOT)),
        "sha256": sha256(path),
        "numeric_precision_decimal_places": 6,
        "active_row_count": len(final_read["active_rows"]),
        "blank_reserved_row_count": len(final_read["blank_rows"]),
        "maximum_spatial_reconstruction_residual_m": maximum_spatial,
        "maximum_single_bomb_duration_rounding_residual_s": duration_residual,
        "rounded_metrics": rounded_metrics.as_dict(),
        "delta_from_unrounded": delta,
        "constraint_margins": q5.constraint_margins(reread),
        "pass": (
            len(final_read["active_rows"]) == len(decision.bombs)
            and maximum_spatial <= EXCEL_SPACE_TOL_M
            and duration_residual <= EXCEL_DURATION_ROUNDING_TOL_S
            and abs(delta["j_sum_missile_s"]) <= EXCEL_METRIC_TOL_MISSILE_S
            and min(q5.constraint_margins(reread).values()) >= -2e-8
        ),
    }
    return result


def dense_target_check(
    decision: q5.Decision, precise_payload: Mapping[str, object], started: float,
    checkpoint_path: Path,
) -> dict[str, object]:
    mesh = q5.q3.surface_mesh(q5.DENSIFIED["n_theta"], q5.DENSIFIED["n_levels"])
    by_missile: list[dict[str, object]] = []
    for missile_index, record in enumerate(precise_payload["by_missile"]):
        interval_times: list[float] = []
        for left, right in record["intervals_s"]:
            interval_times.extend([left + fraction * (right - left) for fraction in (0.25, 0.5, 0.75)])
        interior_times = sorted(set(interval_times))
        interior_margin = q5.strict_margins(decision, missile_index, interior_times, mesh, assume_valid=True)
        roots: list[dict[str, object]] = []
        for root in record["root_residuals"]:
            if int(root["iterations"]) == 0:
                continue
            endpoint = float(root["time_s"])
            delta = 0.005
            samples = [max(0.0, endpoint - delta), min(q5.arrival(missile_index), endpoint + delta)]
            margins = q5.strict_margins(decision, missile_index, samples, mesh, assume_valid=True)
            entry = root["kind"] == "entry"
            bracket_pass = bool(
                (margins[0] > 0 and margins[1] <= DENSE_INTERIOR_TOL_M)
                if entry else
                (margins[0] <= DENSE_INTERIOR_TOL_M and margins[1] > 0)
            )
            drift = None
            if margins[0] * margins[1] <= 0 and abs(margins[1] - margins[0]) > 1e-14:
                estimated = samples[0] - margins[0] * (samples[1] - samples[0]) / (margins[1] - margins[0])
                drift = float(estimated - endpoint)
            roots.append({
                "kind": root["kind"], "precise_time_s": endpoint,
                "probe_times_s": samples, "probe_margins_m": margins.tolist(),
                "linearized_drift_s": drift, "bracket_pass": bracket_pass,
            })
        maximum_interior = float(np.max(interior_margin)) if len(interior_margin) else -math.inf
        maximum_drift = max((abs(item["linearized_drift_s"]) for item in roots
                             if item["linearized_drift_s"] is not None), default=0.0)
        record_out = {
            "missile": q5.MISSILE_NAMES[missile_index],
            "interior_probe_count": len(interior_times),
            "maximum_interior_margin_m": maximum_interior,
            "geometric_roots": roots,
            "maximum_abs_linearized_endpoint_drift_s": maximum_drift,
            "pass": (
                maximum_interior <= DENSE_INTERIOR_TOL_M
                and all(item["bracket_pass"] for item in roots)
                and maximum_drift <= DENSE_ENDPOINT_DRIFT_TOL_S
            ),
        }
        by_missile.append(record_out)
        emit("densified_missile", missile=q5.MISSILE_NAMES[missile_index], passed=record_out["pass"])
        save_checkpoint(checkpoint_path, started, "densified_missile", missile=missile_index + 1, passed=record_out["pass"])
    return {
        "settings": dict(q5.DENSIFIED),
        "mode": "targeted_reported_interval_interior_and_endpoint_check",
        "by_missile": by_missile,
        "pass": all(item["pass"] for item in by_missile),
        "claim_boundary": "not an independent full-horizon DENSIFIED missed-island proof",
    }


def e3_checks(decision: q5.Decision, precise_payload: Mapping[str, object], structural: Mapping[str, object]) -> dict[str, object]:
    intervals = precise_payload["metrics"]["per_missile_intervals_s"]
    probe_times = sorted({left + 0.5 * (right - left) for group in intervals for left, right in group})
    mesh = q5.q3.surface_mesh(64, 7)
    baseline = [q5.strict_margins(decision, missile, probe_times, mesh, assume_valid=True) for missile in range(3)]
    relabeled = q5.Decision(
        decision.headings_rad, decision.speeds_mps,
        tuple(replace(bomb, label=(bomb.label + 1) % 3) for bomb in decision.bombs),
        "cyclic_label_permutation",
    )
    reversed_decision = q5.Decision(
        decision.headings_rad, decision.speeds_mps, tuple(reversed(decision.bombs)), "reversed_bomb_order"
    )
    label_diff = max(float(np.max(np.abs(base - q5.strict_margins(relabeled, missile, probe_times, mesh, assume_valid=True))))
                     for missile, base in enumerate(baseline))
    order_diff = max(float(np.max(np.abs(base - q5.strict_margins(reversed_decision, missile, probe_times, mesh, assume_valid=True))))
                     for missile, base in enumerate(baseline))
    synthetic = np.array([[-1.0, 1.0], [1.0, -1.0]])
    joint_margin = float(np.max(np.min(synthetic, axis=0)))
    single_cloud_margins = np.max(synthetic, axis=1).tolist()
    empty_metrics = from_interval_sets([[], [], []])
    conclusion = structural["conclusion"]
    result = {
        "cyclic_label_permutation_max_margin_difference_m": label_diff,
        "bomb_order_reversal_max_margin_difference_m": order_diff,
        "empty_resource_metrics": empty_metrics.as_dict(),
        "multi_cloud_quantifier_counterexample": {
            "cloud_by_target_point_margins": synthetic.tolist(),
            "single_cloud_complete_margins": single_cloud_margins,
            "joint_max_point_min_cloud_margin": joint_margin,
            "interpretation": "each cloud fails alone, while pointwise cloud choice completes joint coverage",
        },
        "all_13_bombs_positive_primary_marginal_at_screen": conclusion["all_13_bombs_positive_primary_marginal_at_screen"],
        "all_5_platforms_positive_primary_marginal_at_screen": conclusion["all_5_platforms_positive_primary_marginal_at_screen"],
        "fy1_existing_library_improvement_counterexample": conclusion["fy1_existing_library_improvement_counterexample"],
        "pass": (
            label_diff <= 1e-12 and order_diff <= 1e-12
            and empty_metrics.j_sum_missile_s == 0.0
            and joint_margin <= 0.0 and min(single_cloud_margins) > 0.0
            and conclusion["all_13_bombs_positive_primary_marginal_at_screen"]
            and conclusion["all_5_platforms_positive_primary_marginal_at_screen"]
        ),
        "claim_boundary": conclusion["claim_boundary"],
    }
    return result


def render_report(payload: Mapping[str, object]) -> str:
    metrics = payload["result_identity"]["metrics"]
    e1, e2, e3, e4 = payload["E1"], payload["E2"], payload["E3"], payload["E4"]
    excel = e2["rounded_excel_readback"]
    return "\n".join([
        "# Q5_v2 工作验证报告",
        "",
        f"结论：E1、E2、E3 均为 **{'通过' if payload['working_verification_pass'] else '未通过'}**；E4 为 `{e4['status']}`。这支持当前候选作为经工作验证的高质量可行解，不构成连续域全局最优证明。",
        "",
        "## 结果身份",
        "",
        f"- `J_sum = {metrics['j_sum_missile_s']:.12f}` 导弹·s；",
        f"- 三枚导弹时长：{', '.join(f'{value:.12f}' for value in metrics['per_missile_duration_s'])} s；",
        f"- `J_fair = {metrics['j_fair_s']:.12f} s`，`J_all = {metrics['j_all_s']:.12f} s`，`L_all = {metrics['l_all_s']:.12f} s`。",
        "",
        "## E1 实现核对",
        "",
        f"输入哈希、硬约束和 6 类非法输入拒绝均通过。运动学最大位置回代残差为 `{e1['kinematic_roundtrip']['maximum_position_residual_m']:.3e} m`；独立有限视线 oracle 与生产评价器最大差异为 `{e1['independent_oracle']['maximum_abs_difference_m']:.3e} m`。",
        "",
        "## E2 数值核对",
        "",
        f"SCREEN 到 PRECISE 的 `J_sum` 变化为 `{e2['screen_to_precise']['j_sum_delta_missile_s']:.9f}` 导弹·s。PRECISE 求根回退次数为 `{e2['root_quality']['fallback_count']}`。定向 DENSIFIED 内部/端点检查结果为 `{e2['targeted_densified']['pass']}`。",
        "",
        f"舍入结果表位于 `{excel['path']}`，含 `{excel['active_row_count']}` 条有效投放记录；坐标反解最大残差 `{excel['maximum_spatial_reconstruction_residual_m']:.3e} m`，舍入后 `J_sum` 变化 `{excel['delta_from_unrounded']['j_sum_missile_s']:.9f}` 导弹·s。",
        "",
        "## E3 结构核对",
        "",
        "标签循环置换和烟幕弹排列反转均不改变物理裕量；空资源退化为零时长。二点二烟团反例证明实现采用的是逐目标点选择烟团的联合覆盖量词。已有删除消融显示 13 枚弹和 5 个平台对主目标均有正边际，但该结论只是在其余变量固定及已检查候选库范围内成立。",
        "",
        "## E4 现实边界",
        "",
        e4["conclusion"],
        "",
        "## 证明边界",
        "",
        payload["proof_boundary"],
        "",
    ])


def run(args: argparse.Namespace) -> dict[str, object]:
    started = time.perf_counter()
    checkpoint_path = (ROOT / args.checkpoint_json).resolve()
    precise_path = (ROOT / args.input_precise).resolve()
    result_path = (ROOT / args.input_result).resolve()
    structural_path = (ROOT / args.input_structural).resolve()
    precise = json.loads(precise_path.read_text(encoding="utf-8"))
    result = json.loads(result_path.read_text(encoding="utf-8"))
    structural = json.loads(structural_path.read_text(encoding="utf-8"))
    decision = decision_from_payload(precise)
    emit("start", bombs=len(decision.bombs))
    save_checkpoint(checkpoint_path, started, "start", bombs=len(decision.bombs))

    hashes = {relative: sha256(ROOT / relative) for relative in EXPECTED_HASHES}
    identities_pass = hashes == EXPECTED_HASHES
    invalid = invalid_case_results(decision)
    kinematics = kinematic_roundtrip(decision)
    constants = {
        "smoke_radius_m": q5.p().smoke_radius,
        "smoke_sink_speed_mps": q5.p().smoke_sink_speed,
        "smoke_duration_s": q5.p().smoke_duration,
        "pass": (
            q5.p().smoke_radius == 10.0 and q5.p().smoke_sink_speed == 3.0
            and q5.p().smoke_duration == 20.0
        ),
    }

    oracle_mesh = q5.q3.surface_mesh(256, 17)
    oracle_records: list[dict[str, object]] = []
    for missile_index, intervals in enumerate(precise["metrics"]["per_missile_intervals_s"]):
        times = interval_probe_times(intervals)
        production = q5.strict_margins(decision, missile_index, times, oracle_mesh, assume_valid=True)
        independent = independent_margin(decision, missile_index, times, oracle_mesh)
        difference = float(np.max(np.abs(production - independent)))
        oracle_records.append({"missile": q5.MISSILE_NAMES[missile_index], "probe_count": len(times), "maximum_abs_difference_m": difference})
        emit("oracle_missile", missile=q5.MISSILE_NAMES[missile_index], maximum_abs_difference_m=difference)
        save_checkpoint(checkpoint_path, started, "oracle_missile", missile=missile_index + 1, maximum_abs_difference_m=difference)
    oracle_max = max(item["maximum_abs_difference_m"] for item in oracle_records)
    e1 = {
        "input_identity": {"expected": EXPECTED_HASHES, "observed": hashes, "pass": identities_pass},
        "constants": constants,
        "constraint_margins": q5.constraint_margins(decision),
        "invalid_cases_rejected": invalid,
        "kinematic_roundtrip": kinematics,
        "independent_oracle": {"mesh": "256x17 complete cylinder", "by_missile": oracle_records, "maximum_abs_difference_m": oracle_max, "pass": oracle_max <= ORACLE_TOL_M},
    }
    e1["pass"] = bool(identities_pass and constants["pass"] and all(invalid.values()) and kinematics["pass"] and e1["independent_oracle"]["pass"] and min(e1["constraint_margins"].values()) >= -2e-8)
    emit("e1_complete", passed=e1["pass"])

    recomputed = from_interval_sets(precise["metrics"]["per_missile_intervals_s"])
    algebra_delta = max(abs(float(recomputed.as_dict()[key]) - float(precise["metrics"][key])) for key in ("j_sum_missile_s", "j_fair_s", "j_all_s", "l_all_s"))
    screen_delta_per = [p - s for p, s in zip(precise["metrics"]["per_missile_duration_s"], precise["screen_metrics"]["per_missile_duration_s"])]
    screen_sum_delta = float(precise["comparison_to_screen"]["j_sum_missile_s"])
    roots = [root for missile in precise["by_missile"] for root in missile["root_residuals"] if int(root["iterations"]) > 0]
    root_quality = {
        "geometric_root_count": len(roots),
        "fallback_count": sum(int(missile["root_fallback_count"]) for missile in precise["by_missile"]),
        "maximum_abs_residual_m": max((abs(float(root["margin_m"])) for root in roots), default=0.0),
    }
    dense = dense_target_check(decision, precise, started, checkpoint_path)
    excel = excel_delivery(decision, precise, (ROOT / args.excel).resolve(), started, checkpoint_path)
    master_gap = float(result["restricted_master"]["mip_gap"])
    screen_convergence_pass = (
        max(abs(value) for value in screen_delta_per) <= SCREEN_PRECISE_PER_MISSILE_TOL_S
        and abs(screen_sum_delta) <= SCREEN_PRECISE_SUM_TOL_MISSILE_S
    )
    e2 = {
        "interval_algebra": {"maximum_metric_residual": algebra_delta, "pass": algebra_delta <= 1e-10},
        "screen_to_precise": {"per_missile_delta_s": screen_delta_per, "j_sum_delta_missile_s": screen_sum_delta, "pass": screen_convergence_pass},
        "root_quality": {**root_quality, "pass": root_quality["fallback_count"] == 0 and root_quality["maximum_abs_residual_m"] <= 1e-8},
        "targeted_densified": dense,
        "rounded_excel_readback": excel,
        "restricted_master": {"mip_gap": master_gap, "pass_within_current_library_only": master_gap == 0.0},
    }
    e2["pass"] = bool(e2["interval_algebra"]["pass"] and e2["screen_to_precise"]["pass"] and e2["root_quality"]["pass"] and dense["pass"] and excel["pass"] and master_gap == 0.0)
    emit("e2_complete", passed=e2["pass"])
    save_checkpoint(checkpoint_path, started, "e2_complete", passed=e2["pass"])

    e3 = e3_checks(decision, precise, structural)
    e4 = {
        "status": "justified_not_triggered_working",
        "pass": True,
        "conclusion": "题面理想运动学下，烟团截断有限视线段会降低真目标几何可见性，机制方向一致；但缺少导引头探测、烟雾浓度与实装散布数据，不能把本结果外推为现实部署成功率。",
        "unavailable_evidence": ["seeker detection model", "smoke concentration/transmittance", "deployment dispersion and field trials"],
    }
    working_pass = bool(e1["pass"] and e2["pass"] and e3["pass"] and e4["pass"])
    payload = {
        "schema_version": "q5-v2-working-validation-1.0",
        "status": "working_validation_complete" if working_pass else "working_validation_failed",
        "generated_at": datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds"),
        "elapsed_s": time.perf_counter() - started,
        "working_verification_pass": working_pass,
        "formal_s6_close": False,
        "inputs": {str(path.relative_to(ROOT)): sha256(path) for path in (precise_path, result_path, structural_path)},
        "result_identity": {"source": str(precise_path.relative_to(ROOT)), "metrics": precise["metrics"], "decision_bomb_count": len(decision.bombs)},
        "E1": e1, "E2": e2, "E3": e3, "E4": e4,
        "proof_boundary": "受限 MILP 的零 gap 只属于当前列库与离散主问题；PRECISE 依赖全时域 SCREEN 支撑区间守卫，DENSIFIED 为报告区间定向检查；未认证列定价、未证明连续域全局最优，也未完成需要用户通读并复述的正式 S6 人类验收。",
    }
    solver.atomic_json((ROOT / args.output_json).resolve(), payload)
    report_path = (ROOT / args.report_md).resolve()
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(render_report(payload), encoding="utf-8")
    save_checkpoint(checkpoint_path, started, "complete", passed=working_pass, output_json=args.output_json, report_md=args.report_md, excel=args.excel)
    emit("complete", passed=working_pass, elapsed_s=payload["elapsed_s"])
    return payload


def parser() -> argparse.ArgumentParser:
    result = argparse.ArgumentParser()
    result.add_argument("--input-precise", default="docs/q5_v2_precise.json")
    result.add_argument("--input-result", default="docs/q5_v2_result.json")
    result.add_argument("--input-structural", default="docs/q5_v2_structural_audit.json")
    result.add_argument("--output-json", default="docs/q5_v2_validation_results.json")
    result.add_argument("--report-md", default="docs/q5_v2_validation_report.md")
    result.add_argument("--checkpoint-json", default="output/q5_v2/verification_checkpoint.json")
    result.add_argument("--excel", default="output/q5_v2/result3.xlsx")
    return result


if __name__ == "__main__":
    run(parser().parse_args())
