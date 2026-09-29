"""Q5 strict shared-track solver for the longest continuous joint-obscuration window."""
from __future__ import annotations

import argparse
import concurrent.futures
import hashlib
import itertools
import json
import math
import os
import platform
import shutil
import sys
import tempfile
import time
from dataclasses import asdict, dataclass, replace
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
from src.q3 import selection_pilot as q3  # noqa: E402
from src.q5 import geometry_pilot as geometry  # noqa: E402
from src.q5 import selection_pilot as pilot  # noqa: E402

SPEC_ID = "SPEC-Q5-WORKING-4.0"
RESULT_ID = "Q5-SPEC-Q5-WORKING-4.0-LONGEST-CONTINUOUS-WORKING-004"
NAMES = pilot.UN
MISSILE_NAMES = pilot.MN
ORIGINS = pilot.ORIGINS
TIME_TOL = 1e-9
HASHES = {
    "data/A题.pdf": "a37f6aad30b16ea09da9cb320ce194b12d3a28874008cf9ff19549f6c6079447",
    "data/附件/result3.xlsx": "b648c82d63e459ba6e6b3711ae79875e373521cd543b45571c4d8ff1ad5ec54a",
    "planning/35_q5_longest_continuous_candidate_set.md": "c823deb5459793cac41ef6e70795b430f22bcbbd25a1cd51e78562ecb85d6eb7",
    "planning/36_q5_longest_continuous_model_spec.md": "0fc18833b14a62ad298a84a43a1d9860eb4919603ee19c910ac4a1c6ccaf3962",
    "docs/q5_intersection_seed.json": "70e25a837f41c92057aa19cfd8090b9d0ed57527370a5f861d82a875a3157324",
    "src/q1/model.py": "cb36870ca55ed3b1e3c3adfe02cd91667c89249fa1aeb730cb42f0a1bbd4f7f1",
    "src/q3/selection_pilot.py": "33f67bee3a033ffebc48dd7a6275846c991977f6035554b75cd055de80430f64",
    "src/q5/geometry_pilot.py": "d6a639bba67b076f6338a6e03539971c88706d67e4458d8bd00b6e255609f878",
}

FAST = {"name": "fast_strict", "scan_step_s": .05, "n_theta": 64, "n_levels": 7,
        "root_time_tolerance_s": 1e-6, "margin_tolerance_m": 1e-6}
SCREEN = {"name": "screen_strict", "scan_step_s": .02, "n_theta": 128, "n_levels": 9,
          "root_time_tolerance_s": 2e-6, "margin_tolerance_m": 2e-6}
PRECISE = {"name": "precise_strict", "scan_step_s": .005, "n_theta": 512, "n_levels": 33,
           "root_time_tolerance_s": 1e-8, "margin_tolerance_m": 1e-8}
DENSIFIED = {"name": "densified_strict", "scan_step_s": .0025, "n_theta": 1024, "n_levels": 65,
             "root_time_tolerance_s": 5e-9, "margin_tolerance_m": 5e-9}
PROBE = {"name": "refine_probe", "scan_step_s": .10, "n_theta": 32, "n_levels": 5,
         "root_time_tolerance_s": 5e-6, "margin_tolerance_m": 5e-6}
RANK_TOLERANCE = {"fast_strict": 1e-3, "screen_strict": 2e-4,
                  "precise_strict": 3e-5, "densified_strict": 3e-5,
                  "refine_probe": 2e-3}
LEGACY_LONGEST_CONTINUOUS_S = 7.060098194887
LEGACY_TOTAL_INTERSECTION_S = 9.733419947347
SEED_PATH = ROOT / "docs/q5_intersection_seed.json"

GEOMETRY_BUDGET = {
    "main_time_step_s": .25,
    "offset_time_s": .125,
    "local_time_step_s": .05,
    "local_half_width_s": .25,
    "mesh_theta": 48,
    "mesh_levels": 7,
    "c3_track_limit": 48,
    "c2_guard_track_limit": 12,
    "event_limit": 18,
    "plan_limit": 96,
    "beam_widths": (256, 1024, 3000, 6000, 10000),
    "complete_limit": 240,
    "guard_fraction": .20,
    "parallel_workers": 4,
}

REFINE_BUDGET = {
    "wall_time_s_per_start": 90.,
    "probe_evaluations_per_start": 180,
    "fast_evaluations_per_start": 70,
    "progress_every_probe_evaluations": 20,
}

PARAMETERS = q3.base_parameters()
ARRIVAL_TIMES = tuple(float(np.linalg.norm(geometry.MISSILES[j]) / PARAMETERS.missile_speed)
                      for j in range(3))


class Q5Error(RuntimeError):
    def __init__(self, code, message):
        super().__init__(message)
        self.code = code


@dataclass(frozen=True)
class Bomb:
    platform: int
    label: int
    release_s: float
    fuse_s: float


@dataclass(frozen=True)
class Decision:
    headings_rad: tuple[float, float, float, float, float]
    speeds_mps: tuple[float, float, float, float, float]
    bombs: tuple[Bomb, ...]
    source: str


def sha(path):
    h = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for block in iter(lambda: handle.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def ready(value):
    if isinstance(value, np.ndarray):
        return value.tolist()
    if isinstance(value, np.generic):
        return value.item()
    if isinstance(value, dict):
        return {k: ready(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [ready(v) for v in value]
    return value


def p():
    return PARAMETERS


def arrival(j):
    return ARRIVAL_TIMES[j]


def verify_identities():
    got = {key: sha(ROOT / key) for key in HASHES}
    bad = {key: (HASHES[key], got[key]) for key in HASHES if HASHES[key] != got[key]}
    if bad:
        raise Q5Error("Q5_IDENTITY_MISMATCH", str(bad))
    return got


def by_platform(x, i):
    return tuple(sorted((b for b in x.bombs if b.platform == i), key=lambda b: b.release_s))


def explosion_time(b):
    return b.release_s + b.fuse_s


def direction(x, i):
    return np.array([math.cos(x.headings_rad[i]), math.sin(x.headings_rad[i]), 0.])


def release_point(x, b):
    return ORIGINS[b.platform] + x.speeds_mps[b.platform] * b.release_s * direction(x, b.platform)


def explosion_point(x, b):
    point = ORIGINS[b.platform] + x.speeds_mps[b.platform] * explosion_time(b) * direction(x, b.platform)
    point[2] = ORIGINS[b.platform, 2] - .5 * p().gravity * b.fuse_s ** 2
    return point


def cloud(x, b, times):
    points = np.repeat(explosion_point(x, b).reshape(1, 3), len(times), axis=0)
    points[:, 2] -= p().smoke_sink_speed * (times - explosion_time(b))
    return points


def constraint_margins(x):
    gaps = [v.release_s - u.release_s for i in range(5) for u, v in zip(by_platform(x, i)[:-1], by_platform(x, i)[1:])]
    heights = [explosion_point(x, b)[2] for b in x.bombs]
    arrivals = [max(arrival(j) for j in range(3)) - explosion_time(b) for b in x.bombs]
    return {
        "speed_low_mps": float(min(np.asarray(x.speeds_mps) - 70.)),
        "speed_high_mps": float(min(140. - np.asarray(x.speeds_mps))),
        "release_time_s": float(min((b.release_s for b in x.bombs), default=math.inf)),
        "fuse_delay_s": float(min((b.fuse_s for b in x.bombs), default=math.inf)),
        "explosion_height_m": float(min(heights, default=math.inf)),
        "one_second_gap_s": float(min(gaps, default=math.inf) - 1.),
        "latest_arrival_s": float(min(arrivals, default=math.inf)),
        "capacity_bombs": float(min(3 - len(by_platform(x, i)) for i in range(5))),
    }


def validate(x):
    numbers = [*x.headings_rad, *x.speeds_mps, *[v for b in x.bombs for v in (b.release_s, b.fuse_s)]]
    if not numbers or not np.all(np.isfinite(numbers)):
        raise Q5Error("Q5_CONSTRAINT_FAILURE", "nonfinite or empty decision")
    if any(not 0 <= a < 2 * math.pi for a in x.headings_rad):
        raise Q5Error("Q5_CONSTRAINT_FAILURE", "heading")
    if len({(b.platform, round(b.release_s, 10), round(b.fuse_s, 10), b.label) for b in x.bombs}) != len(x.bombs):
        raise Q5Error("Q5_CONSTRAINT_FAILURE", "duplicate bomb")
    if any(not 0 <= b.platform < 5 or not 0 <= b.label < 3 for b in x.bombs):
        raise Q5Error("Q5_CONSTRAINT_FAILURE", "indices")
    margins = constraint_margins(x)
    if min(margins.values()) < -2e-8:
        raise Q5Error("Q5_CONSTRAINT_FAILURE", str(margins))


def decision_record(x):
    validate(x)
    platforms = []
    for i in range(5):
        rows = []
        for number, b in enumerate(by_platform(x, i), 1):
            rows.append({**asdict(b), "bomb_number": number, "label": MISSILE_NAMES[b.label],
                         "explosion_time_s": explosion_time(b), "release_point_m": release_point(x, b).tolist(),
                         "explosion_point_m": explosion_point(x, b).tolist()})
        platforms.append({"platform": NAMES[i], "heading_rad": x.headings_rad[i],
                          "heading_deg": math.degrees(x.headings_rad[i]) % 360,
                          "speed_mps": x.speeds_mps[i], "bombs": rows})
    return {"source": x.source, "active_bomb_count": len(x.bombs), "platforms": platforms,
            "constraint_margins": constraint_margins(x)}


def strict_margins(x, j, times, mesh, assume_valid=False):
    if not assume_valid:
        validate(x)
    times = np.asarray(times, float)
    out = np.empty(len(times))
    batch = max(2, min(96, int(1_600_000 / max(1, len(mesh.points)))))
    for left in range(0, len(times), batch):
        right = min(len(times), left + batch)
        block = times[left:right]
        missiles = pilot.missile(j, block)
        visible = q3.visibility_mask(missiles, mesh)
        minimum = np.full((len(block), len(mesh.points)), np.inf)
        for b in x.bombs:
            et = explosion_time(b)
            active = (block >= et - TIME_TOL) & (block <= min(et + p().smoke_duration, arrival(j)) + TIME_TOL)
            if np.any(active):
                dist = q3._distance_to_segments(cloud(x, b, block), missiles, mesh)
                dist[~active] = np.inf
                minimum = np.minimum(minimum, dist)
        values = minimum - p().smoke_radius
        values[~visible] = -np.inf
        margin = np.max(values, axis=1)
        margin[~np.any(visible, axis=1)] = np.inf
        out[left:right] = margin
    if np.any(np.isnan(out)):
        raise Q5Error("Q5_SURFACE_EVALUATOR_FAILURE", "NaN")
    return out


def time_grid(x, j, step, restrict=None):
    windows = q3.merge_intervals([[explosion_time(b), min(explosion_time(b) + p().smoke_duration, arrival(j))]
                                  for b in x.bombs if explosion_time(b) <= arrival(j)])
    if restrict is not None:
        windows = q3.merge_intervals([[max(a, c), min(b, d)] for a, b in windows for c, d in restrict
                                      if min(b, d) >= max(a, c)])
    boundaries = sorted({v for b in x.bombs for v in (explosion_time(b), min(explosion_time(b) + p().smoke_duration, arrival(j)))})
    values = []
    for a, b in windows:
        cuts = sorted({a, b, *[v for v in boundaries if a < v < b]})
        for left, right in zip(cuts[:-1], cuts[1:]):
            values.extend(np.linspace(left, right, max(1, math.ceil((right - left) / step)) + 1))
    return np.array(sorted(set(float(v) for v in values))), np.asarray(boundaries)


def bisect_root(x, j, left, right, mesh, settings):
    fl = float(strict_margins(x, j, [left], mesh, assume_valid=True)[0])
    fr = float(strict_margins(x, j, [right], mesh, assume_valid=True)[0])
    if not math.isfinite(fl) or not math.isfinite(fr) or fl * fr > 0:
        return (left, fl, 0, True) if fl <= 0 else (right, fr, 0, True)
    mid, fm = (left + right) / 2, math.nan
    for iteration in range(1, 101):
        mid = (left + right) / 2
        fm = float(strict_margins(x, j, [mid], mesh, assume_valid=True)[0])
        if right - left <= settings["root_time_tolerance_s"] and abs(fm) <= settings["margin_tolerance_m"]:
            return mid, fm, iteration, False
        if fl * fm <= 0:
            right, fr = mid, fm
        else:
            left, fl = mid, fm
    return mid, fm, 100, True


def solve_intervals(x, j, settings, mesh=None, restrict=None):
    validate(x)
    mesh = mesh or q3.surface_mesh(settings["n_theta"], settings["n_levels"])
    times, boundaries = time_grid(x, j, settings["scan_step_s"], restrict)
    if not len(times):
        return {"missile": MISSILE_NAMES[j], "settings": dict(settings), "surface_point_count": len(mesh.points),
                "scan_count": 0, "intervals_s": [], "duration_s": 0., "minimum_scan_margin_m": math.inf,
                "root_residuals": [], "root_fallback_count": 0, "recovered_unsampled_interval_count": 0}
    margins = strict_margins(x, j, times, mesh, assume_valid=True)
    inside = margins <= 0
    intervals, residuals = [], []
    fallbacks, i = 0, 0
    gap = settings["scan_step_s"] * (1 + 1e-6)
    while i < len(times):
        if not inside[i]:
            i += 1
            continue
        first = i
        while i + 1 < len(times) and inside[i + 1] and times[i + 1] - times[i] <= gap:
            i += 1
        last = i
        if first == 0 or times[first] - times[first - 1] > gap or np.any(np.isclose(times[first], boundaries, atol=TIME_TOL, rtol=0)):
            entry = (float(times[first]), float(margins[first]), 0, False)
        else:
            entry = bisect_root(x, j, float(times[first - 1]), float(times[first]), mesh, settings)
        if last == len(times) - 1 or times[last + 1] - times[last] > gap or np.any(np.isclose(times[last], boundaries, atol=TIME_TOL, rtol=0)):
            exit_ = (float(times[last]), float(margins[last]), 0, False)
        else:
            exit_ = bisect_root(x, j, float(times[last]), float(times[last + 1]), mesh, settings)
        fallbacks += int(entry[3]) + int(exit_[3])
        intervals.append([entry[0], exit_[0]])
        residuals.extend([{"kind": "entry", "time_s": entry[0], "margin_m": entry[1], "iterations": entry[2], "fallback": entry[3]},
                          {"kind": "exit", "time_s": exit_[0], "margin_m": exit_[1], "iterations": exit_[2], "fallback": exit_[3]}])
        i += 1
    evaluator = lambda t: float(strict_margins(x, j, [float(t)], mesh, assume_valid=True)[0])
    try:
        recovered = q3.q1.recover_unsampled_intervals(times, margins, evaluator,
                                                       time_tolerance=settings["root_time_tolerance_s"],
                                                       margin_tolerance=settings["margin_tolerance_m"], maximum_gap=gap)
    except q3.q1.Q1Error as exc:
        raise Q5Error("Q5_TIME_EVENT_FAILURE", str(exc)) from exc
    intervals.extend(recovered["recovered_intervals_s"])
    merged = q3.merge_intervals(intervals)
    return {"missile": MISSILE_NAMES[j], "settings": dict(settings), "surface_point_count": len(mesh.points),
            "scan_count": len(times), "intervals_s": merged, "duration_s": q3.interval_measure(merged),
            "minimum_scan_margin_m": float(np.min(margins)), "root_residuals": residuals,
            "root_fallback_count": fallbacks, "local_minimum_checks": recovered["checked_local_minima"],
            "recovered_unsampled_interval_count": len(recovered["recovered_intervals_s"])}


def longest_component(intervals):
    if not intervals:
        return None, 0.
    interval = max(intervals, key=lambda item: item[1] - item[0])
    return [float(interval[0]), float(interval[1])], float(interval[1] - interval[0])


def evaluate(x, settings, mesh=None, restricts=None):
    mesh = mesh or q3.surface_mesh(settings["n_theta"], settings["n_levels"])
    records = [solve_intervals(x, j, settings, mesh, None if restricts is None else restricts[j]) for j in range(3)]
    durations = [r["duration_s"] for r in records]
    common = pilot.intersect_intervals([r["intervals_s"] for r in records])
    total = q3.interval_measure(common)
    longest, longest_duration = longest_component(common)
    if total > min(durations) + 2e-7:
        raise Q5Error("Q5_INTERVAL_INTERSECTION_FAILURE", f"intersection={total}, durations={durations}")
    return {"by_missile": records, "durations_s": durations,
            "intersection_intervals_s": common,
            "longest_interval_s": longest,
            "longest_continuous_s": longest_duration,
            "total_intersection_s": total,
            "objective_s": longest_duration,
            "sum_duration_missile_s": float(sum(durations)),
            "minimum_missile_duration_s": float(min(durations))}


def ranked(records, result_field, tolerance):
    del tolerance  # Numerical tolerances never authorize a shorter primary window.
    def key(row):
        result = row[result_field]
        return (result["longest_continuous_s"], result["total_intersection_s"],
                result["sum_duration_missile_s"], result["minimum_missile_duration_s"])
    return sorted(records, key=key, reverse=True)


def better_result(candidate, current, numerical_tolerance=1e-9):
    """Strict lexicographic acceptance without cumulative primary loss."""
    if candidate["longest_continuous_s"] > current["longest_continuous_s"] + numerical_tolerance:
        return True
    if candidate["longest_continuous_s"] < current["longest_continuous_s"]:
        return False
    for field in ("total_intersection_s", "sum_duration_missile_s", "minimum_missile_duration_s"):
        if candidate[field] > current[field] + numerical_tolerance:
            return True
        if candidate[field] < current[field] - numerical_tolerance:
            return False
    return False


def from_geometry_choice(choice, source):
    headings = [0.] * 5
    speeds = [70.] * 5
    bombs = []
    for plan in choice:
        i = plan.platform
        headings[i], speeds[i] = plan.heading % (2 * math.pi), plan.speed
        bombs.extend(Bomb(i, event.label, event.release, event.fuse) for event in plan.events)
    x = Decision(tuple(headings), tuple(speeds),
                 tuple(sorted(bombs, key=lambda b: (b.platform, b.release_s))), source)
    validate(x)
    return x


def registered_incumbent():
    data = json.loads(SEED_PATH.read_text(encoding="utf-8"))
    if data.get("source_result_path") == "docs/q5_result.json" and data.get("source_spec_id") == SPEC_ID:
        raise Q5Error("Q5_SELF_REFERENTIAL_INCUMBENT", "seed points to the current output specification")
    decision = data["decision"]
    headings = [float(value) for value in decision["headings_rad"]]
    speeds = [float(value) for value in decision["speeds_mps"]]
    bombs = [Bomb(int(row["platform"]), int(row["label"]), float(row["release_s"]), float(row["fuse_s"]))
             for row in decision["bombs"]]
    x = Decision(tuple(headings), tuple(speeds), tuple(sorted(bombs, key=lambda b: (b.platform, b.release_s))),
                 "historical_intersection_seed_rescored")
    validate(x)
    return x, sha(SEED_PATH)


def stratified_surface_indices(mesh, n_azimuth=8, n_second=4):
    indices = []
    targets = np.linspace(0., 2 * math.pi, n_azimuth, endpoint=False)
    for kind in (0, 1):
        pool = np.flatnonzero(mesh.kinds == kind)
        if not len(pool):
            continue
        second_values = np.unique(mesh.second[pool])
        chosen_second = second_values[np.linspace(0, len(second_values) - 1,
                                                  min(n_second, len(second_values)), dtype=int)]
        for second in chosen_second:
            layer = pool[np.isclose(mesh.second[pool], second, atol=1e-12, rtol=0)]
            for angle in targets:
                angular = np.abs(np.angle(np.exp(1j * (mesh.theta[layer] - angle))))
                indices.append(int(layer[int(np.argmin(angular))]))
    top_centres = np.flatnonzero((mesh.kinds == 1) & np.isclose(mesh.second, 0., atol=1e-12, rtol=0))
    if len(top_centres):
        indices.append(int(top_centres[0]))
    return sorted(set(indices))


def inverse_track_cached(i, t, age, cloud):
    te = t - age
    if te <= 0.:
        return None
    ez = cloud[2] + p().smoke_sink_speed * age
    low = max(0., ORIGINS[i, 2] - .5 * p().gravity * te * te)
    if not low - TIME_TOL <= ez <= ORIGINS[i, 2] + TIME_TOL:
        return None
    fuse = math.sqrt(max(0., 2. * (ORIGINS[i, 2] - ez) / p().gravity))
    release = te - fuse
    horizontal = cloud[:2] - ORIGINS[i, :2]
    speed = float(np.hypot(horizontal[0], horizontal[1]) / te)
    if release < -TIME_TOL or not 70. - TIME_TOL <= speed <= 140. + TIME_TOL:
        return None
    heading = float(math.atan2(horizontal[1], horizontal[0]) % (2 * math.pi))
    return heading, min(140., max(70., speed)), max(0., release), fuse


def c3_tracks(i, sl):
    raw = {}
    proposed = feasible = 0
    representatives = [(index, sl.mesh.points[index])
                       for index in stratified_surface_indices(sl.mesh, 8, 4)]
    for j in range(3):
        m = sl.missiles[j]
        for surface_index, point in representatives:
            if not sl.visible[j, surface_index]:
                continue
            vector = point - m
            normal_h = np.array([-vector[1], vector[0]])
            norm_h = np.linalg.norm(normal_h)
            normal_h = normal_h / norm_h if norm_h > 0 else np.zeros(2)
            for fraction, age, transverse in itertools.product(
                    geometry.FRACTIONS, geometry.AGES, geometry.TRANSVERSE):
                if age >= sl.t:
                    continue
                proposed += 1
                cloud = m + fraction * vector
                cloud[:2] += transverse * normal_h
                solved = inverse_track_cached(i, sl.t, age, cloud)
                if solved is None:
                    continue
                feasible += 1
                heading, speed, _, _ = solved
                key = (round(heading / .004), round(speed / .5))
                raw.setdefault(key, geometry.Track(i, heading, speed, "C3_surface_tube", j, surface_index))
    ordered = sorted(raw.values(), key=lambda track: (
        track.source_label, track.source_surface, track.heading, track.speed))
    selected = []
    used_keys = set()
    used_strata = set()
    for track in ordered:
        stratum = (track.source_label, track.source_surface % 6)
        key = (round(track.heading, 12), round(track.speed, 9))
        if stratum not in used_strata and key not in used_keys:
            selected.append(track)
            used_strata.add(stratum)
            used_keys.add(key)
    for track in ordered:
        key = (round(track.heading, 12), round(track.speed, 9))
        if key not in used_keys:
            selected.append(track)
            used_keys.add(key)
        if len(selected) >= GEOMETRY_BUDGET["c3_track_limit"]:
            break
    selected = selected[:GEOMETRY_BUDGET["c3_track_limit"]]
    return selected, {"proposed": proposed, "kinematically_feasible": feasible,
                      "unique_tracks": len(raw), "selected_tracks": len(selected)}


def c2_guard_tracks(i, sl, occupied):
    surface_ids = stratified_surface_indices(sl.mesh, 5, 2)
    phase = (i * 7 + round(sl.t * 8)) % 48
    headings = (2 * math.pi * (np.arange(48) + phase / 48.) / 48.) % (2 * math.pi)
    speeds = np.linspace(70., 140., 4)
    track_speeds, track_headings = np.meshgrid(speeds, headings, indexing="ij")
    track_speeds = track_speeds.ravel()
    track_headings = track_headings.ravel()
    directions = np.column_stack((np.cos(track_headings), np.sin(track_headings)))
    ages = np.asarray([age for age in geometry.AGES if age < sl.t], float)
    candidates = []
    if len(ages):
        te = sl.t - ages
        clouds = ORIGINS[i, :2][None, None, :] + \
                 track_speeds[:, None, None] * te[None, :, None] * directions[:, None, :]
        starts, ends = [], []
        for j in range(3):
            for sid in surface_ids:
                if sl.visible[j, sid]:
                    starts.append(sl.missiles[j, :2])
                    ends.append(sl.mesh.points[sid, :2])
        if starts:
            starts = np.asarray(starts, float)
            vectors = np.asarray(ends, float) - starts
            denominator = np.einsum("ij,ij->i", vectors, vectors)
            offset = clouds[:, :, None, :] - starts[None, None, :, :]
            lam = np.clip(np.einsum("tasi,si->tas", offset, vectors) /
                          denominator[None, None, :], 0., 1.)
            nearest = starts[None, None, :, :] + lam[:, :, :, None] * vectors[None, None, :, :]
            distances = np.linalg.norm(clouds[:, :, None, :] - nearest, axis=3)
            age_minimum = np.min(distances, axis=2)
            minima = np.min(age_minimum, axis=1)
            useful = np.sum(age_minimum <= p().smoke_radius + 2., axis=1)
            for minimum, useful_ages, heading, speed in zip(
                    minima, useful, track_headings, track_speeds):
                key = (round(float(heading), 12), round(float(speed), 9))
                if key not in occupied:
                    candidates.append((float(minimum), -int(useful_ages), float(heading), float(speed)))
    candidates.sort()
    chosen = [geometry.Track(i, heading, speed, "C2_complete_guard", -1, -2)
              for _, _, heading, speed in candidates[:GEOMETRY_BUDGET["c2_guard_track_limit"]]]
    return chosen, {"sampled_tracks": len(candidates), "selected_tracks": len(chosen),
                    "minimum_horizontal_tube_distance_m": candidates[0][0] if candidates else math.inf}


def event_library(track, sl):
    h = np.array([math.cos(track.heading), math.sin(track.heading)])
    surface_ids = stratified_surface_indices(sl.mesh, 4, 4)
    states = []
    proposed = 0
    for age in geometry.AGES:
        te = sl.t - age
        if te <= 0:
            continue
        cloud_h = ORIGINS[track.platform, :2] + track.speed * te * h
        low_ez = max(0., ORIGINS[track.platform, 2] - .5 * p().gravity * te * te)
        for j in range(3):
            m = sl.missiles[j]
            for sid in surface_ids:
                if not sl.visible[j, sid]:
                    continue
                point = sl.mesh.points[sid]
                dh = point[:2] - m[:2]
                denominator = float(dh @ dh)
                lam = float(np.clip(((cloud_h - m[:2]) @ dh) / denominator, 0., 1.)) if denominator > 0 else 0.
                horizontal_near = m[:2] + lam * dh
                if np.linalg.norm(cloud_h - horizontal_near) > p().smoke_radius + 2.:
                    continue
                line_z = float(m[2] + lam * (point[2] - m[2]))
                for dz in (-6., 0., 6.):
                    proposed += 1
                    cloud_z = line_z + dz
                    ez = cloud_z + p().smoke_sink_speed * age
                    if not low_ez - TIME_TOL <= ez <= ORIGINS[track.platform, 2] + TIME_TOL:
                        continue
                    fuse = math.sqrt(max(0., 2. * (ORIGINS[track.platform, 2] - ez) / p().gravity))
                    release = te - fuse
                    if release < -TIME_TOL:
                        continue
                    states.append((age, max(0., release), fuse, (float(cloud_h[0]), float(cloud_h[1]), cloud_z)))
    if not states:
        return [], {"proposed_vertical_states": proposed, "useful_events": 0, "selected_events": 0}
    clouds = np.asarray([state[3] for state in states], float)
    masks = [0] * len(states)
    count_rows = [[0, 0, 0] for _ in states]
    for j in range(3):
        missiles = np.repeat(sl.missiles[j:j + 1], len(states), axis=0)
        distances = q3._distance_to_segments(clouds, missiles, sl.mesh)
        hits = distances <= p().smoke_radius + 1e-12
        hits &= sl.visible[j][None, :]
        offset = sl.offsets[j]
        for index, row in enumerate(hits):
            local = geometry.bits(row)
            masks[index] |= local << offset
            count_rows[index][j] = int(np.sum(row))
    candidates = {}
    for state, mask, counts in zip(states, masks, count_rows):
        if mask == 0:
            continue
        age, release, fuse, cloud = state
        label = int(np.argmax(counts))
        event = geometry.Event(track.platform, track.heading, track.speed, age, release, fuse,
                               cloud, mask, tuple(counts), label, track.source)
        key = (round(release / .02), round(fuse / .02))
        old = candidates.get(key)
        if old is None or (event.mask.bit_count(), min(event.counts)) > (old.mask.bit_count(), min(old.counts)):
            candidates[key] = event
    pool = list(candidates.values())
    selected = []
    union = 0
    while pool and len(selected) < GEOMETRY_BUDGET["event_limit"]:
        event = max(pool, key=lambda item: (
            (item.mask & ~union).bit_count(), item.mask.bit_count(), min(item.counts), -item.release))
        selected.append(event)
        union |= event.mask
        pool.remove(event)
    return selected, {"proposed_vertical_states": proposed, "useful_events": len(candidates),
                      "selected_events": len(selected)}


def raw_plans(track, sl):
    events, stats = event_library(track, sl)
    plans = []
    for size in range(1, min(3, len(events)) + 1):
        for group in itertools.combinations(events, size):
            ordered = tuple(sorted(group, key=lambda event: event.release))
            if any(v.release - u.release < 1. - TIME_TOL for u, v in zip(ordered[:-1], ordered[1:])):
                continue
            mask = 0
            for event in ordered:
                mask |= event.mask
            counts = tuple(((mask >> (j * sl.n)) & ((1 << sl.n) - 1)).bit_count() for j in range(3))
            plans.append(geometry.Plan(track.platform, track.heading, track.speed, ordered,
                                       mask, counts, track.source))
    return plans, stats


def plan_signature(plan):
    return (round(plan.heading / .004), round(plan.speed / .5),
            tuple((round(event.release / .02), round(event.fuse / .02), event.label)
                  for event in plan.events))


def plan_time_margin(plan):
    if not plan.events:
        return 0.
    return min(min(event.age, p().smoke_duration - event.age) for event in plan.events)


def select_plans(i, sl, c3, guards):
    pools = []
    event_stats = []
    for track in [*c3, *guards]:
        plans, stats = raw_plans(track, sl)
        pools.extend(plans)
        event_stats.append({"source": track.source, **stats})
    best = {}
    for plan in pools:
        key = (plan.mask, plan.source, plan_signature(plan))
        old = best.get(key)
        if old is None or plan_time_margin(plan) > plan_time_margin(old):
            best[key] = plan
    guards_pool = [plan for plan in best.values() if plan.source == "C2_complete_guard"]
    all_pool = list(best.values())
    selected = []
    union = 0
    track_use = {}

    def take(pool, limit):
        nonlocal union
        while pool and len(selected) < limit:
            plan = max(pool, key=lambda item: (
                (item.mask & ~union).bit_count(), item.mask.bit_count(), min(item.counts),
                -track_use.get((round(item.heading / .004), round(item.speed / .5)), 0),
                plan_time_margin(item), len(item.events)))
            selected.append(plan)
            union |= plan.mask
            track_key = (round(plan.heading / .004), round(plan.speed / .5))
            track_use[track_key] = track_use.get(track_key, 0) + 1
            pool.remove(plan)

    guard_reserve = math.ceil(GEOMETRY_BUDGET["plan_limit"] * GEOMETRY_BUDGET["guard_fraction"])
    take(guards_pool, guard_reserve)
    remaining = [plan for plan in all_pool if plan not in selected]
    take(remaining, GEOMETRY_BUDGET["plan_limit"])
    selected.append(geometry.Plan(i, 0., 70., (), 0, (0, 0, 0), "empty"))
    return selected, {"raw_plans": len(pools), "unique_temporal_plans": len(best),
                      "selected_nonempty_plans": len(selected) - 1,
                      "selected_guard_plans": sum(plan.source == "C2_complete_guard" for plan in selected),
                      "event_library_count": len(event_stats),
                      "useful_event_count": sum(row["useful_events"] for row in event_stats),
                      "selected_event_count": sum(row["selected_events"] for row in event_stats)}


def combine_geometry(libraries, sl):
    future = [0] * (len(libraries) + 1)
    for level in range(len(libraries) - 1, -1, -1):
        future[level] = future[level + 1]
        for plan in libraries[level]:
            future[level] |= plan.mask
    states = [(0, ())]
    pruned = 0
    widths = GEOMETRY_BUDGET["beam_widths"]
    per_mask = 8
    for level, library in enumerate(libraries):
        buckets = {}
        for mask, choice in states:
            for plan in library:
                new = mask | plan.mask
                if ((new | future[level + 1]) & sl.universe) != sl.universe:
                    pruned += 1
                    continue
                candidate = choice + (plan,)
                buckets.setdefault(new, []).append(candidate)
        states = []
        for mask, choices in buckets.items():
            distinct = {}
            for choice in choices:
                temporal_key = tuple(plan_signature(plan) for plan in choice)
                distinct.setdefault(temporal_key, choice)
            choices = list(distinct.values())
            choices.sort(key=lambda choice: (
                min((plan_time_margin(plan) for plan in choice if plan.events), default=0.),
                sum(plan.source == "C2_complete_guard" for plan in choice),
                sum(len(plan.events) for plan in choice)), reverse=True)
            states.extend((mask, choice) for choice in choices[:per_mask])
        def state_key(row):
            mask, choice = row
            per_missile = [((mask >> (j * sl.n)) & ((1 << sl.n) - 1)).bit_count() for j in range(3)]
            return ((mask & sl.universe) == sl.universe, mask.bit_count(), min(per_missile),
                    sum(plan.source == "C2_complete_guard" for plan in choice),
                    sum(len(plan.events) for plan in choice))
        states.sort(key=state_key, reverse=True)
        states = states[:widths[level]]
        if not states:
            break
    complete = [row for row in states if (row[0] & sl.universe) == sl.universe]
    best_ratio = max((row[0].bit_count() / sl.universe_count for row in states), default=0.)
    return complete, states[:8], pruned, best_ratio


def solve_geometry_slice(t, mesh):
    sl = geometry.Slice(float(t), mesh)
    libraries = []
    platform_stats = []
    for i in range(5):
        c3, c3_stats = c3_tracks(i, sl)
        occupied = {(round(track.heading, 12), round(track.speed, 9)) for track in c3}
        guards, guard_stats = c2_guard_tracks(i, sl, occupied)
        plans, plan_stats = select_plans(i, sl, c3, guards)
        libraries.append(plans)
        platform_stats.append({"platform": NAMES[i], "c3": c3_stats, "c2_guard": guard_stats,
                               "plans": plan_stats})
    complete, near, pruned, best_ratio = combine_geometry(libraries, sl)
    return {"time_s": float(t), "universe_bits": sl.universe_count,
            "complete_count_after_beam": len(complete), "best_coverage_ratio": best_ratio,
            "pruned_branch_count": pruned, "platform_stats": platform_stats,
            "complete_choices": [choice for _, choice in complete[:8]],
            "near_masks": [mask for mask, _ in near]}


WORKER_MESH = None


def initialise_geometry_worker():
    global WORKER_MESH
    WORKER_MESH = q3.surface_mesh(GEOMETRY_BUDGET["mesh_theta"],
                                  GEOMETRY_BUDGET["mesh_levels"])


def solve_geometry_task(t):
    if WORKER_MESH is None:
        initialise_geometry_worker()
    return solve_geometry_slice(float(t), WORKER_MESH)


def time_grid_values(step, offset, cap):
    return [float(v) for v in np.arange(offset, cap + 1e-12, step) if v < cap]


def choose_geometry_decisions(records):
    rows = []
    for record in records:
        for ordinal, choice in enumerate(record["complete_choices"], 1):
            guard_count = sum(plan.source == "C2_complete_guard" for plan in choice)
            decision = from_geometry_choice(choice,
                f"geometry_common_time_{record['time_s']:.3f}_choice_{ordinal}")
            rows.append({"decision": decision, "time_s": record["time_s"],
                         "guard_plan_count": guard_count})
    unique = {}
    for row in rows:
        unique.setdefault(signature(row["decision"]), row)
    rows = list(unique.values())
    guard = [row for row in rows if row["guard_plan_count"] > 0]
    other = [row for row in rows if row["guard_plan_count"] == 0]
    limit = GEOMETRY_BUDGET["complete_limit"]
    final_size = min(limit, len(rows))
    guard_needed = math.ceil(final_size * GEOMETRY_BUDGET["guard_fraction"])
    if len(guard) < guard_needed:
        raise Q5Error("Q5_GEOMETRY_CANDIDATE_ADMISSION_FAILURE",
                      f"C2 guard choices {len(guard)} below required {guard_needed}")

    def spread(pool, count):
        pool = sorted(pool, key=lambda row: (row["time_s"], row["decision"].source))
        if len(pool) <= count:
            return pool
        indices = np.linspace(0, len(pool) - 1, count, dtype=int)
        return [pool[index] for index in sorted(set(indices.tolist()))]

    selected = spread(guard, min(guard_needed, len(guard)))
    selected_signatures = {signature(row["decision"]) for row in selected}
    remainder = [row for row in [*guard, *other] if signature(row["decision"]) not in selected_signatures]
    selected.extend(spread(remainder, max(0, limit - len(selected))))
    return selected[:limit], {"complete_choice_count": len(rows), "guard_choice_count": len(guard),
                              "selected_count": min(limit, len(selected)),
                              "selected_guard_count": sum(row["guard_plan_count"] > 0 for row in selected[:limit]),
                              "guard_target_count": guard_needed}


def generate_candidates(repair_checkpoint=None):
    cap = min(arrival(j) for j in range(3))
    initial_times = sorted(set(time_grid_values(GEOMETRY_BUDGET["main_time_step_s"], 0., cap) +
                               time_grid_values(GEOMETRY_BUDGET["main_time_step_s"],
                                                GEOMETRY_BUDGET["offset_time_s"], cap)))
    records = []
    cache = {}

    def calculate(executor, times, phase):
        pending = [float(t) for t in times if round(float(t), 9) not in cache]
        for index, (t, record) in enumerate(zip(
                pending, executor.map(solve_geometry_task, pending, chunksize=1)), 1):
            key = round(t, 9)
            cache[key] = record
            records.append(record)
            if index % 20 == 0 or index == len(pending):
                print(json.dumps({"stage": "geometry_candidates", "phase": phase,
                                  "completed": index, "total": len(pending),
                                  "complete_times": sum(row["complete_count_after_beam"] > 0 for row in records)},
                                 ensure_ascii=False), flush=True)

    workers = min(GEOMETRY_BUDGET["parallel_workers"], max(1, os.cpu_count() or 1))
    with concurrent.futures.ProcessPoolExecutor(
            max_workers=workers, initializer=initialise_geometry_worker) as executor:
        calculate(executor, initial_times, "main_and_offset")
        qualified = [row["time_s"] for row in records
                     if row["complete_count_after_beam"] > 0 or row["best_coverage_ratio"] >= .98]
        local_times = sorted({round(value, 9) for t in qualified
                              for value in np.arange(max(0., t - GEOMETRY_BUDGET["local_half_width_s"]),
                                                     min(cap, t + GEOMETRY_BUDGET["local_half_width_s"]) + 1e-12,
                                                     GEOMETRY_BUDGET["local_time_step_s"])})
        calculate(executor, local_times, "local_densification")
    records.sort(key=lambda row: row["time_s"])
    selected, selection_stats = choose_geometry_decisions(records)
    decisions = [row["decision"] for row in selected]
    incumbent, incumbent_identity = registered_incumbent()
    decisions.append(incumbent)
    if repair_checkpoint is not None:
        decisions.append(repair_checkpoint)
    if len(decisions) == 1 + int(repair_checkpoint is not None):
        raise Q5Error("Q5_GEOMETRY_CANDIDATE_ADMISSION_FAILURE", "no complete geometry candidate")
    return decisions, {"budget": ready(GEOMETRY_BUDGET),
                       "main_and_offset_time_count": len(initial_times),
                       "locally_densified_time_count": len(local_times),
                       "evaluated_time_count": len(cache),
                       "complete_timeslice_hit_count": sum(row["complete_count_after_beam"] > 0 for row in records),
                       "near_complete_timeslice_count": sum(row["best_coverage_ratio"] >= .98 for row in records),
                       "selection": selection_stats,
                       "slice_summaries": [{k: v for k, v in row.items() if k not in ("complete_choices", "near_masks", "platform_stats")}
                                           for row in records],
                       "registered_incumbent_forced_into_shortlist": True,
                       "registered_incumbent_input_sha256": incumbent_identity,
                       "repair_checkpoint_used": repair_checkpoint is not None}


def workbook_checkpoint(path):
    """Recover only complete active rows left by a failed final-stage run."""
    if not path.exists():
        return None
    try:
        sheet = load_workbook(path, data_only=True).active
    except Exception:
        return None
    headings, speeds, bombs = [None] * 5, [None] * 5, []
    for row in range(2, 17):
        values = [sheet.cell(row=row, column=col).value for col in range(1, 13)]
        if values[1] is None:
            continue
        if any(v is None for v in values):
            return None
        i = NAMES.index(str(values[0]))
        headings[i] = math.radians(float(values[1])) % (2 * math.pi)
        speeds[i] = float(values[2])
        d = np.array([math.cos(headings[i]), math.sin(headings[i])])
        tau = float((np.asarray(values[4:6], float) - ORIGINS[i, :2]) @ d / speeds[i])
        et = float((np.asarray(values[7:9], float) - ORIGINS[i, :2]) @ d / speeds[i])
        bombs.append(Bomb(i, MISSILE_NAMES.index(str(values[11])), tau, et - tau))
    if not bombs or any(v is None for v in headings) or any(v is None for v in speeds):
        return None
    candidate = Decision(tuple(headings), tuple(speeds), tuple(sorted(bombs, key=lambda b: (b.platform, b.release_s))),
                         "repair_checkpoint_incumbent")
    try:
        validate(candidate)
    except Q5Error:
        return None
    return candidate


def signature(x):
    return (tuple(round(v, 11) for v in x.headings_rad), tuple(round(v, 9) for v in x.speeds_mps),
            tuple((b.platform, b.label, round(b.release_s, 9), round(b.fuse_s, 9)) for b in x.bombs))


def window_nodes(result, pad=.25):
    interval = result["longest_interval_s"]
    if interval is None:
        return np.linspace(0., min(arrival(j) for j in range(3)), 7).tolist()
    left, right = map(float, interval)
    span = right - left
    values = [left - pad, left, left + .25 * span, .5 * (left + right),
              left + .75 * span, right, right + pad]
    cap = min(arrival(j) for j in range(3))
    return sorted(set(min(cap, max(0., float(value))) for value in values))


def event_bomb(event):
    return Bomb(event.platform, event.label, event.release, event.fuse)


def replace_platform_plan(x, plan, source):
    i = plan.platform
    headings, speeds = list(x.headings_rad), list(x.speeds_mps)
    headings[i], speeds[i] = plan.heading % (2 * math.pi), plan.speed
    bombs = [b for b in x.bombs if b.platform != i]
    bombs.extend(event_bomb(event) for event in plan.events)
    y = Decision(tuple(headings), tuple(speeds),
                 tuple(sorted(bombs, key=lambda b: (b.platform, b.release_s))), source)
    try:
        validate(y)
    except Q5Error:
        return None
    return y


def replace_bomb(x, index, bomb, source):
    bombs = list(x.bombs)
    bombs[index] = bomb
    y = Decision(x.headings_rad, x.speeds_mps,
                 tuple(sorted(bombs, key=lambda b: (b.platform, b.release_s))), source)
    try:
        validate(y)
    except Q5Error:
        return None
    return y


def add_bomb(x, bomb, source):
    y = Decision(x.headings_rad, x.speeds_mps,
                 tuple(sorted((*x.bombs, bomb), key=lambda b: (b.platform, b.release_s))), source)
    try:
        validate(y)
    except Q5Error:
        return None
    return y


def delete_bomb(x, index, source):
    y = Decision(x.headings_rad, x.speeds_mps,
                 tuple(b for k, b in enumerate(x.bombs) if k != index), source)
    try:
        validate(y)
    except Q5Error:
        return None
    return y


def proxy_score(x, nodes, mesh):
    margins = np.vstack([strict_margins(x, j, nodes, mesh) for j in range(3)])
    joint = np.max(margins, axis=0)
    finite = np.where(np.isfinite(joint), joint, 1e9)
    return (int(np.sum(finite <= 0.)), float(-np.max(finite)),
            float(-np.mean(np.maximum(finite, 0.))), float(-np.mean(finite)))


def generate_window_candidates(start):
    """Generate multi-time add/delete/replace/plan-swap neighbours around a strict seed."""
    bit_mesh = q3.surface_mesh(GEOMETRY_BUDGET["mesh_theta"], GEOMETRY_BUDGET["mesh_levels"])
    probe_mesh = q3.surface_mesh(PROBE["n_theta"], PROBE["n_levels"])
    seed_result = evaluate(start, FAST)
    nodes = window_nodes(seed_result)
    event_pools = [[] for _ in range(5)]
    plan_pools = [[] for _ in range(5)]
    pool_stats = []

    for t in nodes:
        sl = geometry.Slice(float(t), bit_mesh)
        for i in range(5):
            current = geometry.Track(i, start.headings_rad[i], start.speeds_mps[i],
                                     "incumbent_window_track", -1, -1)
            c3, c3_stats = c3_tracks(i, sl)
            occupied = {(round(track.heading, 12), round(track.speed, 9)) for track in c3}
            guards, guard_stats = c2_guard_tracks(i, sl, occupied)
            tracks = [current, *c3[:3], *guards[:1]]
            for track in tracks:
                plans, stats = raw_plans(track, sl)
                plans.sort(key=lambda plan: (plan.mask.bit_count(), min(plan.counts),
                                             plan_time_margin(plan), len(plan.events)), reverse=True)
                plan_pools[i].extend(plans[:3])
                for plan in plans[:6]:
                    event_pools[i].extend(plan.events)
                pool_stats.append({"time_s": float(t), "platform": NAMES[i], "source": track.source,
                                   **stats, "retained_plans": min(3, len(plans))})
            pool_stats.append({"time_s": float(t), "platform": NAMES[i],
                               "c3_track_stats": c3_stats, "guard_track_stats": guard_stats})

    for i in range(5):
        unique_events = {}
        for event in event_pools[i]:
            key = (round(event.release / .02), round(event.fuse / .02), event.label)
            old = unique_events.get(key)
            if old is None or (event.mask.bit_count(), min(event.counts)) > (old.mask.bit_count(), min(old.counts)):
                unique_events[key] = event
        centre = float(np.mean(nodes))
        event_pools[i] = sorted(unique_events.values(), key=lambda event: (
            event.mask.bit_count(), min(event.counts),
            -abs(event.release + event.fuse + .5 * p().smoke_duration - centre)), reverse=True)[:12]
        unique_plans = {}
        for plan in plan_pools[i]:
            unique_plans.setdefault(plan_signature(plan), plan)
        plan_pools[i] = sorted(unique_plans.values(), key=lambda plan: (
            plan.mask.bit_count(), min(plan.counts), plan_time_margin(plan), len(plan.events)), reverse=True)[:4]

    candidates = {signature(start): start}
    move_counts = {name: 0 for name in ("add", "delete", "replace", "swap_plan", "exchange")}

    for index in range(len(start.bombs)):
        y = delete_bomb(start, index, f"topology_delete_{index}")
        if y is not None:
            candidates.setdefault(signature(y), y)
            move_counts["delete"] += 1

    for i in range(5):
        if len(by_platform(start, i)) < 3:
            for event in event_pools[i]:
                y = add_bomb(start, event_bomb(event), f"topology_add_{NAMES[i]}")
                if y is not None:
                    candidates.setdefault(signature(y), y)
                    move_counts["add"] += 1

    for index, old in enumerate(start.bombs):
        for event in event_pools[old.platform][:4]:
            y = replace_bomb(start, index, event_bomb(event), f"topology_replace_{index}")
            if y is not None:
                candidates.setdefault(signature(y), y)
                move_counts["replace"] += 1

    swap_candidates = []
    for plans in plan_pools:
        for plan in plans:
            y = replace_platform_plan(start, plan, f"topology_swap_plan_{NAMES[plan.platform]}")
            if y is not None:
                candidates.setdefault(signature(y), y)
                swap_candidates.append(y)
                move_counts["swap_plan"] += 1

    for first, second in itertools.combinations(swap_candidates[:10], 2):
        changed_first = {i for i in range(5) if by_platform(first, i) != by_platform(start, i)}
        changed_second = {i for i in range(5) if by_platform(second, i) != by_platform(start, i)}
        if len(changed_first) != 1 or len(changed_second) != 1 or changed_first == changed_second:
            continue
        i, k = next(iter(changed_first)), next(iter(changed_second))
        headings, speeds = list(start.headings_rad), list(start.speeds_mps)
        headings[i], speeds[i] = first.headings_rad[i], first.speeds_mps[i]
        headings[k], speeds[k] = second.headings_rad[k], second.speeds_mps[k]
        bombs = [b for b in start.bombs if b.platform not in (i, k)]
        bombs.extend(by_platform(first, i))
        bombs.extend(by_platform(second, k))
        y = Decision(tuple(headings), tuple(speeds), tuple(sorted(bombs, key=lambda b: (b.platform, b.release_s))),
                     f"topology_exchange_{NAMES[i]}_{NAMES[k]}")
        try:
            validate(y)
        except Q5Error:
            continue
        candidates.setdefault(signature(y), y)
        move_counts["exchange"] += 1
        if move_counts["exchange"] >= 10:
            break

    proxy_rows = [{"decision": x, "proxy": proxy_score(x, nodes, bit_mesh)} for x in candidates.values()]
    proxy_rows.sort(key=lambda row: row["proxy"], reverse=True)
    proxy_rows = proxy_rows[:64]
    probe_rows = [{"decision": row["decision"], "proxy": row["proxy"],
                   "result": evaluate(row["decision"], PROBE, probe_mesh)} for row in proxy_rows]
    probe_rows = ranked(probe_rows, "result", RANK_TOLERANCE[PROBE["name"]])[:20]
    return [row["decision"] for row in probe_rows], {
        "anchor_nodes_s": nodes,
        "seed_longest_continuous_s": seed_result["longest_continuous_s"],
        "seed_total_intersection_s": seed_result["total_intersection_s"],
        "event_pool_sizes": [len(pool) for pool in event_pools],
        "plan_pool_sizes": [len(pool) for pool in plan_pools],
        "move_counts": move_counts,
        "unique_neighbour_count": len(candidates) - 1,
        "proxy_shortlist_count": len(proxy_rows),
        "probe_shortlist_count": len(probe_rows),
        "pool_diagnostics": pool_stats,
    }


def mutate(x, kind, index, delta, source):
    headings, speeds, bombs = list(x.headings_rad), list(x.speeds_mps), list(x.bombs)
    if kind == "heading":
        headings[index] = (headings[index] + delta) % (2 * math.pi)
    elif kind == "speed":
        speeds[index] = min(140., max(70., speeds[index] + delta))
    elif kind == "platform_shift":
        for k, b in enumerate(bombs):
            if b.platform == index:
                bombs[k] = replace(b, release_s=b.release_s + delta)
    elif kind == "release":
        bombs[index] = replace(bombs[index], release_s=bombs[index].release_s + delta)
    elif kind == "fuse":
        bombs[index] = replace(bombs[index], fuse_s=bombs[index].fuse_s + delta)
    y = Decision(tuple(headings), tuple(speeds), tuple(bombs), source)
    try:
        validate(y)
    except Q5Error:
        return None
    return y


def refine(start, mesh, probe_mesh, ordinal):
    started = time.perf_counter()
    current = start
    score = evaluate(current, FAST, mesh)
    cache = {signature(current): score}
    probe_cache = {signature(current): evaluate(current, PROBE, probe_mesh)}
    stages = []
    terminated_by = "completed_scales"
    print(json.dumps({"stage": "refine_start", "ordinal": ordinal, "source": start.source,
                      "longest_continuous_s": score["longest_continuous_s"],
                      "total_intersection_s": score["total_intersection_s"],
                      "budget": REFINE_BUDGET}, ensure_ascii=False), flush=True)
    coordinates = [("heading", i) for i in range(5)] + [("speed", i) for i in range(5)] + \
                  [("platform_shift", i) for i in range(5)] + [("release", i) for i in range(len(start.bombs))] + \
                  [("fuse", i) for i in range(len(start.bombs))]
    scales = [(.003, .5, .05, .05, .05), (.001, .2, .02, .02, .02),
              (.0003, .05, .005, .005, .005)]
    exhausted = False
    for stage, scale in enumerate(scales, 1):
        steps = dict(zip(("heading", "speed", "platform_shift", "release", "fuse"), scale))
        sweeps = []
        for sweep in range(1, 4):
            accepted = primary_accepted = 0
            start_objective = score["objective_s"]
            for kind, index in coordinates:
                elapsed = time.perf_counter() - started
                if elapsed >= REFINE_BUDGET["wall_time_s_per_start"]:
                    terminated_by, exhausted = "wall_time_budget", True
                    break
                if len(probe_cache) >= REFINE_BUDGET["probe_evaluations_per_start"]:
                    terminated_by, exhausted = "probe_evaluation_budget", True
                    break
                if len(cache) >= REFINE_BUDGET["fast_evaluations_per_start"]:
                    terminated_by, exhausted = "fast_evaluation_budget", True
                    break
                probe_options = [(current, probe_cache[signature(current)])]
                for sign in (-1., 1.):
                    y = mutate(current, kind, index, sign * steps[kind],
                               f"original_variable_refine_{ordinal}")
                    if y is None:
                        continue
                    key = signature(y)
                    if key not in probe_cache:
                        probe_cache[key] = evaluate(y, PROBE, probe_mesh)
                        if len(probe_cache) % REFINE_BUDGET["progress_every_probe_evaluations"] == 0:
                            print(json.dumps({"stage": "refine_progress", "ordinal": ordinal,
                                              "scale_stage": stage, "sweep": sweep,
                                              "probe_evaluations": len(probe_cache),
                                              "fast_evaluations": len(cache),
                                              "elapsed_s": time.perf_counter() - started,
                                              "longest_continuous_s": score["longest_continuous_s"],
                                              "total_intersection_s": score["total_intersection_s"]},
                                             ensure_ascii=False), flush=True)
                    probe_options.append((y, probe_cache[key]))
                probe_best = ranked([{"decision": decision, "result": result}
                                     for decision, result in probe_options],
                                    "result", RANK_TOLERANCE[PROBE["name"]])[0]
                candidate = probe_best["decision"]
                if signature(candidate) == signature(current):
                    continue
                key = signature(candidate)
                if key not in cache:
                    cache[key] = evaluate(candidate, FAST, mesh)
                candidate_score = cache[key]
                primary_improves = (
                    candidate_score["longest_continuous_s"] > score["longest_continuous_s"] + 1e-9)
                if better_result(candidate_score, score):
                    current, score = candidate, candidate_score
                    accepted += 1
                    primary_accepted += int(primary_improves)
            sweeps.append({"sweep": sweep, "accepted_coordinate_moves": accepted,
                           "primary_improving_moves": primary_accepted,
                           "longest_continuous_s": score["longest_continuous_s"],
                           "total_intersection_s": score["total_intersection_s"]})
            print(json.dumps({"stage": "refine_sweep_complete", "ordinal": ordinal,
                              "scale_stage": stage, **sweeps[-1],
                              "probe_evaluations": len(probe_cache),
                              "fast_evaluations": len(cache),
                              "elapsed_s": time.perf_counter() - started,
                              "terminated_by": terminated_by if exhausted else None},
                             ensure_ascii=False), flush=True)
            if exhausted:
                break
            if score["longest_continuous_s"] <= start_objective + 1e-9:
                break
        stages.append({"stage": stage, "steps": steps, "sweeps": sweeps,
                       "longest_continuous_s": score["longest_continuous_s"],
                       "total_intersection_s": score["total_intersection_s"]})
        if exhausted:
            break
    print(json.dumps({"stage": "refine_complete", "ordinal": ordinal,
                      "terminated_by": terminated_by,
                      "probe_evaluations": len(probe_cache), "fast_evaluations": len(cache),
                      "elapsed_s": time.perf_counter() - started,
                      "longest_continuous_s": score["longest_continuous_s"],
                      "total_intersection_s": score["total_intersection_s"]},
                     ensure_ascii=False), flush=True)
    return {"start": decision_record(start), "chosen": decision_record(current), "fast": score,
            "fast_function_evaluations": len(cache), "probe_function_evaluations": len(probe_cache),
            "elapsed_s": time.perf_counter() - started, "terminated_by": terminated_by,
            "stages": stages, "_decision": current}


def expanded(intervals, pad, cap):
    return q3.merge_intervals([[max(0., a - pad), min(cap, b + pad)] for a, b in intervals])


def diagnostics(x, precise_result, precise_mesh, dense_mesh):
    label_only = []
    for j in range(3):
        subset = tuple(b for b in x.bombs if b.label == j)
        y = Decision(x.headings_rad, x.speeds_mps, subset, "label_only_diagnostic")
        label_only.append(evaluate(y, PRECISE, precise_mesh)["by_missile"][j] if subset else {"duration_s": 0., "intervals_s": []})
    durations = [r["duration_s"] for r in label_only]
    dense_restrict = [expanded(record["intervals_s"], .10, arrival(j)) for j, record in enumerate(precise_result["by_missile"])]
    dense_result = evaluate(x, DENSIFIED, dense_mesh, dense_restrict)
    longest_change = dense_result["longest_continuous_s"] - precise_result["longest_continuous_s"]
    total_change = dense_result["total_intersection_s"] - precise_result["total_intersection_s"]
    precise_interval = precise_result["longest_interval_s"]
    dense_interval = dense_result["longest_interval_s"]
    endpoint_changes = ([dense_interval[k] - precise_interval[k] for k in (0, 1)]
                        if precise_interval is not None and dense_interval is not None else [math.inf, math.inf])
    cross_raw = precise_result["sum_duration_missile_s"] - sum(durations)
    gaps = [v.release_s - u.release_s for i in range(5) for u, v in zip(by_platform(x, i)[:-1], by_platform(x, i)[1:])]
    return {"label_only_by_missile": label_only, "label_only_durations_s": durations,
            "cross_label_gain_raw_missile_s": cross_raw, "cross_label_gain_missile_s": max(0., cross_raw),
            "cross_label_monotonicity_pass": all(a + 2e-6 >= b for a, b in zip(precise_result["durations_s"], durations)),
            "minimum_release_gap_s": min(gaps) if gaps else None, "densified": dense_result,
            "densified_change_s": longest_change,
            "densified_longest_change_s": longest_change,
            "densified_total_intersection_change_s": total_change,
            "densified_longest_endpoint_changes_s": endpoint_changes,
            "precision_convergence_pair": "precise_to_densified",
            "precision_convergence_change_s": longest_change,
            "precision_convergence_pass": (abs(longest_change) <= 3e-5 and abs(total_change) <= 3e-5 and
                                             max(abs(value) for value in endpoint_changes) <= 3e-5)}


def single_bomb_duration(x, bomb, mesh):
    y = Decision(x.headings_rad, x.speeds_mps, (bomb,), "excel_single_bomb_semantics")
    return solve_intervals(y, bomb.label, PRECISE, mesh)["duration_s"]


def write_excel(x, path, mesh, logical_path=None):
    path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(ROOT / "data/附件/result3.xlsx", path)
    book = load_workbook(path)
    sheet = book[book.sheetnames[0]]
    written = []
    for i in range(5):
        bombs = by_platform(x, i)
        for slot in range(3):
            row = 2 + 3 * i + slot
            if slot >= len(bombs):
                for col in range(1, 13):
                    sheet.cell(row=row, column=col).value = None
                continue
            b = bombs[slot]
            rp, ep = release_point(x, b), explosion_point(x, b)
            duration = single_bomb_duration(x, b, mesh)
            values = [NAMES[i], math.degrees(x.headings_rad[i]) % 360, x.speeds_mps[i], slot + 1,
                      *rp.tolist(), *ep.tolist(), duration, MISSILE_NAMES[b.label]]
            for col, value in enumerate(values, 1):
                sheet.cell(row=row, column=col, value=value if col in (1, 12) else float(value))
            written.append({"row": row, "platform": NAMES[i], "bomb_number": slot + 1,
                            "label": MISSILE_NAMES[b.label], "independent_duration_s": duration})
    book.save(path)

    sheet = load_workbook(path, data_only=True).active
    headings, speeds = list(x.headings_rad), list(x.speeds_mps)
    bombs = []
    duration_residuals = []
    nonblank_rows = 0
    for row in range(2, 17):
        values = [sheet.cell(row=row, column=col).value for col in range(1, 13)]
        if values[0] is None:
            if any(v is not None for v in values):
                raise Q5Error("Q5_EXCEL_ROUNDTRIP_FAILURE", f"partially blank row {row}")
            continue
        if any(v is None for v in values):
            raise Q5Error("Q5_EXCEL_ROUNDTRIP_FAILURE", f"incomplete active row {row}")
        nonblank_rows += 1
        i = NAMES.index(str(values[0]))
        headings[i] = math.radians(float(values[1])) % (2 * math.pi)
        speeds[i] = float(values[2])
        d = np.array([math.cos(headings[i]), math.sin(headings[i])])
        release_xy = np.asarray(values[4:6], float)
        explosion_xy = np.asarray(values[7:9], float)
        tau = float((release_xy - ORIGINS[i, :2]) @ d / speeds[i])
        et = float((explosion_xy - ORIGINS[i, :2]) @ d / speeds[i])
        bombs.append(Bomb(i, MISSILE_NAMES.index(str(values[11])), tau, et - tau))
    y = Decision(tuple(headings), tuple(speeds), tuple(sorted(bombs, key=lambda b: (b.platform, b.release_s))), "excel_roundtrip")
    validate(y)
    for row in range(2, 17):
        if sheet.cell(row=row, column=1).value is None:
            continue
        i = NAMES.index(str(sheet.cell(row=row, column=1).value))
        number = int(sheet.cell(row=row, column=4).value)
        b = by_platform(y, i)[number - 1]
        recomputed = single_bomb_duration(y, b, mesh)
        duration_residuals.append(abs(recomputed - float(sheet.cell(row=row, column=11).value)))
    full = evaluate(y, PRECISE, mesh)
    residual = max(duration_residuals, default=0.)
    if nonblank_rows != len(x.bombs) or residual > 2e-6:
        raise Q5Error("Q5_EXCEL_ROUNDTRIP_FAILURE", f"rows={nonblank_rows}, residual={residual}")
    output_path = logical_path if logical_path is not None else path
    try:
        output_name = output_path.relative_to(ROOT).as_posix()
    except ValueError:
        output_name = str(output_path)
    return {"path": output_name, "sha256": sha(path), "status": "pass",
            "active_row_count": nonblank_rows, "unused_reserved_row_count": 15 - nonblank_rows,
            "row_duration_max_residual_s": residual, "reconstructed_decision": decision_record(y),
            "recomputed_full_objective": full, "written_rows": written}


def make_report(result, result_path):
    best = result["formal_best"]
    lines = ["# Q5 正式工作求解实现报告", "", f"- 规格：`{SPEC_ID}`",
             f"- 最长连续三导弹同时严格全遮蔽时长：`{best['precise']['longest_continuous_s']:.9f} s`",
             f"- 最长连续区间：`{best['precise']['longest_interval_s']}`",
             f"- 三导弹同时全遮蔽总测度：`{best['precise']['total_intersection_s']:.9f} s`",
             f"- 分导弹时长：`{best['precise']['durations_s']}`",
             f"- 逐导弹覆盖总量：`{best['precise']['sum_duration_missile_s']:.9f} 导弹·s`",
             f"- 活跃烟幕弹：`{best['decision']['active_bomb_count']}` 枚",
             f"- 机器可读结果：`{result_path.relative_to(ROOT).as_posix()}`", "", "## 求解路线", "",
             "以历史严格可行解的最长公共分支为首个锚定窗，在端点、四分点、中点和外扩点生成完整可达事件池；显式枚举增弹、删弹、换弹、整个平台计划替换和双平台交换邻域。代理筛选后全部候选进入连续严格核，按最长连续全遮蔽时长、总交集测度、逐导弹覆盖总量的字典序排名，再回到原变量精修。逐弹标签只说明主要干扰对象，严格评价始终让全部活跃烟幕参与三枚导弹。", "", "## 最终工作方案", ""]
    for row in best["decision"]["platforms"]:
        labels = [b["label"] for b in row["bombs"]]
        lines.append(f"- {row['platform']}：航向 {row['heading_deg']:.9f}°，速度 {row['speed_mps']:.9f} m/s，标签序列 {labels}。")
    lines.extend(["", "## 数值核验", "",
                  f"- 最小相邻投放间隔：`{result['verification']['minimum_release_gap_s']:.9f} s`。",
                  f"- 标签外烟幕对逐导弹覆盖总量的增量：`{result['verification']['cross_label_gain_missile_s']:.9f} 导弹·s`。",
                  f"- 精确层到加密层的最长连续时长变化：`{result['verification']['densified_longest_change_s']:.9e} s`。",
                  f"- 精确层到加密层的总交集测度变化：`{result['verification']['densified_total_intersection_change_s']:.9e} s`。",
                  f"- 最终相邻加密层：`{result['verification']['precision_convergence_pair']}`，变化 `{result['verification']['precision_convergence_change_s']:.9e} s`。",
                  f"- Excel 回读：`{result['excel_roundtrip']['status']}`，未用预留行 {result['excel_roundtrip']['unused_reserved_row_count']} 行。", "",
                  "该数值是冻结模型与计算预算下的可复核工作解，不宣称全局最优；旧交集总测度方案只作为历史 seed，经新评价器重算最长连续时长后才参与选择。"])
    return "\n".join(lines) + "\n"


def run_pilot(output):
    started = time.perf_counter()
    identities = verify_identities()
    incumbent, incumbent_identity = registered_incumbent()
    decisions, generation = generate_window_candidates(incumbent)
    if not any(signature(x) == signature(incumbent) for x in decisions):
        decisions.append(incumbent)
    fast_mesh = q3.surface_mesh(FAST["n_theta"], FAST["n_levels"])
    fast = ranked([{"decision": x, "result": evaluate(x, FAST, fast_mesh)} for x in decisions],
                  "result", RANK_TOLERANCE[FAST["name"]])
    screen_mesh = q3.surface_mesh(SCREEN["n_theta"], SCREEN["n_levels"])
    screen_source = fast[:5]
    if not any(signature(row["decision"]) == signature(incumbent) for row in screen_source):
        screen_source.append(next(row for row in fast if signature(row["decision"]) == signature(incumbent)))
    screened = ranked([{**row, "screen": evaluate(row["decision"], SCREEN, screen_mesh)}
                       for row in screen_source], "screen", RANK_TOLERANCE[SCREEN["name"]])
    baseline = next(row for row in screened if signature(row["decision"]) == signature(incumbent))
    best = screened[0]
    payload = {
        "schema_version": "1.0",
        "spec_id": SPEC_ID,
        "status": "pilot_pass",
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "baseline": {"source": incumbent.source, **baseline["screen"]},
        "best_before_refinement": {"source": best["decision"].source,
                                   "decision": decision_record(best["decision"]), **best["screen"]},
        "improvement_over_baseline_s": (best["screen"]["longest_continuous_s"] -
                                         baseline["screen"]["longest_continuous_s"]),
        "generation": generation,
        "candidate_count": len(decisions),
        "screen_candidate_count": len(screened),
        "historical_seed_sha256": incumbent_identity,
        "identity": identities,
        "elapsed_s": time.perf_counter() - started,
        "claim_boundary": "Early candidate-layer diagnostic only; not a precise or densified formal result."
    }
    output = output.resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile("w", encoding="utf-8", delete=False,
                                     dir=output.parent, prefix=output.name + ".", suffix=".tmp") as handle:
        handle.write(json.dumps(ready(payload), ensure_ascii=False, indent=2))
        temporary = Path(handle.name)
    os.replace(temporary, output)
    return payload


def run(output, report_path, excel_path):
    started = time.perf_counter()
    identities = verify_identities()
    repair = workbook_checkpoint(excel_path)
    incumbent, incumbent_identity = registered_incumbent()
    decisions, generation = generate_window_candidates(incumbent)
    if not any(signature(x) == signature(incumbent) for x in decisions):
        decisions.append(incumbent)
    if repair is not None and not any(signature(x) == signature(repair) for x in decisions):
        decisions.append(repair)
    generation["historical_seed_sha256"] = incumbent_identity
    generation["repair_checkpoint_used"] = repair is not None
    print(json.dumps({"stage": "strict_fast", "candidate_count": len(decisions)}, ensure_ascii=False), flush=True)
    fast_mesh = q3.surface_mesh(FAST["n_theta"], FAST["n_levels"])
    fast = [{"decision": x, "result": evaluate(x, FAST, fast_mesh)} for x in decisions]
    fast = ranked(fast, "result", RANK_TOLERANCE[FAST["name"]])
    print(json.dumps({"stage": "strict_fast_complete", "best_source": fast[0]["decision"].source,
                      "longest_continuous_s": fast[0]["result"]["longest_continuous_s"],
                      "total_intersection_s": fast[0]["result"]["total_intersection_s"]},
                     ensure_ascii=False), flush=True)

    print(json.dumps({"stage": "strict_screen", "candidate_count": min(20, len(fast))}, ensure_ascii=False), flush=True)
    screen_mesh = q3.surface_mesh(SCREEN["n_theta"], SCREEN["n_levels"])
    shortlist = fast[:20]
    if not any(signature(v["decision"]) == signature(incumbent) for v in shortlist):
        shortlist.append(next(v for v in fast if signature(v["decision"]) == signature(incumbent)))
    if repair is not None and not any(signature(v["decision"]) == signature(repair) for v in shortlist):
        repair_row = next((v for v in fast if signature(v["decision"]) == signature(repair)), None)
        if repair_row is not None:
            shortlist.append(repair_row)
    screened = [{**v, "screen": evaluate(v["decision"], SCREEN, screen_mesh)} for v in shortlist]
    screened = ranked(screened, "screen", RANK_TOLERANCE[SCREEN["name"]])
    print(json.dumps({"stage": "strict_screen_complete", "best_source": screened[0]["decision"].source,
                      "longest_continuous_s": screened[0]["screen"]["longest_continuous_s"],
                      "total_intersection_s": screened[0]["screen"]["total_intersection_s"]},
                     ensure_ascii=False), flush=True)

    print(json.dumps({"stage": "original_variable_refinement"}, ensure_ascii=False), flush=True)
    refinement_starts = screened[:4]
    incumbent_screened = next(v for v in screened if signature(v["decision"]) == signature(incumbent))
    if not any(signature(v["decision"]) == signature(incumbent) for v in refinement_starts):
        refinement_starts.append(incumbent_screened)
    probe_mesh = q3.surface_mesh(PROBE["n_theta"], PROBE["n_levels"])
    refinements = [refine(v["decision"], fast_mesh, probe_mesh, index)
                   for index, v in enumerate(refinement_starts, 1)]
    refined_decisions = [v["_decision"] for v in refinements] + [v["decision"] for v in screened]
    unique = {signature(v): v for v in refined_decisions}
    rescreened = [{"decision": x, "screen": evaluate(x, SCREEN, screen_mesh)} for x in unique.values()]
    rescreened = ranked(rescreened, "screen", RANK_TOLERANCE[SCREEN["name"]])

    print(json.dumps({"stage": "strict_precise", "candidate_count": min(12, len(rescreened))}, ensure_ascii=False), flush=True)
    precise_mesh = q3.surface_mesh(PRECISE["n_theta"], PRECISE["n_levels"])
    screen_pool = rescreened[:12]
    if not any(signature(v["decision"]) == signature(incumbent) for v in screen_pool):
        screen_pool.append({"decision": incumbent, "screen": evaluate(incumbent, SCREEN, screen_mesh)})
    precise_pool = screen_pool[:6]
    if not any(signature(v["decision"]) == signature(incumbent) for v in precise_pool):
        precise_pool.append(next(v for v in screen_pool if signature(v["decision"]) == signature(incumbent)))
    precise = []
    for index, row in enumerate(precise_pool, 1):
        evaluated = evaluate(row["decision"], PRECISE, precise_mesh)
        precise.append({**row, "precise": evaluated})
        print(json.dumps({"stage": "strict_precise_progress", "completed": index,
                          "total": len(precise_pool), "source": row["decision"].source,
                          "longest_continuous_s": evaluated["longest_continuous_s"],
                          "total_intersection_s": evaluated["total_intersection_s"]},
                         ensure_ascii=False), flush=True)
    precise = ranked(precise, "precise", RANK_TOLERANCE[PRECISE["name"]])
    winner = precise[0]

    print(json.dumps({"stage": "densified_and_roundtrip"}, ensure_ascii=False), flush=True)
    dense_mesh = q3.surface_mesh(DENSIFIED["n_theta"], DENSIFIED["n_levels"])
    verification = diagnostics(winner["decision"], winner["precise"], precise_mesh, dense_mesh)
    print(json.dumps({"stage": "densified_complete",
                      "longest_continuous_s": verification["densified"]["longest_continuous_s"],
                      "total_intersection_s": verification["densified"]["total_intersection_s"],
                      "longest_change_s": verification["densified_longest_change_s"],
                      "total_change_s": verification["densified_total_intersection_change_s"]},
                     ensure_ascii=False), flush=True)
    incumbent_precise = next(v["precise"] for v in precise if signature(v["decision"]) == signature(incumbent))
    floor = max(LEGACY_LONGEST_CONTINUOUS_S, incumbent_precise["longest_continuous_s"])
    if winner["precise"]["longest_continuous_s"] + 3e-5 < floor:
        raise Q5Error("Q5_LONGEST_INCUMBENT_LOSS", "formal winner below retained longest-window incumbent")
    if not verification["precision_convergence_pass"]:
        raise Q5Error("Q5_LONGEST_PRECISION_STABILITY_FAILURE", str(verification))

    with tempfile.TemporaryDirectory(prefix="q5_spec4_") as temporary:
        temporary = Path(temporary)
        staged_excel = temporary / "result3.xlsx"
        excel = write_excel(winner["decision"], staged_excel, precise_mesh, logical_path=excel_path)
        result = {
            "schema_version": "4.0", "spec_id": SPEC_ID, "result_id": RESULT_ID,
            "status": "working_verified", "generated_at": datetime.now(timezone.utc).isoformat(),
            "objective": "longest connected interval in the intersection of the three strict complete-cylinder obscuration sets",
            "objective_hierarchy": ["longest_continuous_s", "total_intersection_s",
                                    "sum_duration_missile_s", "minimum_missile_duration_s"],
            "formal_best": {"decision": decision_record(winner["decision"]), "precise": winner["precise"],
                            "densified": verification["densified"]},
            "single_chain": {"generation": generation,
                             "fast_records": [{"source": v["decision"].source, **v["result"]} for v in fast],
                             "screen_records": [{"source": v["decision"].source, **v["screen"]} for v in screened],
                             "refinements": [{k: value for k, value in v.items() if k != "_decision"} for v in refinements],
                             "precise_records": [{"source": v["decision"].source, **v["precise"]} for v in precise],
                             "incumbent_precise_longest_continuous_s": incumbent_precise["longest_continuous_s"],
                             "legacy_longest_continuous_floor_s": LEGACY_LONGEST_CONTINUOUS_S,
                             "legacy_total_intersection_diagnostic_s": LEGACY_TOTAL_INTERSECTION_S,
                             "incumbent_preserved": winner["precise"]["longest_continuous_s"] + 3e-5 >= floor},
            "verification": verification, "excel_roundtrip": excel,
            "identity": {**identities, "src/q5/solve.py": sha(Path(__file__))},
            "precision_layers": {"bitset": {"n_theta": GEOMETRY_BUDGET["mesh_theta"],
                                             "n_levels": GEOMETRY_BUDGET["mesh_levels"]},
                                 "fast": FAST, "screen": SCREEN, "precise": PRECISE,
                                 "densified": DENSIFIED},
            "software": {"python": platform.python_version(), "numpy": np.__version__,
                         "deterministic": True, "random_seed": None},
            "elapsed_s": time.perf_counter() - started,
            "claims": {"global_optimality": False, "continuous_candidate_domain_exhausted": False,
                       "legacy_value_promoted_without_recomputation": False,
                       "labels_filter_physical_clouds": False,
                       "longest_continuous_all_three_missiles_is_primary": True,
                       "total_intersection_is_secondary_objective": True,
                       "sum_duration_is_primary_objective": False,
                       "same_platform_release_gap_enforced": True,
                       "smoke_sink_speed_mps": p().smoke_sink_speed,
                       "smoke_duration_s": p().smoke_duration}
        }
        staged_result = temporary / "q5_result.json"
        staged_report = temporary / "q5_implementation_report.md"
        staged_result.write_text(json.dumps(ready(result), ensure_ascii=False, indent=2), encoding="utf-8")
        staged_report.write_text(make_report(result, output), encoding="utf-8")
        output.parent.mkdir(parents=True, exist_ok=True)
        report_path.parent.mkdir(parents=True, exist_ok=True)
        excel_path.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(staged_result, output)
        shutil.copyfile(staged_report, report_path)
        shutil.copyfile(staged_excel, excel_path)
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--pilot-output", type=Path)
    parser.add_argument("--output", type=Path, default=ROOT / "docs/q5_result.json")
    parser.add_argument("--report", type=Path, default=ROOT / "docs/q5_implementation_report.md")
    parser.add_argument("--excel", type=Path, default=ROOT / "output/q5/result3.xlsx")
    args = parser.parse_args()
    try:
        if args.pilot_output is not None:
            result = run_pilot(args.pilot_output)
            print(json.dumps({"status": "pass", "pilot": True,
                              "best_longest_continuous_s": result["best_before_refinement"]["longest_continuous_s"],
                              "baseline_longest_continuous_s": result["baseline"]["longest_continuous_s"],
                              "elapsed_s": result["elapsed_s"]}, ensure_ascii=False))
            return 0
        result = run(args.output.resolve(), args.report.resolve(), args.excel.resolve())
    except Q5Error as exc:
        print(json.dumps({"status": "fail", "code": exc.code, "message": str(exc)}, ensure_ascii=False))
        return 2
    print(json.dumps({"status": "pass", "result_id": result["result_id"],
                      "longest_continuous_s": result["formal_best"]["precise"]["longest_continuous_s"],
                      "total_intersection_s": result["formal_best"]["precise"]["total_intersection_s"],
                      "durations_s": result["formal_best"]["precise"]["durations_s"],
                      "elapsed_s": result["elapsed_s"]}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
