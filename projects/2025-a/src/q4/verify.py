"""Independent Q4 E1/E2 verifier for SPEC-Q4-VERIFY-1.0."""
from __future__ import annotations

import argparse, hashlib, json, math, platform, sys, time
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
import scipy
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT))
from src.q4 import solve  # noqa: E402

TOL={"space_m":1e-7,"time_s":1e-10,"margin_m":1e-8,"duration_s":2e-4,"root_margin_m":2e-6}
EXPECTED={
    "data/A题.pdf":"a37f6aad30b16ea09da9cb320ce194b12d3a28874008cf9ff19549f6c6079447",
    "data/附件/result2.xlsx":"c681d5e378538f71c77fca199a3ca8303a04dbcfc7bd95f870ae22f01ab69f91",
    "planning/26_q4_candidate_set.md":"f612eb8b92283bbb1ee533f639d467595cf73613450575216b47c6ae595eec33",
    "src/q3/selection_pilot.py":"33f67bee3a033ffebc48dd7a6275846c991977f6035554b75cd055de80430f64",
}

class VerificationFailure(RuntimeError):pass
def require(ok,msg):
    if not ok:raise VerificationFailure(msg)
def sha(path):
    h=hashlib.sha256()
    with Path(path).open("rb") as f:
        for b in iter(lambda:f.read(1<<20),b""):h.update(b)
    return h.hexdigest()

def load():
    payload=json.loads((ROOT/"docs/q4_result.json").read_text(encoding="utf-8"));d=payload["formal_best"]["decision"]
    x=solve.Decision(tuple(d["headings_rad"]),tuple(d["speeds_mps"]),tuple(d["release_times_s"]),tuple(d["fuse_delays_s"]),"q4_e1e2",0)
    solve.validate(x);return payload,x

def independent_mesh(n_theta,n_levels):
    p=solve.p();angles=np.arange(n_theta)*2*math.pi/n_theta;heights=np.linspace(0,p.target_height,n_levels)
    side_theta=np.tile(angles,n_levels);side_z=np.repeat(heights,n_theta)
    side=np.column_stack((p.target_base_center[0]+p.target_radius*np.cos(side_theta),p.target_base_center[1]+p.target_radius*np.sin(side_theta),side_z))
    side_normals=np.column_stack((np.cos(side_theta),np.sin(side_theta),np.zeros_like(side_theta)))
    radii=np.linspace(0,p.target_radius,n_levels)[1:];top_theta=np.tile(angles,len(radii));top_radius=np.repeat(radii,n_theta)
    top=np.column_stack((p.target_base_center[0]+top_radius*np.cos(top_theta),p.target_base_center[1]+top_radius*np.sin(top_theta),np.full_like(top_radius,p.target_height)))
    top=np.vstack((top,[p.target_base_center[0],p.target_base_center[1],p.target_height]));top_theta=np.r_[top_theta,0.];top_radius=np.r_[top_radius,0.]
    normals=np.vstack((side_normals,np.repeat([[0.,0.,1.]],len(top),axis=0)));points=np.vstack((side,top));kinds=np.r_[np.zeros(len(side),np.int8),np.ones(len(top),np.int8)]
    return solve.q3.SurfaceMesh(points,normals,kinds,np.r_[side_theta,top_theta],np.r_[side_z,top_radius],np.einsum("ij,ij->i",points,points),n_theta,n_levels)

def independent_states(x):
    p=solve.p();a=np.asarray(x.headings_rad);direction=np.column_stack((np.cos(a),np.sin(a),np.zeros(3)));tau=np.asarray(x.release_times_s);delta=np.asarray(x.fuse_delays_s);ets=tau+delta
    release=solve.ORIGINS+np.asarray(x.speeds_mps)[:,None]*tau[:,None]*direction
    explosion=solve.ORIGINS+np.asarray(x.speeds_mps)[:,None]*ets[:,None]*direction;explosion[:,2]=solve.ORIGINS[:,2]-.5*p.gravity*delta**2
    return release,explosion,ets

def missile_at(t):
    p=solve.p();direction=p.missile_initial/np.linalg.norm(p.missile_initial);return p.missile_initial-p.missile_speed*t*direction

def independent_margin(x,t,mesh):
    p=solve.p();missile=missile_at(t);release,explosion,ets=independent_states(x);worst=-math.inf;seen=False
    for index,point in enumerate(mesh.points):
        if mesh.kinds[index]==0:
            normal=mesh.point_normals[index]
            if float(np.dot(normal,missile-point)) < -1e-12:continue
        elif missile[2] < point[2]-1e-12:continue
        seen=True;best=math.inf;segment=point-missile;den=float(np.dot(segment,segment))
        for k in range(3):
            if not ets[k]-solve.TIME_TOL<=t<=min(ets[k]+p.smoke_duration,solve.t_m())+solve.TIME_TOL:continue
            center=explosion[k].copy();center[2]-=p.smoke_sink_speed*(t-ets[k]);w=center-missile;lam=min(1.,max(0.,float(np.dot(w,segment)/den)));distance=float(np.linalg.norm(center-(missile+lam*segment)));best=min(best,distance-p.smoke_radius)
        worst=max(worst,best)
    return worst if seen else math.inf

def e1(payload,x):
    p=solve.p();stored=payload["formal_best"]["decision"];release,explosion,ets=independent_states(x)
    stored_release=np.array([r["release_point_m"] for r in stored["platforms"]]);stored_explosion=np.array([r["explosion_point_m"] for r in stored["platforms"]]);stored_ets=np.array([r["explosion_time_s"] for r in stored["platforms"]])
    spatial=max(float(np.max(np.abs(release-stored_release))),float(np.max(np.abs(explosion-stored_explosion))));temporal=float(np.max(np.abs(ets-stored_ets)))
    require(np.allclose(solve.ORIGINS,[[17800,0,1800],[12000,1400,1400],[6000,-3000,700]],atol=0,rtol=0),"official origins mismatch")
    require(spatial<=TOL["space_m"] and temporal<=TOL["time_s"],"kinematic reconstruction mismatch")
    require(len(x.headings_rad)==len(x.speeds_mps)==len(x.release_times_s)==len(x.fuse_delays_s)==3,"not twelve variables")
    equal_release=solve.Decision(x.headings_rad,x.speeds_mps,(x.release_times_s[0],x.release_times_s[0],x.release_times_s[2]),x.fuse_delays_s,"no_q3_gap",0);solve.validate(equal_release)
    failures={}
    cases={
        "speed_below_70":solve.Decision(x.headings_rad,(69.,x.speeds_mps[1],x.speeds_mps[2]),x.release_times_s,x.fuse_delays_s,"invalid",0),
        "negative_release":solve.Decision(x.headings_rad,x.speeds_mps,(-.1,x.release_times_s[1],x.release_times_s[2]),x.fuse_delays_s,"invalid",0),
        "negative_explosion_height":solve.Decision(x.headings_rad,x.speeds_mps,x.release_times_s,(math.sqrt(2*solve.ORIGINS[0,2]/p.gravity)+1,x.fuse_delays_s[1],x.fuse_delays_s[2]),"invalid",0),
    }
    for name,candidate in cases.items():
        try:solve.validate(candidate);failures[name]=False
        except solve.Q4Error:failures[name]=True
    require(all(failures.values()),"invalid inputs did not fail")
    intervals=payload["formal_best"]["densified_intervals_s"];duration=sum(b-a for a,b in intervals)
    require(abs(duration-payload["formal_best"]["densified_duration_s"])<=TOL["time_s"],"union measure mismatch")
    claims=payload["claims"];require(not claims["three_clouds_simultaneously_active_required"] and not claims["positive_marginal_required"] and not claims["positive_pure_synergy_required"],"extra objective constraint present")
    require(abs(p.smoke_sink_speed-3.)<=1e-15 and abs(p.smoke_duration-20.)<=1e-15,"smoke semantics mismatch")
    return {"status":"pass","kinematic_max_spatial_residual_m":spatial,"kinematic_max_temporal_residual_s":temporal,"invalid_input_failures":failures,"equal_release_times_accepted":True,"smoke_sink_speed_mps":p.smoke_sink_speed,"smoke_duration_s":p.smoke_duration,"joint_interval_union_residual_s":abs(duration-payload["formal_best"]["densified_duration_s"])}

def excel_check(x):
    rows=[];sheet=load_workbook(ROOT/"output/q4/result2.xlsx",data_only=True).active
    for r in range(2,5):rows.append([sheet.cell(r,c).value for c in range(1,11)])
    require(tuple(str(r[0]) for r in rows)==solve.NAMES,"Excel platform order mismatch")
    headings=np.radians([r[1] for r in rows]);speeds=np.array([r[2] for r in rows]);release=np.array([r[3:6] for r in rows]);explosion=np.array([r[6:9] for r in rows]);tau=[];ets=[]
    for i in range(3):
        d=np.array([math.cos(headings[i]),math.sin(headings[i])]);tau.append(float((release[i,:2]-solve.ORIGINS[i,:2])@d/speeds[i]));ets.append(float((explosion[i,:2]-solve.ORIGINS[i,:2])@d/speeds[i]))
    delta=np.array(ets)-tau;time_res=max(float(np.max(np.abs(np.array(tau)-x.release_times_s))),float(np.max(np.abs(delta-x.fuse_delays_s))))
    rp,ep,_=independent_states(x);space_res=max(float(np.max(np.abs(release-rp))),float(np.max(np.abs(explosion-ep))))
    require(time_res<=TOL["time_s"] and space_res<=TOL["space_m"],"Excel reconstruction mismatch")
    return {"status":"pass","time_residual_s":time_res,"space_residual_m":space_res,"row_durations_s":[float(r[9]) for r in rows]}

def e2(payload,x):
    mesh=independent_mesh(256,17);intervals=payload["formal_best"]["densified_intervals_s"];probe=[]
    for a,b in intervals:probe.extend(np.linspace(a,b,5).tolist())
    for value in independent_states(x)[2]:probe.extend([float(value),float(value+1e-4)])
    probe=np.array(sorted(set(probe)));production=solve.joint_margins(x,probe,mesh);independent=np.array([independent_margin(x,float(t),mesh) for t in probe]);margin_diff=float(np.max(np.abs(production-independent)))
    require(margin_diff<=TOL["margin_m"],f"independent margin mismatch {margin_diff}")
    settings={"name":"e2_ultra","scan_step_s":.005,"n_theta":1024,"n_levels":65,"root_time_tolerance_s":1e-7,"margin_tolerance_m":2e-7,"surface_refinement":False}
    ultra_mesh=independent_mesh(1024,65);windows=[[max(0,a-.02),min(solve.t_m(),b+.02)] for a,b in intervals];ultra=solve.solve_intervals(x,settings,mesh=ultra_mesh,restrict=windows)
    change=float(ultra["effective_duration_s"]-payload["formal_best"]["densified_duration_s"]);require(abs(change)<=TOL["duration_s"],f"densification change {change}")
    dense=payload["verification"]["densified_joint"];root_max=max(abs(float(r["margin_m"])) for r in dense["root_residuals"]);require(root_max<=TOL["root_margin_m"] and dense["root_fallback_count"]==0,"production roots failed")
    excel=excel_check(x);individual=payload["verification"]["individuals"];require(payload["verification"]["resource_monotonicity_pass"],"resource monotonicity failed")
    return {"status":"pass","independent_probe_count":len(probe),"independent_margin_max_difference_m":margin_diff,"ultra_settings":settings,"ultra_surface_point_count":len(ultra_mesh.points),"ultra_intervals_s":ultra["intervals_s"],"ultra_duration_s":ultra["effective_duration_s"],"ultra_minus_formal_duration_s":change,"production_root_max_abs_margin_m":root_max,"production_root_fallback_count":dense["root_fallback_count"],"excel_roundtrip":excel,"individual_durations_s":[v["effective_duration_s"] for v in individual],"resource_monotonicity_pass":True}

def report(result):
    a=result["E1"];b=result["E2"]
    return f"""# Q4 E1/E2 独立验证报告

验证器角色：`auxiliary_validator`。本报告不参与选解，也不建立第二套正式算法。

## E1：实现符合性

- 状态：`{a['status']}`；三平台运动学最大空间残差 `{a['kinematic_max_spatial_residual_m']:.3e} m`，时间残差 `{a['kinematic_max_temporal_residual_s']:.3e} s`。
- FY1/FY2/FY3 使用官方三个不同起点和三组独立控制量；等投放时刻测试可通过，说明没有误带 Q3 的同机一秒间隔。
- 速度越界、负投放时刻和落地后起爆均触发确定性失败。
- 烟团下沉速度为 `3 m/s`、有效期为 `20 s`；正式时长与联合区间并集测度残差 `{a['joint_interval_union_residual_s']:.3e} s`。

## E2：数值可信度

- 状态：`{b['status']}`；独立逐点有限线段 oracle 共检查 `{b['independent_probe_count']}` 个时刻，最大余量差 `{b['independent_margin_max_difference_m']:.3e} m`。
- `1024×65` 表面、`0.005 s` 扫描的加密时长为 `{b['ultra_duration_s']:.12f} s`，相对正式加密结果变化 `{b['ultra_minus_formal_duration_s']:.3e} s`。
- 正式根最大绝对余量 `{b['production_root_max_abs_margin_m']:.3e} m`，fallback 数 `{b['production_root_fallback_count']}`。
- Excel 独立回读时间残差 `{b['excel_roundtrip']['time_residual_s']:.3e} s`、空间残差 `{b['excel_roundtrip']['space_residual_m']:.3e} m`。

## 证明边界

E1/E2 支持当前 Q4 实现与冻结规格同义，并支持 `11.610102348586 s` 工作解在预设离散加密下数值稳定。它们不证明十二维连续域全局最优，不完成 E3/E4，不支持概率稳健或现实部署有效，也不产生 Q5 结论。
"""

def run(output,report_path):
    started=time.perf_counter();payload,x=load();identities={k:sha(ROOT/k) for k in EXPECTED};require(all(identities[k]==v for k,v in EXPECTED.items()),"protected identity mismatch")
    result={"schema_version":"1.0","spec_id":"SPEC-Q4-VERIFY-1.0","result_id":"Q4-E1E2-20260828","status":"pass","method_role":"auxiliary_validator","generated_at":datetime.now(timezone.utc).isoformat(),"input_identity":{**identities,"src/q4/solve.py":sha(ROOT/"src/q4/solve.py"),"docs/q4_result.json":sha(ROOT/"docs/q4_result.json"),"output/q4/result2.xlsx":sha(ROOT/"output/q4/result2.xlsx")},"tolerances":TOL,"E1":e1(payload,x),"E2":e2(payload,x),"proof_boundary":{"global_optimality":False,"E3_E4_complete":False,"q5_solved":False},"software":{"python":platform.python_version(),"numpy":np.__version__,"scipy":scipy.__version__},"elapsed_s":time.perf_counter()-started}
    output.parent.mkdir(parents=True,exist_ok=True);output.write_text(json.dumps(solve.ready(result),ensure_ascii=False,indent=2),encoding="utf-8");report_path.write_text(report(result),encoding="utf-8");return result

def main():
    parser=argparse.ArgumentParser();parser.add_argument("--output",type=Path,required=True);parser.add_argument("--report",type=Path,required=True);a=parser.parse_args()
    try:r=run(a.output.resolve(),a.report.resolve())
    except VerificationFailure as e:print(json.dumps({"status":"fail","message":str(e)},ensure_ascii=False));return 2
    print(json.dumps({"status":"pass","E1":r["E1"]["status"],"E2":r["E2"]["status"],"elapsed_s":r["elapsed_s"]},ensure_ascii=False));return 0
if __name__=="__main__":raise SystemExit(main())
