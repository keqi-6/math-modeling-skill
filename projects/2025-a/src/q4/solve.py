"""Q4 single chain: physical events, strict joint screen, 12-D refine, verify."""
from __future__ import annotations

import argparse, hashlib, json, math, platform, shutil, sys, time
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable, Sequence

import numpy as np
import scipy
from openpyxl import load_workbook
from scipy.optimize import minimize

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
from src.q3 import selection_pilot as q3  # noqa: E402

SPEC_ID = "SPEC-Q4-WORKING-1.0"
RESULT_ID = "Q4-SPEC-Q4-WORKING-1.0-UNIFIED-SERVICE-EVENT-CHAIN"
HASHES = {
    "data/A题.pdf": "a37f6aad30b16ea09da9cb320ce194b12d3a28874008cf9ff19549f6c6079447",
    "data/附件/result2.xlsx": "c681d5e378538f71c77fca199a3ca8303a04dbcfc7bd95f870ae22f01ab69f91",
    "planning/26_q4_candidate_set.md": "f612eb8b92283bbb1ee533f639d467595cf73613450575216b47c6ae595eec33",
    "src/q3/selection_pilot.py": "33f67bee3a033ffebc48dd7a6275846c991977f6035554b75cd055de80430f64",
}
NAMES = ("FY1", "FY2", "FY3")
ORIGINS = np.array([[17800., 0., 1800.], [12000., 1400., 1400.], [6000., -3000., 700.]])
AGES = (0., .5, 1.5, 3., 6., 10., 15.)
FRACTIONS = (.10, .25, .40, .55, .70, .85, .95)
TIME_TOL, LENGTH_TOL = q3.TIME_TOL, q3.LENGTH_TOL
FAST = {"name":"fast_strict","scan_step_s":.08,"n_theta":32,"n_levels":5,"root_time_tolerance_s":4e-6,"margin_tolerance_m":4e-6,"surface_refinement":False}
SCREEN = {"name":"screen_strict","scan_step_s":.03,"n_theta":96,"n_levels":9,"root_time_tolerance_s":2e-6,"margin_tolerance_m":2e-6,"surface_refinement":False}
PRECISE = {"name":"precise_strict","scan_step_s":.01,"n_theta":256,"n_levels":17,"root_time_tolerance_s":5e-7,"margin_tolerance_m":5e-7,"surface_refinement":True}
DENSIFIED = {"name":"densified_strict","scan_step_s":.01,"n_theta":512,"n_levels":33,"root_time_tolerance_s":2e-7,"margin_tolerance_m":2e-7,"surface_refinement":True}

class Q4Error(RuntimeError):
    def __init__(self, code, message): super().__init__(message); self.code = code

@dataclass(frozen=True)
class Event:
    platform_index:int; heading_rad:float; speed_mps:float; observation_time_s:float
    explosion_time_s:float; sight_fraction:float; smoke_age_s:float
    release_time_s:float; fuse_delay_s:float; explosion_height_m:float
    horizontal_residual_m:float; vertical_residual_m:float; source:str

@dataclass(frozen=True)
class Decision:
    headings_rad:tuple[float,float,float]; speeds_mps:tuple[float,float,float]
    release_times_s:tuple[float,float,float]; fuse_delays_s:tuple[float,float,float]
    source:str; proposal_index:int

def sha(path):
    h=hashlib.sha256()
    with Path(path).open("rb") as f:
        for block in iter(lambda:f.read(1<<20),b""): h.update(block)
    return h.hexdigest()

def ready(v):
    if isinstance(v,np.ndarray): return v.tolist()
    if isinstance(v,np.generic): return v.item()
    if isinstance(v,dict): return {k:ready(x) for k,x in v.items()}
    if isinstance(v,(list,tuple)): return [ready(x) for x in v]
    return v

def verify_identities():
    got={k:sha(ROOT/k) for k in HASHES}
    bad={k:(HASHES[k],got[k]) for k in HASHES if HASHES[k]!=got[k]}
    if bad: raise Q4Error("Q4_IDENTITY_MISMATCH",str(bad))
    return got

def p(): return q3.base_parameters()
def t_m(): return q3.arrival_time()
def te(x): return np.asarray(x.release_times_s)+np.asarray(x.fuse_delays_s)

def release_points(x):
    a=np.asarray(x.headings_rad); d=np.column_stack((np.cos(a),np.sin(a),np.zeros(3)))
    return ORIGINS+np.asarray(x.speeds_mps)[:,None]*np.asarray(x.release_times_s)[:,None]*d

def explosion_points(x):
    a=np.asarray(x.headings_rad); d=np.column_stack((np.cos(a),np.sin(a),np.zeros(3)))
    z=ORIGINS+np.asarray(x.speeds_mps)[:,None]*te(x)[:,None]*d
    z[:,2]=ORIGINS[:,2]-.5*p().gravity*np.asarray(x.fuse_delays_s)**2
    return z

def constraint_margins(x):
    a=np.asarray(x.headings_rad); v=np.asarray(x.speeds_mps); tau=np.asarray(x.release_times_s); delta=np.asarray(x.fuse_delays_s)
    return {"heading":float(np.min(np.minimum(a,2*math.pi-a))),"speed_low":float(np.min(v-70)),"speed_high":float(np.min(140-v)),"release":float(np.min(tau)),"fuse":float(np.min(delta)),"height":float(np.min(ORIGINS[:,2]-.5*p().gravity*delta**2)),"arrival":float(np.min(t_m()-tau-delta))}

def validate(x):
    values=np.array([*x.headings_rad,*x.speeds_mps,*x.release_times_s,*x.fuse_delays_s])
    if not np.all(np.isfinite(values)) or any(not 0<=a<2*math.pi for a in x.headings_rad): raise Q4Error("Q4_CONSTRAINT_FAILURE","invalid values")
    m=constraint_margins(x)
    if min(v for k,v in m.items() if k!="heading") < -max(TIME_TOL,LENGTH_TOL): raise Q4Error("Q4_CONSTRAINT_FAILURE",str(m))

def decision_record(x):
    validate(x); rp,ep,et=release_points(x),explosion_points(x),te(x)
    return {**asdict(x),"platforms":[{"platform":NAMES[i],"initial_position_m":ORIGINS[i].tolist(),"heading_rad":x.headings_rad[i],"heading_deg":math.degrees(x.headings_rad[i])%360,"speed_mps":x.speeds_mps[i],"release_time_s":x.release_times_s[i],"fuse_delay_s":x.fuse_delays_s[i],"explosion_time_s":float(et[i]),"release_point_m":rp[i].tolist(),"explosion_point_m":ep[i].tolist()} for i in range(3)],"constraint_margins":constraint_margins(x)}

def cloud(x,k,times):
    z=np.repeat(explosion_points(x)[k].reshape(1,3),len(times),axis=0); z[:,2]-=p().smoke_sink_speed*(times-te(x)[k]); return z

def joint_margins(x,times,mesh,indices=(0,1,2)):
    validate(x); times=np.asarray(times,float); out=np.empty(len(times)); ets=te(x)
    batch=max(4,min(128,int(2_500_000/max(1,len(mesh.points)))))
    for left in range(0,len(times),batch):
        right=min(len(times),left+batch); block=times[left:right]
        missiles=np.vstack([q3.q1.missile_position(float(t),p()) for t in block]); visible=q3.visibility_mask(missiles,mesh)
        minimum=np.full((len(block),len(mesh.points)),np.inf)
        for k in indices:
            active=(block>=ets[k]-TIME_TOL)&(block<=min(ets[k]+p().smoke_duration,t_m())+TIME_TOL)
            if np.any(active):
                dist=q3._distance_to_segments(cloud(x,k,block),missiles,mesh); dist[~active]=np.inf; minimum=np.minimum(minimum,dist)
        m=minimum-p().smoke_radius; m[~visible]=-np.inf; r=np.max(m,axis=1); r[~np.any(visible,axis=1)]=np.inf; out[left:right]=r
    if np.any(np.isnan(out)): raise Q4Error("Q4_SURFACE_EVALUATOR_FAILURE","NaN")
    return out

def point_margin(x,t,point,indices):
    missile=q3.q1.missile_position(t,p()); best=math.inf; ets=te(x)
    for k in indices:
        if ets[k]-TIME_TOL<=t<=min(ets[k]+p().smoke_duration,t_m())+TIME_TOL:
            d,_=q3.q1.segment_distance(cloud(x,k,np.array([t]))[0],missile,point.reshape(1,3)-missile); best=min(best,float(d[0])-p().smoke_radius)
    return best

def refined_margin(x,t,mesh,indices):
    missile=q3.q1.missile_position(t,p()).reshape(1,3); visible=q3.visibility_mask(missile,mesh)[0]; minimum=np.full(len(mesh.points),np.inf)
    for k in indices:
        if te(x)[k]-TIME_TOL<=t<=min(te(x)[k]+p().smoke_duration,t_m())+TIME_TOL: minimum=np.minimum(minimum,q3._distance_to_segments(cloud(x,k,np.array([t])),missile,mesh)[0])
    vals=minimum-p().smoke_radius; vals[~visible]=-np.inf; grid=float(np.max(vals))
    if not math.isfinite(grid): return grid
    candidates=[grid]; da=2*math.pi/mesh.n_theta; dz=p().target_height/max(1,mesh.n_levels-1); dr=p().target_radius/mesh.n_levels
    for idx in np.argsort(-vals)[:4]:
        kind=int(mesh.kinds[idx]); a0=float(mesh.theta[idx]); s0=float(mesh.second[idx])
        if kind==0:
            bounds=[(a0-da,a0+da),(max(0,s0-dz),min(p().target_height,s0+dz))]
            make=lambda y:np.array([p().target_base_center[0]+p().target_radius*math.cos(y[0]),p().target_base_center[1]+p().target_radius*math.sin(y[0]),y[1]])
        else:
            bounds=[(a0-da,a0+da),(max(0,s0-dr),min(p().target_radius,s0+dr))]
            make=lambda y:np.array([p().target_base_center[0]+y[1]*math.cos(y[0]),p().target_base_center[1]+y[1]*math.sin(y[0]),p().target_height])
        def objective(y):
            point=make(y)
            if kind==0 and np.dot(np.array([math.cos(y[0]),math.sin(y[0]),0]),missile[0]-point)<-1e-9: return 1e3
            if kind==1 and missile[0,2]<p().target_height-1e-9: return 1e3
            value=point_margin(x,t,point,indices); return -value if math.isfinite(value) else 1e3
        opt=minimize(objective,np.array([a0,s0]),method="Powell",bounds=bounds,options={"maxiter":35,"maxfev":100,"xtol":1e-9,"ftol":1e-10})
        if opt.success and math.isfinite(float(opt.fun)): candidates.append(-float(opt.fun))
    return max(candidates)

def time_grid(x,step,indices,restrict=None):
    ets=te(x); events=np.array(sorted({*[float(ets[k]) for k in indices],*[float(min(ets[k]+p().smoke_duration,t_m())) for k in indices],t_m()}))
    windows=q3.merge_intervals([[ets[k],min(ets[k]+p().smoke_duration,t_m())] for k in indices])
    if restrict is not None:
        windows=q3.merge_intervals([[max(a,c),min(b,d)] for a,b in windows for c,d in restrict if min(b,d)>=max(a,c)])
    values=[]
    for a,b in windows:
        bounds=sorted({a,b,*[float(e) for e in events if a<e<b]})
        for left,right in zip(bounds[:-1],bounds[1:]): values.extend(np.linspace(left,right,max(1,math.ceil((right-left)/step))+1).tolist())
    return np.array(sorted(set(values))),events

def bisect_root(x,left,right,mesh,indices,settings):
    evaluator=(lambda t:refined_margin(x,t,mesh,indices)) if settings["surface_refinement"] else (lambda t:float(joint_margins(x,[t],mesh,indices)[0]))
    fl,fr=evaluator(left),evaluator(right)
    if not math.isfinite(fl) or not math.isfinite(fr) or fl*fr>0:
        return (left,fl,0,True) if fl<=0 else (right,fr,0,True)
    mid,fm=(left+right)/2,math.nan
    for it in range(1,101):
        mid=(left+right)/2; fm=evaluator(mid)
        if right-left<=settings["root_time_tolerance_s"] and abs(fm)<=settings["margin_tolerance_m"]: return mid,fm,it,False
        if fl==0:return left,fl,it,False
        if fr==0:return right,fr,it,False
        if fl*fm<=0:right,fr=mid,fm
        else:left,fl=mid,fm
    return mid,fm,100,True

def solve_intervals(x,settings,indices=(0,1,2),mesh=None,restrict=None):
    mesh=mesh or q3.surface_mesh(settings["n_theta"],settings["n_levels"]); times,events=time_grid(x,settings["scan_step_s"],indices,restrict)
    if len(times)==0:return {"settings":dict(settings),"surface_point_count":len(mesh.points),"scan_count":0,"intervals_s":[],"effective_duration_s":0.,"minimum_scan_margin_m":float("inf"),"root_residuals":[],"root_fallback_count":0}
    margins=joint_margins(x,times,mesh,indices); inside=margins<=0; intervals=[]; residuals=[]; fallbacks=0; i=0; gap=settings["scan_step_s"]*(1+1e-6)
    while i<len(times):
        if not inside[i]:i+=1;continue
        first=i
        while i+1<len(times) and inside[i+1] and times[i+1]-times[i]<=gap:i+=1
        last=i
        if first==0 or times[first]-times[first-1]>gap or np.any(np.isclose(times[first],events,atol=TIME_TOL,rtol=0)): entry=(float(times[first]),float(margins[first]),0,False)
        else: entry=bisect_root(x,float(times[first-1]),float(times[first]),mesh,indices,settings)
        if last==len(times)-1 or times[last+1]-times[last]>gap or np.any(np.isclose(times[last],events,atol=TIME_TOL,rtol=0)): exit_=(float(times[last]),float(margins[last]),0,False)
        else: exit_=bisect_root(x,float(times[last]),float(times[last+1]),mesh,indices,settings)
        fallbacks+=int(entry[3])+int(exit_[3]); intervals.append([entry[0],exit_[0]])
        residuals.extend([{"kind":"entry","time_s":entry[0],"margin_m":entry[1],"iterations":entry[2],"fallback":entry[3]},{"kind":"exit","time_s":exit_[0],"margin_m":exit_[1],"iterations":exit_[2],"fallback":exit_[3]}]);i+=1
    evaluator=(lambda t:refined_margin(x,float(t),mesh,indices)) if settings["surface_refinement"] else (lambda t:float(joint_margins(x,[float(t)],mesh,indices)[0]))
    try: tang=q3.q1.recover_unsampled_intervals(times,margins,evaluator,time_tolerance=settings["root_time_tolerance_s"],margin_tolerance=settings["margin_tolerance_m"],maximum_gap=gap)
    except q3.q1.Q1Error as e: raise Q4Error("Q4_TIME_EVENT_FAILURE",str(e)) from e
    intervals.extend(tang["recovered_intervals_s"]); merged=q3.merge_intervals(intervals)
    return {"settings":dict(settings),"surface_point_count":len(mesh.points),"scan_count":len(times),"intervals_s":merged,"effective_duration_s":q3.interval_measure(merged),"minimum_scan_margin_m":float(np.min(margins)),"root_residuals":residuals,"root_fallback_count":fallbacks,"local_minimum_checks":tang["checked_local_minima"],"recovered_unsampled_interval_count":len(tang["recovered_intervals_s"])}

def anchor_event(i,t,age,fraction):
    et=t-age
    if not 0<et<=t<t_m():return None
    missile=q3.q1.missile_position(t,p()); target=p().target_base_center+np.array([0.,0.,.5*p().target_height]); desired=missile+fraction*(target-missile)
    ez=float(desired[2]+p().smoke_sink_speed*age)
    if not 0<=ez<=ORIGINS[i,2]:return None
    delta=math.sqrt(max(0,2*(ORIGINS[i,2]-ez)/p().gravity));tau=et-delta
    if tau<-TIME_TOL:return None
    d=desired[:2]-ORIGINS[i,:2];speed=float(np.linalg.norm(d)/et)
    if not 70<=speed<=140:return None
    heading=float(math.atan2(d[1],d[0])%(2*math.pi));direction=np.array([math.cos(heading),math.sin(heading)])
    hr=float(np.linalg.norm(ORIGINS[i,:2]+speed*et*direction-desired[:2]));vr=abs(float(ORIGINS[i,2]-.5*p().gravity*delta**2-p().smoke_sink_speed*age-desired[2]))
    return Event(i,heading,speed,t,et,fraction,age,max(0.,tau),delta,ez,hr,vr,"physical_anchor")

def track_event(i,heading,speed,t):
    missile=q3.q1.missile_position(t,p());target=p().target_base_center+np.array([0.,0.,.5*p().target_height]); sight=target[:2]-missile[:2];direction=np.array([math.cos(heading),math.sin(heading)])
    matrix=np.column_stack((speed*direction,-sight))
    if abs(np.linalg.det(matrix))<=1e-8:return None
    et,fraction=np.linalg.solve(matrix,missile[:2]-ORIGINS[i,:2]);et,fraction=float(et),float(fraction);age=t-et
    if not 0<fraction<1 or not 0<=age<=p().smoke_duration or not 0<et<=t:return None
    desired=missile+fraction*(target-missile);ez=float(desired[2]+p().smoke_sink_speed*age)
    if not 0<=ez<=ORIGINS[i,2]:return None
    delta=math.sqrt(max(0,2*(ORIGINS[i,2]-ez)/p().gravity));tau=et-delta
    if tau<-TIME_TOL:return None
    hr=float(np.linalg.norm(ORIGINS[i,:2]+speed*et*direction-desired[:2]));vr=abs(float(ORIGINS[i,2]-.5*p().gravity*delta**2-p().smoke_sink_speed*age-desired[2]))
    if hr>1e-6 or vr>1e-7:return None
    return Event(i,heading%(2*math.pi),speed,t,et,fraction,age,max(0.,tau),delta,ez,hr,vr,"service_curve")

def proxy(event,times):
    direction=np.array([math.cos(event.heading_rad),math.sin(event.heading_rad),0.]);e=ORIGINS[event.platform_index]+event.speed_mps*event.explosion_time_s*direction;e[2]=event.explosion_height_m
    centers=np.repeat(e.reshape(1,3),len(times),axis=0);centers[:,2]-=p().smoke_sink_speed*(times-event.explosion_time_s);missiles=np.vstack([q3.q1.missile_position(float(t),p()) for t in times]);target=p().target_base_center+np.array([0.,0.,.5*p().target_height]);vec=target-missiles;den=np.einsum("ij,ij->i",vec,vec);off=centers-missiles;lam=np.clip(np.einsum("ij,ij->i",off,vec)/den,0,1);near=missiles+lam[:,None]*vec
    margin=np.linalg.norm(centers-near,axis=1)+lam*math.hypot(p().target_radius,.5*p().target_height)-p().smoke_radius;active=(times>=event.explosion_time_s-TIME_TOL)&(times<=min(event.explosion_time_s+p().smoke_duration,t_m())+TIME_TOL)
    return active&(margin<=0),margin

def track_key(e):return round(e.heading_rad/5e-4),round(e.speed_mps/.1)
def item_key(x):e=x["event"];return round(e.heading_rad/5e-4),round(e.speed_mps/.1),round(e.release_time_s/.02),round(e.fuse_delay_s/.02)

def diverse_select(values,limit,key):
    chosen={key(x):x for x in sorted(values,key=lambda x:(-x["mask_count"],x["minimum_proxy_margin_m"]))[:12]}
    ordered=sorted(values,key=lambda x:x["event"].observation_time_s)
    for group in np.array_split(np.arange(len(ordered)),max(1,limit-len(chosen))):
        if len(group):
            x=max((ordered[int(j)] for j in group),key=lambda y:(y["mask_count"],-y["minimum_proxy_margin_m"]));chosen[key(x)]=x
    return sorted(chosen.values(),key=lambda x:(-x["mask_count"],x["minimum_proxy_margin_m"],x["event"].observation_time_s))[:limit]

def generate_tracks(i,times):
    raw={};feasible=0
    for t in np.arange(.5,t_m(),.5):
        for age in AGES:
            if age>=t:continue
            for f in FRACTIONS:
                e=anchor_event(i,float(t),age,f)
                if e is None:continue
                feasible+=1;m,margin=proxy(e,times);x={"event":e,"mask_count":int(np.sum(m)),"minimum_proxy_margin_m":float(np.min(margin))};k=track_key(e)
                if k not in raw or (x["mask_count"],-x["minimum_proxy_margin_m"])>(raw[k]["mask_count"],-raw[k]["minimum_proxy_margin_m"]):raw[k]=x
    # Deterministic physical coverage guard.  The registered q-anchor grid can
    # miss a platform when the feasible sight fraction lies close to a LOS
    # endpoint.  These tracks still enter through the same 2x2 service-event
    # inversion; they are coverage seeds, not a competing solver.
    target_xy=p().target_base_center[:2]
    centres=(
        math.atan2(target_xy[1]-ORIGINS[i,1],target_xy[0]-ORIGINS[i,0]),
        math.atan2(p().missile_initial[1]-ORIGINS[i,1],p().missile_initial[0]-ORIGINS[i,0]),
    )
    for centre in centres:
        for offset in np.linspace(-.35,.35,15):
            heading=float((centre+offset)%(2*math.pi))
            for speed in np.linspace(70.,140.,8):
                best=None
                for t in np.arange(.25,t_m(),.25):
                    e=track_event(i,heading,float(speed),float(t))
                    if e is None:continue
                    mask,margin=proxy(e,times);candidate={"event":e,"mask_count":int(np.sum(mask)),"minimum_proxy_margin_m":float(np.min(margin))}
                    if best is None or (candidate["mask_count"],-candidate["minimum_proxy_margin_m"])>(best["mask_count"],-best["minimum_proxy_margin_m"]):best=candidate
                if best is not None:
                    feasible+=1;k=track_key(best["event"])
                    if k not in raw or (best["mask_count"],-best["minimum_proxy_margin_m"])>(raw[k]["mask_count"],-raw[k]["minimum_proxy_margin_m"]):raw[k]=best
    selected=diverse_select(list(raw.values()),48,lambda x:track_key(x["event"]))
    return selected,{"platform":NAMES[i],"raw_anchor_state_count":feasible,"unique_anchor_track_count":len(raw),"selected_track_count":len(selected)}

def event_library(i,tracks,times):
    raw={};curve_count=0
    for track in tracks:
        a=track["event"]
        # The physical anchor that created a track is an exact member of that
        # track's service curve.  Keep it explicitly so the later uniform
        # observation grid cannot discard the generating state through a
        # determinant/round-off edge case.
        curve=[a]
        curve.extend(
            e for e in (
                track_event(i,a.heading_rad,a.speed_mps,float(t))
                for t in np.arange(.25,t_m(),.25)
            ) if e is not None
        )
        for e in curve:
            curve_count+=1;m,margin=proxy(e,times);x={"event":e,"mask":m,"mask_count":int(np.sum(m)),"minimum_proxy_margin_m":float(np.min(margin))};k=item_key(x)
            if k not in raw or (x["mask_count"],-x["minimum_proxy_margin_m"])>(raw[k]["mask_count"],-raw[k]["minimum_proxy_margin_m"]):raw[k]=x
    selected=diverse_select(list(raw.values()),36,item_key)
    if not selected:raise Q4Error("Q4_EVENT_LIBRARY_EMPTY",NAMES[i])
    return selected,{"platform":NAMES[i],"service_curve_event_count":curve_count,"unique_event_count":len(raw),"selected_event_count":len(selected)}

def decision_from_events(events,source,index):
    if tuple(e.platform_index for e in events)!=(0,1,2):raise Q4Error("Q4_RESOURCE_COMBINATION_FAILURE","one event per platform")
    x=Decision(tuple(e.heading_rad for e in events),tuple(e.speed_mps for e in events),tuple(e.release_time_s for e in events),tuple(e.fuse_delay_s for e in events),source,index);validate(x);return x

def combine(libs,times):
    out=[];ordinal=0
    for a in libs[0]:
        for b in libs[1]:
            pair=a["mask"]|b["mask"]
            for c in libs[2]:
                union=pair|c["mask"];active=times[union];out.append({"events":(a["event"],b["event"],c["event"]),"proxy_union_duration_s":float(np.sum(union)*.1),"proxy_span_s":float(active[-1]-active[0]) if len(active) else 0.,"minimum_individual_proxy_margin_m":min(a["minimum_proxy_margin_m"],b["minimum_proxy_margin_m"],c["minimum_proxy_margin_m"]),"ordinal":ordinal});ordinal+=1
    out.sort(key=lambda x:(x["proxy_union_duration_s"],x["proxy_span_s"],-x["minimum_individual_proxy_margin_m"]),reverse=True);return out[:24]

def to_unit(x):
    z=[]
    for i in range(3):
        cap=min(math.sqrt(2*ORIGINS[i,2]/p().gravity),t_m()-x.release_times_s[i]);z.extend([x.headings_rad[i]/(2*math.pi),(x.speeds_mps[i]-70)/70,x.release_times_s[i]/t_m(),x.fuse_delays_s[i]/cap if cap>0 else 0])
    return np.clip(z,0,1)

def from_unit(z,source,index):
    z=np.clip(np.asarray(z,float),0,1);a=[];v=[];tau=[];delta=[]
    for i in range(3):
        o=4*i;aa=float(2*math.pi*z[o]%(2*math.pi));vv=float(70+70*z[o+1]);tt=float(t_m()*z[o+2]);cap=min(math.sqrt(2*ORIGINS[i,2]/p().gravity),t_m()-tt);dd=float(z[o+3]*max(0,cap));a.append(aa);v.append(vv);tau.append(tt);delta.append(dd)
    x=Decision(tuple(a),tuple(v),tuple(tau),tuple(delta),source,index);validate(x);return x

def refine(start,mesh,ordinal):
    z=to_unit(start);current=start;result=solve_intervals(current,FAST,mesh=mesh);cache={tuple(np.round(z,12)):(current,result)};calls=1;stages=[]
    def get(y):
        nonlocal calls
        y=np.clip(y,0,1);k=tuple(np.round(y,12))
        if k not in cache:
            x=from_unit(y,f"q4_pattern_{ordinal}",ordinal);cache[k]=(x,solve_intervals(x,FAST,mesh=mesh));calls+=1
        return cache[k]
    for radius in (.015,.006):
        accepted=0
        for j in range(12):
            options=[(current,result,z.copy())]
            for sign in (-1,1):
                y=z.copy();y[j]=np.clip(y[j]+sign*radius,0,1);x,r=get(y);options.append((x,r,y))
            best=max(options,key=lambda z:(z[1]["effective_duration_s"],-z[1]["minimum_scan_margin_m"]))
            if (best[1]["effective_duration_s"],-best[1]["minimum_scan_margin_m"])>(result["effective_duration_s"],-result["minimum_scan_margin_m"]):current,result,z=best;accepted+=1
        stages.append({"radius":radius,"accepted_coordinate_moves":accepted,"duration_s":result["effective_duration_s"]})
    return {"start_decision":decision_record(start),"chosen_decision":decision_record(current),"chosen_fast":result,"function_evaluations":calls,"stages":stages,"_decision":current}

def expanded(intervals,pad):return [[max(0,a-pad),min(t_m(),b+pad)] for a,b in intervals]
def diagnostics(x,precise_mesh,dense_mesh):
    joint=solve_intervals(x,PRECISE,mesh=precise_mesh);individual=[solve_intervals(x,PRECISE,(i,),precise_mesh) for i in range(3)];union=q3.merge_intervals(v for item in individual for v in item["intervals_s"])
    dense=solve_intervals(x,DENSIFIED,mesh=dense_mesh,restrict=expanded(joint["intervals_s"],.02)) if joint["intervals_s"] else {"settings":DENSIFIED,"intervals_s":[],"effective_duration_s":0.,"root_fallback_count":0,"scan_count":0}
    leave=[]
    for omitted in range(3):
        ids=tuple(i for i in range(3) if i!=omitted);r=solve_intervals(x,PRECISE,ids,precise_mesh);leave.append({"omitted_platform":NAMES[omitted],"duration_s":r["effective_duration_s"],"intervals_s":r["intervals_s"],"marginal_contribution_s":joint["effective_duration_s"]-r["effective_duration_s"]})
    probes=[]
    for a,b in dense["intervals_s"]:
        ts=np.array([a,(a+b)/2,b]);probes.append({"times_s":ts.tolist(),"margins_m":joint_margins(x,ts,dense_mesh).tolist()})
    return {"precise_joint":joint,"densified_joint":dense,"individuals":individual,"individual_union_intervals_s":union,"individual_union_duration_s":q3.interval_measure(union),"leave_one_out":leave,"resource_monotonicity_pass":joint["effective_duration_s"]+1e-8>=max(v["effective_duration_s"] for v in individual),"densified_change_s":dense["effective_duration_s"]-joint["effective_duration_s"],"boundary_probes":probes}

def write_excel(x,durations,path,mesh):
    path.parent.mkdir(parents=True,exist_ok=True);shutil.copyfile(ROOT/"data/附件/result2.xlsx",path);book=load_workbook(path);sheet=book[book.sheetnames[0]];rec=decision_record(x)
    for i,r in enumerate(rec["platforms"]):
        vals=[r["platform"],r["heading_deg"],r["speed_mps"],*r["release_point_m"],*r["explosion_point_m"],float(durations[i])]
        for col,val in enumerate(vals,1):sheet.cell(row=i+2,column=col,value=val if col==1 else float(val))
    book.save(path);sheet=load_workbook(path,data_only=True).active;rows=[[sheet.cell(row=r,column=c).value for c in range(1,11)] for r in range(2,5)]
    if tuple(str(r[0]) for r in rows)!=NAMES:raise Q4Error("Q4_EXCEL_ROUNDTRIP_FAILURE","labels")
    headings=np.radians([r[1] for r in rows]);speeds=np.array([r[2] for r in rows]);release=np.array([r[3:6] for r in rows]);explosion=np.array([r[6:9] for r in rows]);tau=[];ets=[]
    for i in range(3):
        d=np.array([math.cos(headings[i]),math.sin(headings[i])]);tau.append(float((release[i,:2]-ORIGINS[i,:2])@d/speeds[i]));ets.append(float((explosion[i,:2]-ORIGINS[i,:2])@d/speeds[i]))
    y=Decision(tuple(float(a%(2*math.pi)) for a in headings),tuple(float(v) for v in speeds),tuple(tau),tuple(b-a for a,b in zip(tau,ets)),"excel_roundtrip",0);validate(y)
    joint=solve_intervals(y,PRECISE,mesh=mesh);ind=[solve_intervals(y,PRECISE,(i,),mesh) for i in range(3)];res=float(np.max(np.abs(np.array([r[9] for r in rows])-[v["effective_duration_s"] for v in ind])))
    if res>1e-6:raise Q4Error("Q4_EXCEL_ROUNDTRIP_FAILURE","durations")
    return {"path":path.relative_to(ROOT).as_posix(),"sha256":sha(path),"status":"pass","reconstructed_decision":decision_record(y),"recomputed_joint":joint,"row_duration_max_residual_s":res}

def report(result,result_path):
    best=result["formal_best"];lines=["# Q4 单一主链实现报告","",f"- 规格：`{SPEC_ID}`",f"- 严格联合完整遮蔽时间：`{best['densified_duration_s']:.9f} s`",f"- 连续时间区间：`{best['densified_intervals_s']}`",f"- 机器可读结果：`{result_path.relative_to(ROOT).as_posix()}`","","## 唯一求解路线","","三平台物理服务事件反演 → 单机事件库压缩 → 每平台各取一事件组合 → 完整可见圆柱与有限视线段严格筛选 → 十二原变量精修 → 加密复核与 Excel 回读。历史构造没有作为并列算法运行。","","## 三平台决策",""]
    for r in best["decision"]["platforms"]:lines.append(f"- {r['platform']}：航向 {r['heading_deg']:.9f}°，速度 {r['speed_mps']:.9f} m/s，投放 {r['release_time_s']:.9f} s，引信 {r['fuse_delay_s']:.9f} s。")
    lines.extend(["","## 验证","",f"- 资源单调性：`{result['verification']['resource_monotonicity_pass']}`",f"- 精确层到加密层变化：`{result['verification']['densified_change_s']:.9e} s`",f"- Excel 回读：`{result['excel_roundtrip']['status']}`","","该结果是冻结预算下的可复核工作解，不宣称全局最优；唯一正式目标不附加三弹同时生效、正边际或纯协同约束。"]) ;return "\n".join(lines)+"\n"

def run(output,report_path,excel):
    started=time.perf_counter();identities=verify_identities();proxy_times=np.arange(0,t_m()+.05,.1);tracks=[];track_stats=[]
    for i in range(3):x,s=generate_tracks(i,proxy_times);tracks.append(x);track_stats.append(s)
    libs=[];lib_stats=[]
    for i in range(3):x,s=event_library(i,tracks[i],proxy_times);libs.append(x);lib_stats.append(s)
    combos=combine(libs,proxy_times)
    if not combos:raise Q4Error("Q4_RESOURCE_COMBINATION_FAILURE","empty")
    fast_mesh=q3.surface_mesh(FAST["n_theta"],FAST["n_levels"]);strict=[]
    for rank,c in enumerate(combos,1):x=decision_from_events(c["events"],"unified_service_combination",rank);strict.append({"decision":x,"fast":solve_intervals(x,FAST,mesh=fast_mesh),"proxy":c})
    strict.sort(key=lambda z:(z["fast"]["effective_duration_s"],-z["fast"]["minimum_scan_margin_m"]),reverse=True)
    refinements=[refine(v["decision"],fast_mesh,i) for i,v in enumerate(strict[:4],1)];refined=[{"decision":v["_decision"],"fast":v["chosen_fast"],"refinement":v} for v in refinements];refined.sort(key=lambda z:(z["fast"]["effective_duration_s"],-z["fast"]["minimum_scan_margin_m"]),reverse=True)
    screen_mesh=q3.surface_mesh(SCREEN["n_theta"],SCREEN["n_levels"]);screened=[{**v,"screen":solve_intervals(v["decision"],SCREEN,mesh=screen_mesh)} for v in refined[:3]];screened.sort(key=lambda z:(z["screen"]["effective_duration_s"],-z["screen"]["minimum_scan_margin_m"]),reverse=True)
    precise_mesh=q3.surface_mesh(PRECISE["n_theta"],PRECISE["n_levels"]);precise=[{**v,"precise":solve_intervals(v["decision"],PRECISE,mesh=precise_mesh)} for v in screened[:2]];precise.sort(key=lambda z:(z["precise"]["effective_duration_s"],-z["precise"]["minimum_scan_margin_m"]),reverse=True);winner=precise[0]
    dense_mesh=q3.surface_mesh(DENSIFIED["n_theta"],DENSIFIED["n_levels"]);diag=diagnostics(winner["decision"],precise_mesh,dense_mesh);excel_check=write_excel(winner["decision"],[v["effective_duration_s"] for v in diag["individuals"]],excel,precise_mesh)
    result={"schema_version":"1.0","spec_id":SPEC_ID,"result_id":RESULT_ID,"status":"working_verified","generated_at":datetime.now(timezone.utc).isoformat(),"objective":"strict joint complete-obscuration interval-union duration of three smoke clouds","formal_best":{"decision":decision_record(winner["decision"]),"precise_duration_s":diag["precise_joint"]["effective_duration_s"],"precise_intervals_s":diag["precise_joint"]["intervals_s"],"densified_duration_s":diag["densified_joint"]["effective_duration_s"],"densified_intervals_s":diag["densified_joint"]["intervals_s"]},"single_chain":{"track_generation":track_stats,"event_libraries":lib_stats,"strict_combination_count":len(combos),"strict_records":[{"decision":decision_record(v["decision"]),"fast":v["fast"],"proxy":{k:x for k,x in v["proxy"].items() if k!="events"}} for v in strict],"refinements":[{k:x for k,x in v.items() if k!="_decision"} for v in refinements]},"verification":diag,"excel_roundtrip":excel_check,"identity":identities,"precision_layers":{"fast":FAST,"screen":SCREEN,"precise":PRECISE,"densified":DENSIFIED},"software":{"python":platform.python_version(),"numpy":np.__version__,"scipy":scipy.__version__},"elapsed_s":time.perf_counter()-started,"claims":{"global_optimality":False,"three_clouds_simultaneously_active_required":False,"positive_marginal_required":False,"positive_pure_synergy_required":False,"q3_files_modified":False}}
    output.parent.mkdir(parents=True,exist_ok=True);output.write_text(json.dumps(ready(result),ensure_ascii=False,indent=2),encoding="utf-8");report_path.parent.mkdir(parents=True,exist_ok=True);report_path.write_text(report(result,output),encoding="utf-8");return result

def main():
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument("--output",type=Path,default=ROOT/"docs/q4_result.json");parser.add_argument("--report",type=Path,default=ROOT/"docs/q4_implementation_report.md");parser.add_argument("--excel",type=Path,default=ROOT/"output/q4/result2.xlsx");a=parser.parse_args()
    try:r=run(a.output.resolve(),a.report.resolve(),a.excel.resolve())
    except Q4Error as e:print(json.dumps({"status":"fail","code":e.code,"message":str(e)},ensure_ascii=False));return 2
    print(json.dumps({"status":"pass","result_id":r["result_id"],"densified_duration_s":r["formal_best"]["densified_duration_s"],"elapsed_s":r["elapsed_s"]},ensure_ascii=False));return 0
if __name__=="__main__":raise SystemExit(main())
